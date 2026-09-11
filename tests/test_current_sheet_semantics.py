import openpyxl

from action_tracker.excel.reader import ZH_MAP
from action_tracker.excel.writer import _update_or_append_current
from action_tracker.monitor.sku_monitor import SkuStatus
from action_tracker.orchestrator.daily import _build_current_records


def _status(sku: str, status: str) -> SkuStatus:
    present = status != "MISSING_FIRST"
    return SkuStatus(
        sku=sku,
        canonical_id=f"ACT{sku.zfill(7)}",
        status=status,
        source_flag="BOTH" if present else "NONE",
        sitemap_present=present,
        listing_present=present,
        was_yesterday=True,
        ever_seen=True,
        first_seen="2026-08-12",
        previous_status="ACTIVE",
        missing_count=0 if present else 1,
        event=None,
    )


def test_current_records_exclude_missing_and_mark_new_current():
    baseline = {
        "1001": {"sku": "1001", "canonical_id": "ACT0001001", "status": "CURRENT"},
        "1002": {"sku": "1002", "canonical_id": "ACT0001002", "status": "CURRENT"},
    }
    updated = {
        "2002": {"sku": "2002", "canonical_id": "ACT0002002", "status": None},
    }
    statuses = {
        "1001": _status("1001", "ACTIVE"),
        "1002": _status("1002", "MISSING_FIRST"),
        "2002": _status("2002", "NEW"),
    }

    current = _build_current_records(
        baseline, updated, statuses, {"1001", "2002"}, "2026-08-13")

    assert set(current) == {"1001", "2002"}
    assert {record["status"] for record in current.values()} == {"CURRENT"}
    assert {record["last_seen"] for record in current.values()} == {"2026-08-13"}
    assert current["2002"]["first_seen"] == "2026-08-12"


def test_writer_removes_rows_not_in_current_presence():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01_SKU_ZH_CURRENT"
    ws.append(["Canonical_ID", "SKU", "当前状态"])
    ws.append(["ACT0001001", "1001", "CURRENT"])
    ws.append(["ACT0001002", "1002", "MISSING_FIRST"])

    records = {
        "1001": {"canonical_id": "ACT0001001", "sku": "1001", "status": "CURRENT"},
        "2002": {"canonical_id": "ACT0002002", "sku": "2002", "status": "CURRENT"},
    }
    _update_or_append_current(wb, ws.title, ZH_MAP, records, set(ZH_MAP.values()))

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert {str(row[1]) for row in rows} == {"1001", "2002"}
    assert {row[2] for row in rows} == {"CURRENT"}


def test_writer_uses_header_positions_for_zh_current_and_preserves_existing_row():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01_SKU_ZH_CURRENT"
    ws.append(["SKU", "中文品名", "当前状态", "Canonical_ID", "中文描述"])
    ws.append(["1001", "原名称", "CURRENT", "ACT0001001", "已有说明"])
    ws.cell(2, 2).font = openpyxl.styles.Font(bold=True)

    records = {
        "1001": {"canonical_id": "ACT0001001", "sku": "1001", "status": "CURRENT"},
        "2002": {"canonical_id": "ACT0002002", "sku": "2002", "status": "CURRENT"},
    }
    _update_or_append_current(wb, ws.title, ZH_MAP, records, set(ZH_MAP.values()))

    assert ws.max_row == 3
    assert ws.cell(2, 1).value == "1001"
    assert ws.cell(2, 2).value == "原名称"
    assert ws.cell(2, 2).font.bold is True
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:E3"
    assert ws.cell(2, 5).alignment.wrap_text is True
