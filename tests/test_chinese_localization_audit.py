from __future__ import annotations

import importlib.util
import csv
import json
from pathlib import Path
import sys

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
AUDIT_PATH = ROOT / "scripts" / "audit_chinese_localization.py"
SPEC = importlib.util.spec_from_file_location("audit_chinese_localization", AUDIT_PATH)
assert SPEC is not None and SPEC.loader is not None
AUDIT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(AUDIT)


def test_structured_business_flags_take_precedence_over_remarks():
    row = {
        "business_tag_new": "true",
        "business_tag_promotion": False,
        "business_tag_sustainable": "0",
        "remarks": "",
    }

    flags, evidence, invalid = AUDIT._business_tag_flags(row, chinese=False)

    assert flags == {"new": True, "promotion": False, "sustainable": False}
    assert evidence == "STRUCTURED_FIELDS"
    assert invalid == ()


def test_confirmed_brand_loader_excludes_unreviewed_phrases(tmp_path: Path):
    path = tmp_path / "brands.csv"
    path.write_text(
        "canonical_name,aliases_es,review_status\n"
        "GS27,GS27|GS 27,HUMAN_REVIEWED\n"
        "SL-300,SL-300,NEEDS_HUMAN_REVIEW\n",
        encoding="utf-8",
    )

    result = AUDIT._load_confirmed_brand_phrases(path)

    assert result["status"] == "LOADED"
    assert set(result["phrases"]) == {"GS27", "GS 27"}
    assert result["phrase_count"] == 2
    assert result["sha256"] == AUDIT._sha256_file(path)


def test_qa_log_is_deterministic_and_kept_separate_from_product_fields(tmp_path: Path):
    report = {
        "source_sha256": "a" * 64,
        "target_sha256": "b" * 64,
        "issues": [{
            "sku": "1001", "field": "description", "qa_layer": "L5",
            "code": "APPROVED_TERM_UNSUPPORTED_TARGET_REVIEW", "review_only": True,
            "source": "Caja de almacenamiento", "candidate": "木质收纳盒",
            "evidence": {"source_term": "madera"},
        }, {
            "sku": "1001", "field": "description", "qa_layer": "L5",
            "code": "APPROVED_TERM_UNSUPPORTED_TARGET_REVIEW", "review_only": True,
            "source": "Caja de almacenamiento", "candidate": "木质收纳盒",
            "evidence": {"source_term": "madera"},
        }],
    }

    path = AUDIT.write_qa_log(tmp_path / "QA_LOG.csv", report)
    first_bytes = path.read_bytes()
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    AUDIT.write_qa_log(path, report)

    assert path.read_bytes() == first_bytes
    assert len(rows) == 2
    assert len({row["qa_id"] for row in rows}) == 2
    assert rows[0]["qa_id"].startswith("qa-")
    assert rows[0]["sku"] == "1001"
    assert rows[0]["qa_layer"] == "L5"
    assert rows[0]["review_only"] == "True"
    assert "issue_json" in rows[0]
    assert "备注" not in rows[0]


def test_audit_manifest_binds_qa_log_hash_and_row_count(tmp_path: Path, monkeypatch):
    report = {
        "source_sha256": "a" * 64, "target_sha256": "b" * 64, "issue_count": 1,
        "issues": [{"sku": "1001", "field": "description", "qa_layer": "L5", "code": "TEST"}],
    }
    monkeypatch.setattr(AUDIT, "audit_workbooks", lambda *args: report)
    monkeypatch.setattr(sys, "argv", [
        "audit_chinese_localization.py", "--source", str(tmp_path / "source.xlsx"),
        "--target", str(tmp_path / "target.xlsx"), "--output-dir", str(tmp_path / "out"),
        "--sheet", "Catalog",
    ])

    assert AUDIT.main() == 1

    manifest = json.loads((tmp_path / "out/localization_audit.json").read_text(encoding="utf-8"))
    qa_log_path = tmp_path / "out/QA_LOG.csv"
    assert manifest["qa_log_artifact"] == {
        "path": "QA_LOG.csv", "sha256": AUDIT._sha256_file(qa_log_path), "row_count": 1,
    }
    detail_queue_path = tmp_path / "out/DETAIL_RULE_REVIEW.csv"
    detail_queue_manifest = json.loads(
        (tmp_path / "out/DETAIL_RULE_REVIEW.csv.manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["detail_rule_review_queue_artifact"]["auto_apply"] is False
    assert detail_queue_manifest["sha256"] == AUDIT._sha256_file(detail_queue_path)
    assert detail_queue_manifest["row_count"] == 0
    category_queue_path = tmp_path / "out/CATEGORY_RULE_REVIEW.csv"
    category_queue_manifest = json.loads(
        (tmp_path / "out/CATEGORY_RULE_REVIEW.csv.manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["category_rule_review_queue_artifact"]["auto_apply"] is False
    assert category_queue_manifest["sha256"] == AUDIT._sha256_file(category_queue_path)
    assert category_queue_manifest["row_count"] == 0
    term_queue_path = tmp_path / "out/TERM_RULE_REVIEW.csv"
    term_queue_manifest = json.loads(
        (tmp_path / "out/TERM_RULE_REVIEW.csv.manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["term_rule_review_queue_artifact"]["auto_apply"] is False
    assert term_queue_manifest["sha256"] == AUDIT._sha256_file(term_queue_path)
    assert term_queue_manifest["row_count"] == 0


def test_category_rule_review_queue_groups_scoped_findings_without_approval(tmp_path: Path):
    issues = [
        {
            "sku": "1001", "field": "cat2", "code": "CATEGORY_MAPPING_UNAVAILABLE",
            "source_cat1": "Viajes", "source": "Maletas", "candidate_cat1": "旅行用品",
            "candidate": "行李箱",
        },
        {
            "sku": "1002", "field": "cat2", "code": "CATEGORY_MAPPING_UNAVAILABLE",
            "source_cat1": "Viajes", "source": "Maletas", "candidate_cat1": "旅行用品",
            "candidate": "行李箱",
        },
        {
            "sku": "1003", "field": "cat2", "code": "CATEGORY_MAPPING_UNAVAILABLE",
            "source_cat1": "Oficina y papelería", "source": "Maletas",
            "candidate_cat1": "办公文具", "candidate": "行李箱",
        },
        {
            "sku": "1004", "field": "cat2", "code": "CATEGORY_MAPPING_MISMATCH",
            "source_cat1": "Multimedia", "source": "Cascos de teléfono",
            "candidate_cat1": "数码影音", "candidate": "手机配件", "expected": "手机耳机",
        },
    ]

    rows = AUDIT._category_rule_review_rows(issues)
    assert len(rows) == 3
    viajes = next(row for row in rows if row["source_cat1"] == "Viajes")
    assert viajes["occurrences"] == 2
    assert viajes["sku_count"] == 2
    assert viajes["single_observed_target_value"] == "行李箱"
    assert viajes["candidate_is_approved"] is False
    assert any(row["source_cat1"] == "Oficina y papelería" for row in rows)
    mismatch = next(row for row in rows if row["review_kind"] == "CATEGORY_MAPPING_MISMATCH")
    assert json.loads(mismatch["approved_values"]) == ["手机耳机"]

    path, manifest_path = AUDIT.write_category_rule_review_queue(
        tmp_path / "CATEGORY_RULE_REVIEW.csv", rows,
        source_sha256="a" * 64, target_sha256="b" * 64,
        category_dictionary_sha256="c" * 64,
    )
    first = path.read_bytes()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    AUDIT.write_category_rule_review_queue(
        path, rows, source_sha256="a" * 64, target_sha256="b" * 64,
        category_dictionary_sha256="c" * 64,
    )
    assert path.read_bytes() == first
    assert manifest["row_count"] == 3
    assert manifest["sha256"] == AUDIT._sha256_file(path)
    assert manifest["review_only"] is True
    assert manifest["auto_apply"] is False
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        output = list(csv.DictReader(handle))
    assert all(row["candidate_is_approved"] == "False" for row in output)


def test_term_rule_review_queue_groups_by_term_and_field_without_promoting_candidates(tmp_path: Path):
    issues = [
        {
            "sku": "1001", "code": "APPROVED_TERM_CROSS_FIELD_REVIEW", "review_only": True,
            "source_term": "madera", "term_type": "material", "source_field": "details",
            "field": "description", "approved_translation": "木质",
            "source_value": "Material: Madera", "candidate": "木质边框",
        },
        {
            "sku": "1002", "code": "APPROVED_TERM_CROSS_FIELD_REVIEW", "review_only": True,
            "source_term": "madera", "term_type": "material", "source_field": "details",
            "field": "description", "approved_translation": "木质",
            "source_value": "Material: Madera", "candidate": "木质相框",
        },
        {
            "sku": "1003", "code": "APPROVED_TERM_UNSUPPORTED_TARGET_REVIEW", "review_only": True,
            "source_term": "madera", "term_type": "material", "source_field": "",
            "field": "description", "approved_translation": "木质",
            "source_value": "", "candidate": "木质收纳盒",
        },
        {
            "sku": "1004", "code": "TERM_FORBIDDEN_TRANSLATION", "review_only": False,
            "source_term": "madera", "term_type": "material", "source_field": "",
            "field": "name", "approved_translation": "木质", "forbidden_candidate": "木头",
            "source_value": "", "candidate": "木头收纳盒",
        },
    ]

    rows = AUDIT._term_rule_review_rows(issues)
    assert len(rows) == 3
    cross = next(row for row in rows if row["review_kind"] == "APPROVED_TERM_CROSS_FIELD_REVIEW")
    assert cross["occurrences"] == 2
    assert cross["sku_count"] == 2
    assert cross["candidate_variant_count"] == 2
    assert cross["review_only"] is True
    assert cross["candidate_is_approved"] is False
    forbidden = next(row for row in rows if row["review_kind"] == "TERM_FORBIDDEN_TRANSLATION")
    assert forbidden["review_only"] is False

    path, manifest_path = AUDIT.write_term_rule_review_queue(
        tmp_path / "TERM_RULE_REVIEW.csv", rows,
        source_sha256="a" * 64, target_sha256="b" * 64,
        term_dictionary_sha256="c" * 64,
    )
    first = path.read_bytes()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    AUDIT.write_term_rule_review_queue(
        path, rows, source_sha256="a" * 64, target_sha256="b" * 64,
        term_dictionary_sha256="c" * 64,
    )
    assert path.read_bytes() == first
    assert manifest["row_count"] == 3
    assert manifest["sha256"] == AUDIT._sha256_file(path)
    assert manifest["review_only"] is True
    assert manifest["auto_apply"] is False
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        output = list(csv.DictReader(handle))
    assert all(row["candidate_is_approved"] == "False" for row in output)


def test_workbook_reader_loads_operational_badge_boolean_columns(tmp_path: Path):
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Catalog"
    sheet.append([
        "SKU", "name", "cat1", "cat2", "spec", "description", "details",
        "action_new_badge", "promotion_active", "sustainable_badge",
    ])
    sheet.append(["1001", "Product", "Home", "Cleaning", "", "", "", True, False, True])
    workbook.save(path)

    _, rows = AUDIT._read_sheet(path, "Catalog")

    assert rows["1001"]["business_tag_new"] is True
    assert rows["1001"]["business_tag_promotion"] is False
    assert rows["1001"]["business_tag_sustainable"] is True


def test_workbook_reader_supports_actual_action_export_headers(tmp_path: Path):
    source_path = tmp_path / "spanish_export.xlsx"
    target_path = tmp_path / "chinese_export.xlsx"
    source_book = Workbook()
    source_sheet = source_book.active
    source_sheet.title = "商品全量"
    source_sheet.append([
        "图片", "编号", "标题（西语）", "分类1（西语）", "分类2（西语）", "规格（西语）",
        "折后价", "原价", "单价", "描述（西语）", "产品详情（西语）", "图片链接", "商品链接", "备注",
    ])
    source_sheet.append(["", "1001", "Caja", "Hogar", "Almacenaje", "1 unidad",
                         2.5, 3.0, "", "Descripción", "Material: plástico", "", "", ""])
    source_book.save(source_path)

    target_book = Workbook()
    target_sheet = target_book.active
    target_sheet.title = "商品全量"
    target_sheet.append([
        "图片", "编号", "中文品名", "一级类目", "二级类目", "中文规格",
        "折后价", "原价", "单价", "描述", "产品详情", "图片链接", "商品链接", "备注",
    ])
    target_sheet.append(["", "1001", "收纳盒", "家居布置", "收纳", "1件",
                         2.5, 3.0, "", "收纳物品", "材质：塑料", "", "", ""])
    target_book.save(target_path)

    _, source_rows = AUDIT._read_sheet(source_path, "商品全量")
    _, target_rows = AUDIT._read_sheet(target_path, "商品全量")

    assert source_rows["1001"]["name"] == "Caja"
    assert source_rows["1001"]["cat2"] == "Almacenaje"
    assert source_rows["1001"]["details"] == "Material: plástico"
    assert target_rows["1001"]["name"] == "收纳盒"
    assert target_rows["1001"]["cat1"] == "家居布置"
    assert target_rows["1001"]["spec"] == "1件"


def test_business_tag_audit_compares_structured_source_to_exported_remarks():
    source = {
        "business_tag_new": True,
        "business_tag_promotion": False,
        "business_tag_sustainable": False,
    }
    target = {"remarks": "Estado: CURRENT"}
    issues = []

    report = AUDIT._audit_business_tags({"1001"}, {"1001": source}, {"1001": target}, issues)

    assert report["checked_skus"] == 1
    assert report["source_evidence"] == {"STRUCTURED_FIELDS": 1}
    assert [issue["code"] for issue in issues] == ["BUSINESS_TAG_DROPPED"]
    assert issues[0]["source"] == "Nuevo"


def test_chinese_business_tag_audit_checks_rendered_remarks_not_auxiliary_flags():
    target = {
        "business_tag_new": True,
        "business_tag_promotion": False,
        "business_tag_sustainable": False,
        "remarks": "Estado: CURRENT",
    }

    flags, evidence, invalid = AUDIT._business_tag_flags(target, chinese=True)

    assert flags == {"new": False, "promotion": False, "sustainable": False}
    assert evidence == "REMARKS_EXPORT"
    assert invalid == ()


def test_official_tag_identity_is_audited_separately_from_business_badges():
    official_tags = "Una opción más sostenible | Nuevo | -21%"
    source = {
        "raw_tags": official_tags,
        "business_tag_new": True,
        "business_tag_promotion": False,
        "business_tag_sustainable": True,
    }
    target = {
        "remarks": (
            "在售状态：在售；新品；可持续；官网官方标签："
            "Una opción más sostenible（官网可持续选择标签） | Nuevo（官网新品标签）"
        ),
    }
    issues = []

    business_report = AUDIT._audit_business_tags({"1001"}, {"1001": source}, {"1001": target}, issues)
    official_report = AUDIT._audit_official_labels({"1001"}, {"1001": source}, {"1001": target}, issues)

    assert business_report["checked_skus"] == 1
    assert official_report == {"status": "CHECKED", "checked_skus": 1, "unavailable_skus": 0}
    assert issues == []

    dropped = {"remarks": "在售状态：在售；新品；可持续"}
    official_report = AUDIT._audit_official_labels({"1001"}, {"1001": source}, {"1001": dropped}, issues)
    assert official_report["checked_skus"] == 1
    assert [item["code"] for item in issues] == ["OFFICIAL_LABEL_DROPPED"]


def test_invalid_structured_business_flag_does_not_fall_back_to_remarks():
    source = {
        "business_tag_new": "unknown",
        "remarks": "Nuevo",
    }
    target = {"remarks": "Nuevo"}
    issues = []

    report = AUDIT._audit_business_tags({"1001"}, {"1001": source}, {"1001": target}, issues)

    assert report["checked_skus"] == 0
    assert report["unavailable_skus"] == 1
    assert [issue["code"] for issue in issues] == ["BUSINESS_TAG_SOURCE_INVALID"]
    assert "business_tag_promotion" in issues[0]["evidence"]
    assert "business_tag_sustainable" in issues[0]["evidence"]


def test_approved_material_term_without_canonical_phrase_routes_to_review(tmp_path: Path):
    term_path = tmp_path / "terms.csv"
    fields = ["term_es", "term_zh", "term_type", "forbidden_zh", "review_status"]
    with term_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows([
            {"term_es": "madera", "term_zh": "木质", "term_type": "material", "review_status": "APPROVED"},
            {"term_es": "bambú", "term_zh": "竹制", "term_type": "material", "review_status": "APPROVED"},
        ])
    source = {
        "1001": {"description": "Cuchara de madera y bambú"},
        "1002": {"description": "Cuchara de madera"},
    }
    target = {
        "1001": {"description": "木材和竹制勺子"},
        "1002": {"description": "木质勺子"},
    }
    issues = []

    report = AUDIT._audit_term_dictionary(source, target, term_path, issues)

    assert report["review_only_suggestions"] == 1
    assert [(item["sku"], item["source_term"], item["approved_translation"])
            for item in issues if item["code"] == "APPROVED_TERM_CANONICAL_ABSENT_REVIEW"] == [
                ("1001", "madera", "木质"),
            ]
    assert all(item["review_only"] for item in issues)
    policy = AUDIT._load_qa_policy(ROOT / "config/stage5/chinese_gold_qa_policy.json")
    assert AUDIT._issue_layer("APPROVED_TERM_CANONICAL_ABSENT_REVIEW", policy) == "L4"


def test_generic_glossary_terms_are_scoped_outside_structured_details(tmp_path: Path):
    term_path = tmp_path / "terms.csv"
    fields = ["term_es", "term_zh", "term_type", "forbidden_zh", "review_status"]
    with term_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows([
            {"term_es": "diferentes", "term_zh": "不同款式", "term_type": "spec", "review_status": "APPROVED"},
            {"term_es": "madera", "term_zh": "木质", "term_type": "material", "review_status": "APPROVED"},
        ])
    source = {
        "1001": {"spec": "diferentes colores", "description": "Cuchara de madera", "details": ""},
    }
    target = {
        "1001": {"spec": "多种颜色", "description": "木材勺子", "details": ""},
    }
    issues = []
    report = AUDIT._audit_term_dictionary(source, target, term_path, issues)
    assert report["review_policy_id"] == "ACTION_APPROVED_TERM_REVIEW_V1"
    assert not any(
        item.get("source_term") == "diferentes"
        and item.get("code") == "APPROVED_TERM_CANONICAL_ABSENT_REVIEW"
        for item in issues
    )
    assert any(
        item.get("source_term") == "madera"
        and item.get("code") == "APPROVED_TERM_CANONICAL_ABSENT_REVIEW"
        for item in issues
    )


def test_issue_deduplication_keeps_candidate_variants_and_source_identity():
    issues = [
        {"sku": "1001", "field": "details", "code": "X", "source": "Material: madera", "candidate": "木材"},
        {"sku": "1001", "field": "details", "code": "X", "source": "Material: madera", "candidate": "木质"},
        {"sku": "1001", "field": "details", "code": "X", "source": "Color: azul", "candidate": "蓝色"},
    ]
    rows, stats = AUDIT._dedupe_issues(issues)
    assert stats == {"raw_issue_count": 3, "unique_issue_count": 2, "duplicate_issue_count": 1}
    material = next(row for row in rows if row["source"] == "Material: madera")
    assert material["duplicate_count"] == 1
    assert material["candidate_variants"] == ["木质"]


def test_approved_term_cross_field_finding_is_classified_as_l5():
    policy = AUDIT._load_qa_policy(ROOT / "config/stage5/chinese_gold_qa_policy.json")

    assert AUDIT._issue_layer("APPROVED_TERM_CROSS_FIELD_REVIEW", policy) == "L5"
    assert AUDIT._issue_layer("APPROVED_TERM_UNSUPPORTED_TARGET_REVIEW", policy) == "L5"


def test_gold_claim_flag_requires_boolean_and_complete_layer_coverage(tmp_path: Path):
    source_path = ROOT / "config/stage5/chinese_gold_qa_policy.json"
    policy = json.loads(source_path.read_text(encoding="utf-8"))
    policy_path = tmp_path / "gold_policy.json"

    policy["gold_claim_allowed"] = True
    policy_path.write_text(json.dumps(policy), encoding="utf-8")
    try:
        AUDIT._load_qa_policy(policy_path)
    except ValueError as exc:
        assert str(exc) == "GOLD_QA_POLICY_GOLD_COVERAGE_INCOMPLETE"
    else:
        raise AssertionError("partial QA coverage was allowed to enable Gold claim")

    policy["layers"] = [{**layer, "coverage": "COMPLETE"} for layer in policy["layers"]]
    policy_path.write_text(json.dumps(policy), encoding="utf-8")
    try:
        AUDIT._load_qa_policy(policy_path)
    except ValueError as exc:
        assert str(exc) == "GOLD_QA_POLICY_COMPLETE_LAYER_HAS_UNCOVERED:L1"
    else:
        raise AssertionError("a COMPLETE layer with uncovered checks was accepted")

    policy["layers"] = [
        {
            **layer, "coverage": "COMPLETE", "not_covered": [],
            "covered_checks": [] if layer["id"] == "L1" else layer["covered_checks"],
        }
        for layer in policy["layers"]
    ]
    policy_path.write_text(json.dumps(policy), encoding="utf-8")
    try:
        AUDIT._load_qa_policy(policy_path)
    except ValueError as exc:
        assert str(exc) == "GOLD_QA_POLICY_LAYER_HAS_NO_COVERED_CHECKS:L1"
    else:
        raise AssertionError("a COMPLETE layer with no checks was accepted")

    policy = json.loads(source_path.read_text(encoding="utf-8"))
    policy["gold_claim_allowed"] = True
    policy["layers"] = [
        {**layer, "coverage": "COMPLETE", "not_covered": []}
        for layer in policy["layers"]
    ]
    policy_path.write_text(json.dumps(policy), encoding="utf-8")
    assert AUDIT._load_qa_policy(policy_path)["gold_claim_allowed"] is True

    policy["gold_claim_allowed"] = "true"
    policy_path.write_text(json.dumps(policy), encoding="utf-8")
    try:
        AUDIT._load_qa_policy(policy_path)
    except ValueError as exc:
        assert str(exc) == "GOLD_QA_POLICY_GOLD_CLAIM_FLAG_INVALID"
    else:
        raise AssertionError("non-boolean Gold claim flag was accepted")


def test_full_audit_reports_glossary_backed_cross_field_movement(tmp_path: Path):
    term_path = tmp_path / "terms.csv"
    fields = ["term_es", "term_zh", "term_type", "forbidden_zh", "review_status"]
    with term_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerow({
            "term_es": "madera", "term_zh": "木质", "term_type": "material",
            "review_status": "APPROVED",
        })
    source = {
        "1001": {"name": "Caja", "spec": "", "description": "Caja de almacenamiento", "details": "Material: madera"},
    }
    target = {
        "1001": {"name": "收纳盒", "spec": "", "description": "木质收纳盒", "details": "材质：木质"},
    }
    issues = []

    report = AUDIT._audit_term_dictionary(source, target, term_path, issues)

    findings = [item for item in issues if item["code"] == "APPROVED_TERM_CROSS_FIELD_REVIEW"]
    assert report["review_only_suggestions"] == 1
    assert len(findings) == 1
    assert findings[0]["sku"] == "1001"
    assert findings[0]["source_field"] == "details"
    assert findings[0]["field"] == "description"
    assert findings[0]["review_only"] is True


def test_full_audit_reports_glossary_target_without_any_source_field_support(tmp_path: Path):
    term_path = tmp_path / "terms.csv"
    fields = ["term_es", "term_zh", "term_type", "forbidden_zh", "review_status"]
    with term_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerow({
            "term_es": "madera", "term_zh": "木质", "term_type": "material",
            "review_status": "APPROVED",
        })
    source = {
        "1001": {"name": "Caja", "spec": "", "description": "Caja de almacenamiento", "details": "Material: plástico"},
    }
    target = {
        "1001": {"name": "收纳盒", "spec": "", "description": "木质收纳盒", "details": "材质：塑料"},
    }
    issues = []

    report = AUDIT._audit_term_dictionary(source, target, term_path, issues)

    findings = [item for item in issues if item["code"] == "APPROVED_TERM_UNSUPPORTED_TARGET_REVIEW"]
    assert report["review_only_suggestions"] == 1
    assert len(findings) == 1
    assert findings[0]["sku"] == "1001"
    assert findings[0]["field"] == "description"
    assert findings[0]["unsupported_target_is_not_proof_of_hallucination"] is True


def test_full_audit_excludes_conflicting_approved_term_mappings(tmp_path: Path):
    term_path = tmp_path / "terms.csv"
    fields = ["term_es", "term_zh", "term_type", "forbidden_zh", "review_status"]
    with term_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows([
            {"term_es": "madera", "term_zh": "木质", "term_type": "material", "review_status": "APPROVED"},
            {"term_es": "madera", "term_zh": "木材", "term_type": "material", "review_status": "HUMAN_REVIEWED"},
        ])
    source = {"1001": {"name": "Caja", "spec": "madera", "description": "", "details": ""}}
    target = {"1001": {"name": "收纳盒", "spec": "木质", "description": "", "details": ""}}
    issues = []

    report = AUDIT._audit_term_dictionary(source, target, term_path, issues)

    assert report["status"] == "LOADED"
    assert report["ambiguous_normalized_terms"] == 1
    assert report["approved_terms"] == 0
    finding = next(item for item in issues if item["code"] == "TERM_DICTIONARY_AMBIGUOUS")
    assert finding["source_term"] == "madera"
    assert set(finding["approved_targets"]) == {"木材", "木质"}
    assert not any(item["code"] == "APPROVED_TERM_CANONICAL_ABSENT_REVIEW" for item in issues)

    from action_tracker.translation.term_resolver import resolve_exact_spec_term
    rows = (
        {"term_es": "madera", "term_zh": "木质", "term_type": "material", "review_status": "APPROVED"},
        {"term_es": "madera", "term_zh": "木材", "term_type": "material", "review_status": "HUMAN_REVIEWED"},
    )
    assert resolve_exact_spec_term("madera", rows) is None


def test_full_workbook_audit_reports_certification_and_negation_loss(tmp_path: Path):
    source_path = tmp_path / "source.xlsx"
    target_path = tmp_path / "target.xlsx"
    category_path = tmp_path / "categories.csv"
    terms_path = tmp_path / "terms.csv"
    source_book, target_book = Workbook(), Workbook()
    source_sheet, target_sheet = source_book.active, target_book.active
    source_sheet.title = target_sheet.title = "Catalog"
    source_sheet.append(["SKU", "Título", "Categoría 1", "Categoría 2", "Especificación", "Descripción", "Detalles"])
    source_sheet.append([
        "1001", "Caja", "Hogar", "Almacenaje", "20 unidades",
        "Papel FSC® certificado, sin BPA", "Cantidad: 22 unidades; Número del artículo: 1001",
    ])
    target_sheet.append(["编号", "标题", "分类1", "分类2", "规格", "描述", "产品详情"])
    target_sheet.append(["1001", "收纳盒", "家居布置", "收纳", "20件", "纸张", "数量：22件；商品编号：1001"])
    source_book.save(source_path)
    target_book.save(target_path)
    with category_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["cat1_es", "cat2_es", "cat1_zh", "cat2_zh", "review_status"])
        writer.writeheader()
        writer.writerow({
            "cat1_es": "Hogar", "cat2_es": "Almacenaje", "cat1_zh": "家居布置",
            "cat2_zh": "收纳", "review_status": "APPROVED",
        })
    with terms_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["term_es", "term_zh", "term_type", "forbidden_zh", "review_status"])
        writer.writeheader()
        writer.writerow({
            "term_es": "irrelevante", "term_zh": "无关词", "term_type": "attribute",
            "forbidden_zh": "", "review_status": "APPROVED",
        })

    report = AUDIT.audit_workbooks(
        source_path, target_path, "Catalog", ROOT / "config/stage5/detail_terminology_rules.json",
        category_dictionary_path=category_path, term_dictionary_path=terms_path,
    )
    assert report["historical_regression"]["status"] == "PASS"
    assert report["historical_regression"]["matched_sku_count"] == 1

    findings = {
        issue["code"] for issue in report["issues"]
        if issue.get("sku") == "1001" and issue.get("field") == "description"
    }
    assert {"CERTIFICATION_DROPPED", "NEGATION_DROPPED"}.issubset(findings)
    source_conflict = next(
        issue for issue in report["issues"]
        if issue.get("sku") == "1001" and issue.get("field") == "source"
        and issue.get("code") == "NUMERIC_CONFLICT"
    )
    conflict_evidence = json.loads(source_conflict["evidence"])
    assert conflict_evidence[0]["cross_field_comparisons"][0]["left_values"] == [
        {"value": 20.0, "unit": "unidad"}
    ]
    assert conflict_evidence[0]["cross_field_comparisons"][0]["right_values"] == [
        {"value": 22.0, "unit": "unidad"}
    ]

    source_hallucination = tmp_path / "source_hallucination.xlsx"
    target_hallucination = tmp_path / "target_hallucination.xlsx"
    source_book, target_book = Workbook(), Workbook()
    source_sheet, target_sheet = source_book.active, target_book.active
    source_sheet.title = target_sheet.title = "Catalog"
    source_sheet.append(["SKU", "Título", "Categoría 1", "Categoría 2", "Especificación", "Descripción", "Detalles"])
    source_sheet.append(["1001", "Caja", "Hogar", "Almacenaje", "1 unidad", "Papel para manualidades", ""])
    target_sheet.append(["编号", "标题", "分类1", "分类2", "规格", "描述", "产品详情"])
    target_sheet.append(["1001", "收纳盒", "家居布置", "收纳", "1件", "纸张，FSC认证，不含BPA，无硫酸盐", ""])
    source_book.save(source_hallucination)
    target_book.save(target_hallucination)
    hallucination_report = AUDIT.audit_workbooks(
        source_hallucination, target_hallucination, "Catalog",
        ROOT / "config/stage5/detail_terminology_rules.json",
        category_dictionary_path=category_path, term_dictionary_path=terms_path,
    )
    hallucination_findings = {
        issue["code"] for issue in hallucination_report["issues"]
        if issue.get("sku") == "1001" and issue.get("field") == "description"
    }
    assert {"CERTIFICATION_HALLUCINATED", "NEGATION_HALLUCINATED", "UNSUPPORTED_NEGATIVE_ATTRIBUTE"}.issubset(hallucination_findings)


def test_full_workbook_audit_enforces_compatibility_platform_retention(tmp_path: Path):
    source_path, target_path = tmp_path / "source.xlsx", tmp_path / "target.xlsx"
    category_path, terms_path = tmp_path / "categories.csv", tmp_path / "terms.csv"
    source_book, target_book = Workbook(), Workbook()
    source_sheet, target_sheet = source_book.active, target_book.active
    source_sheet.title = target_sheet.title = "Catalog"
    source_sheet.append(["SKU", "Título", "Categoría 1", "Categoría 2", "Especificación", "Descripción", "Detalles"])
    source_sheet.append(["1001", "Funda compatible con Nintendo Switch", "Hogar", "Almacenaje", "", "", ""])
    source_sheet.append(["1002", "Funda compatible con Nintendo Switch", "Hogar", "Almacenaje", "", "", ""])
    target_sheet.append(["编号", "标题", "分类1", "分类2", "规格", "描述", "产品详情"])
    target_sheet.append(["1001", "兼容Switch的保护壳", "家居布置", "收纳", "", "", ""])
    target_sheet.append(["1002", "游戏保护壳", "家居布置", "收纳", "", "", ""])
    source_book.save(source_path)
    target_book.save(target_path)

    with category_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["cat1_es", "cat2_es", "cat1_zh", "cat2_zh", "review_status"])
        writer.writeheader()
        writer.writerow({
            "cat1_es": "Hogar", "cat2_es": "Almacenaje", "cat1_zh": "家居布置",
            "cat2_zh": "收纳", "review_status": "APPROVED",
        })
    with terms_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["term_es", "term_zh", "term_type", "forbidden_zh", "review_status"])
        writer.writeheader()

    report = AUDIT.audit_workbooks(
        source_path, target_path, "Catalog", ROOT / "config/stage5/detail_terminology_rules.json",
        category_dictionary_path=category_path, term_dictionary_path=terms_path,
    )
    platform_findings = [
        issue for issue in report["issues"]
        if issue.get("field") == "name" and issue.get("code") == "TECH_TOKEN_DROPPED"
    ]
    assert [issue["sku"] for issue in platform_findings] == ["1002"]


def test_full_workbook_audit_routes_unmapped_closed_enum_value_to_l4(tmp_path: Path):
    source_path, target_path = tmp_path / "source.xlsx", tmp_path / "target.xlsx"
    category_path, terms_path = tmp_path / "categories.csv", tmp_path / "terms.csv"
    source_book, target_book = Workbook(), Workbook()
    source_sheet, target_sheet = source_book.active, target_book.active
    source_sheet.title = target_sheet.title = "Catalog"
    source_sheet.append(["SKU", "Título", "Categoría 1", "Categoría 2", "Especificación", "Descripción", "Detalles"])
    source_sheet.append(["1001", "Batería", "Hogar", "Almacenaje", "", "", "Tipo de batería: Alcalina"])
    source_sheet.append(["1002", "Batería", "Hogar", "Almacenaje", "", "", "Tipo de batería: Alcalina"])
    target_sheet.append(["编号", "标题", "分类1", "分类2", "规格", "描述", "产品详情"])
    target_sheet.append(["1001", "电池", "家居布置", "收纳", "", "", "电池类型：碱性电池"])
    target_sheet.append(["1002", "电池", "家居布置", "收纳", "", "", "电池类型：碱性"])
    source_book.save(source_path)
    target_book.save(target_path)

    with category_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["cat1_es", "cat2_es", "cat1_zh", "cat2_zh", "review_status"])
        writer.writeheader()
        writer.writerow({
            "cat1_es": "Hogar", "cat2_es": "Almacenaje", "cat1_zh": "家居布置",
            "cat2_zh": "收纳", "review_status": "APPROVED",
        })
    with terms_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["term_es", "term_zh", "term_type", "forbidden_zh", "review_status"])
        writer.writeheader()

    report = AUDIT.audit_workbooks(
        source_path, target_path, "Catalog", ROOT / "config/stage5/detail_terminology_rules.json",
        category_dictionary_path=category_path, term_dictionary_path=terms_path,
    )
    finding = next(
        item for item in report["issues"]
        if item.get("sku") == "1001" and item.get("code") == "DETAIL_SOURCE_VALUE_UNMAPPED_REVIEW:0"
    )
    assert finding["qa_layer"] == "L4"
    assert report["detail_rule_coverage"]["source_key_value_occurrences"] == 2
    assert report["detail_rule_coverage"]["source_key_value_occurrences_with_rule"] == 0
    variants = next(
        item for item in report["issues"]
        if item.get("code") == "SOURCE_VALUE_TRANSLATION_VARIANTS"
    )
    evidence = json.loads(variants["evidence"])
    assert {
        (row["target_value"], row["occurrences"], row["sku_examples"][0])
        for row in evidence["observed_translations"]
    } == {("碱性电池", 1, "1001"), ("碱性", 1, "1002")}


def test_detail_rule_coverage_respects_product_context_per_occurrence():
    source_rows = {
        "1001": {"name": "Libro de ficción", "details": "Género: Ficción"},
        "1002": {"name": "Decoración de hogar", "details": "Género: Ficción"},
    }
    target_rows = {
        "1001": {"details": "题材：虚构类"},
        "1002": {"details": "类别：虚构类"},
    }
    rules = {
        "key_translations": {},
        "contextual_key_translations": [{
            "source_key": "género", "target_key": "题材", "name_es_any": ["libro"],
        }],
        "numeric_suffix_rules": [],
        "value_translations": [{
            "source_key": "género", "source_value": "ficción", "target_value": "虚构类",
            "name_es_any": ["libro"],
        }],
    }

    coverage = AUDIT._detail_rule_coverage(source_rows, target_rows, rules)

    assert coverage["key_occurrences"] == 2
    assert coverage["key_occurrences_with_rule"] == 1
    assert coverage["key_occurrence_coverage"] == 0.5
    assert coverage["source_key_value_occurrences"] == 2
    assert coverage["source_key_value_occurrences_with_rule"] == 1
    assert coverage["source_key_value_occurrence_coverage"] == 0.5
    assert coverage["unmapped_source_key_value_pair_count"] == 1
    unmapped = coverage["unmapped_source_key_value_pairs"][0]
    assert unmapped["uncovered_occurrences"] == 1
    assert unmapped["uncovered_sku_examples"] == ["1002"]
    queue_rows = coverage["review_queue_rows"]
    assert {row["review_kind"] for row in queue_rows} == {"SOURCE_KEY", "SOURCE_VALUE"}
    key_row = next(row for row in queue_rows if row["review_kind"] == "SOURCE_KEY")
    value_row = next(row for row in queue_rows if row["review_kind"] == "SOURCE_VALUE")
    assert key_row["single_observed_candidate"] == ""
    assert value_row["single_observed_candidate"] == ""
    assert key_row["candidate_is_approved"] is False


def test_detail_rule_coverage_reports_translation_variants_even_when_rule_covers_all():
    source_rows = {
        "1001": {"name": "Libro de ficción", "details": "Género: Ficción"},
        "1002": {"name": "Libro de ficción", "details": "Género: Ficción"},
    }
    target_rows = {
        "1001": {"details": "题材：虚构类"},
        "1002": {"details": "题材：小说"},
    }
    rules = {
        "key_translations": {},
        "contextual_key_translations": [{
            "source_key": "género", "target_key": "题材", "name_es_any": ["libro"],
        }],
        "numeric_suffix_rules": [],
        "value_translations": [{
            "source_key": "género", "source_value": "ficción", "target_value": "虚构类",
            "candidate_value_any": ["虚构类", "小说"], "name_es_any": ["libro"],
        }],
    }

    coverage = AUDIT._detail_rule_coverage(source_rows, target_rows, rules)

    assert coverage["source_key_value_occurrence_coverage"] == 1.0
    assert coverage["source_key_value_pairs_with_translation_variants"] == 1
    finding = coverage["source_key_value_translation_variants"][0]
    assert finding["source_key"] == "genero"
    assert finding["source_value"] == "ficcion"
    assert finding["rule_coverage_status"] == "FULL_RULE_COVERAGE"
    assert finding["translation_variant_count"] == 2
    assert {tuple(row["sku_examples"]) for row in finding["variants"]} == {("1001",), ("1002",)}
    assert finding["review_only"] is True
    assert finding["variation_is_not_proof_of_error"] is True
    queue_row = next(
        row for row in coverage["review_queue_rows"]
        if row["review_kind"] == "SOURCE_VALUE"
    )
    assert "TRANSLATION_VARIANTS" in queue_row["review_reasons"]
    assert queue_row["single_observed_candidate"] == ""


def test_detail_rule_review_queue_is_deterministic_and_source_bound(tmp_path: Path):
    rows = [{
        "review_kind": "SOURCE_VALUE", "source_key_normalized": "genero",
        "source_value_normalized": "ficcion", "occurrences": 2,
        "rule_covered_occurrences": 1, "uncovered_occurrences": 1, "sku_count": 2,
        "sku_examples": '["1001", "1002"]',
        "observed_target_candidates": '[{"target_value":"虚构类"},{"target_value":"小说"}]',
        "single_observed_candidate": "", "review_reasons": "TRANSLATION_VARIANTS",
        "candidate_is_approved": True,
    }]
    path, manifest_path = AUDIT.write_detail_rule_review_queue(
        tmp_path / "DETAIL_RULE_REVIEW.csv", rows,
        source_sha256="a" * 64, target_sha256="b" * 64, rules_sha256="c" * 64,
    )
    first = path.read_bytes()
    path, _ = AUDIT.write_detail_rule_review_queue(
        path, rows, source_sha256="a" * 64, target_sha256="b" * 64, rules_sha256="c" * 64,
    )
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        result = list(csv.DictReader(handle))
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))

    assert path.read_bytes() == first
    assert len(result) == 1
    assert result[0]["review_id"].startswith("detail-")
    assert result[0]["candidate_is_approved"] == "False"
    assert result[0]["source_sha256"] == "a" * 64
    assert manifest["sha256"] == AUDIT._sha256_file(path)
    assert manifest["review_only"] is True
    assert manifest["auto_apply"] is False


def test_detail_rule_review_queue_is_not_truncated_to_summary_limit(tmp_path: Path):
    details_es = "; ".join(f"Source key {index}: source value {index}" for index in range(305))
    details_zh = "; ".join(f"中文键 {index}：中文值 {index}" for index in range(305))
    coverage = AUDIT._detail_rule_coverage(
        {"1001": {"name": "", "details": details_es}},
        {"1001": {"details": details_zh}},
        {"key_translations": {}, "contextual_key_translations": [],
         "numeric_suffix_rules": [], "value_translations": []},
    )

    assert coverage["unmapped_source_key_count"] == 305
    assert coverage["unmapped_source_key_value_pair_count"] == 305
    assert coverage["review_queue_row_count"] == 610
    assert len(coverage["review_queue_rows"]) == 610
    assert len(coverage["unmapped_source_keys"]) == 200
    assert len(coverage["unmapped_source_key_value_pairs"]) == 300

    path, manifest_path = AUDIT.write_detail_rule_review_queue(
        tmp_path / "DETAIL_RULE_REVIEW.csv", coverage["review_queue_rows"],
        source_sha256="a" * 64, target_sha256="b" * 64, rules_sha256="c" * 64,
    )
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["row_count"] == 610


def test_detail_pair_alignment_mismatch_is_explicitly_queued_and_classified(tmp_path: Path):
    source_path, target_path = tmp_path / "source.xlsx", tmp_path / "target.xlsx"
    headers = ["SKU", "name", "cat1", "cat2", "spec", "description", "details"]
    source_book = Workbook()
    source_sheet = source_book.active
    source_sheet.title = "Catalog"
    source_sheet.append(headers)
    source_sheet.append(["1001", "Producto", "Hogar", "Casa", "", "", "Material: plástico; Color: azul"])
    source_book.save(source_path)
    target_book = Workbook()
    target_sheet = target_book.active
    target_sheet.title = "Catalog"
    target_sheet.append(headers)
    target_sheet.append(["1001", "产品", "家居布置", "家居", "", "", "材质：塑料"])
    target_book.save(target_path)
    category_path = tmp_path / "categories.csv"
    category_path.write_text("cat1_es,cat2_es,cat1_zh,cat2_zh,review_status\n", encoding="utf-8")
    terms_path = tmp_path / "terms.csv"
    terms_path.write_text("term_es,term_zh,term_type,forbidden_zh,review_status\n", encoding="utf-8")

    report = AUDIT.audit_workbooks(
        source_path, target_path, "Catalog", ROOT / "config/stage5/detail_terminology_rules.json",
        category_dictionary_path=category_path, term_dictionary_path=terms_path,
    )

    finding = next(item for item in report["issues"] if item["code"] == "DETAIL_PAIR_ALIGNMENT_REVIEW")
    assert finding["qa_layer"] == "L4"
    assert json.loads(finding["evidence"]) == {
        "source_pair_count": 2,
        "target_pair_count": 1,
        "automatic_pair_comparison_skipped": True,
    }
    queue_row = next(
        row for row in report["detail_rule_coverage"]["review_queue_rows"]
        if row["review_kind"] == "PAIR_ALIGNMENT"
    )
    assert queue_row["sku"] == "1001"
    assert queue_row["source_details"] == "Material: plástico; Color: azul"
    assert queue_row["target_details"] == "材质：塑料"
    policy = AUDIT._load_qa_policy(ROOT / "config/stage5/chinese_gold_qa_policy.json")
    assert AUDIT._issue_layer("DETAIL_PAIR_ALIGNMENT_REVIEW", policy) == "L4"
