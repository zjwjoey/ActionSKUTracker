"""CI_SAFE: annotated catalog publication against synthetic CURRENT records.

All workbooks and mutations live under tmp_path; no live SQLite, network or
production registry is involved. The annotation implementation is exercised
through the public export service, not replaced with a passing stub.
"""
from copy import deepcopy
from datetime import datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from action_tracker.exporting import service
from action_tracker.exporting.repair_report import ExportRepairReport
from action_tracker.localization.hashes import value_hash
from action_tracker.services.hashing import localization_field_source_hash, localization_source_hash
from test_exporting import _cfg, _record
from test_research_release_gate import _row


HEADERS = ["图片", "编号", "标题", "分类1", "分类2", "规格", "折后价", "原价", "单价",
           "描述", "产品详情", "图片链接", "商品链接", "备注"]
MODE = "ANNOTATED_RELEASE"
DATE = "2026-08-24"


@pytest.fixture
def current_source(tmp_path, monkeypatch):
    cfg = _cfg(tmp_path)
    records = []
    for sku in ('1001', '1002'):
        record = _record(sku)
        detail_es = f'Sustancia: Válido; Número del artículo: {sku}' if sku == '1001' else f'Número del artículo: {sku}'
        detail_zh = f'来源异常：官网字段Sustancia=Válido；商品编号：{sku}' if sku == '1001' else f'商品编号：{sku}'
        record.update(_row(sku=sku, details_es=detail_es, details_zh=detail_zh))
        for field, target_key in {'name':'name_zh','cat1':'cat1_zh','cat2':'cat2_zh',
                                  'spec':'spec_zh','description':'desc_zh','details':'details_zh'}.items():
            record['zh_field_provenance'][field]['value'] = record[target_key]
        if sku == '1001':
            record['zh_field_provenance']['details']['freshness_status'] = 'STALE'
        records.append(record)
    source = service.ExportSource(DATE, 'fixture_run', 'SQLITE_CURRENT', tuple(records), None,
                                  source_commit_id='fixture_committed_head')
    rows = []
    for record in records:
        rows.append({
            '图片':None, '编号':record['sku'], '标题':record['name_zh'],
            '分类1':record['cat1_zh'], '分类2':record['cat2_zh'], '规格':record['spec_zh'],
            '折后价':2.5, '原价':3.0, '单价':'1,25 €/件', '描述':record['desc_zh'],
            '产品详情':record['details_zh'], '图片链接':record['image_url'],
            '商品链接':record['product_url'], '备注':'在售状态：在售',
        })
    finding = {'sku':'1001', 'field':'details_zh', 'code':'SOURCE_DETAILS_TYPED_VALUE_INVALID',
               'blocking':True, 'severity':'BLOCKER', 'evidence':{'source_key':'Sustancia','source_value':'Válido'}}
    audit = {'release_ready':False, 'findings':[finding], 'blocking_affected_fields':1,
             'affected_fields':1, 'affected_skus':1, 'p0_findings':1,
             'unresolved_blocking_count':0, 'audit_findings':1}
    report = ExportRepairReport(run_id=source.run_id, rule_version='fixture')
    report.set_audit(audit)
    projection = service.FinalZhProjection(tuple(rows), {}, None, (), frozenset(), report, audit,
                                          service.canonical_export_rows_hash(rows))
    monkeypatch.setattr(service, 'resolve_formal_source', lambda *args, **kwargs: source)

    def project(*args, release_mode=service.PREVIEW, **kwargs):
        if release_mode == service.PRODUCTION_RELEASE:
            raise service.ExportValidationError('EXPORT_REPAIR_GATE_FAILED:fixture source anomaly')
        return projection

    monkeypatch.setattr(service, 'build_final_zh_projection', project)
    from action_tracker.data_quality import master_gate
    monkeypatch.setattr(master_gate, 'audit_master_quality',
                        lambda *args, **kwargs: SimpleNamespace(as_dict=lambda: {'release_ready':True, 'status':'PASS'}))
    plan = {
        'contract':'ANNOTATED_DELIVERY_V1', 'base_commit_id':source.source_commit_id,
        'source_hash':service.canonical_source_hash(source.records), 'export_date':DATE,
        'authorization_ref':'user:synthetic-annotated-delivery',
        'authorization_text':'Authorize this exact test plan with visible source uncertainty; do not approve facts.',
        'approval_mode':'OWNER_AUTHORIZED_PROXY',
        'approved_by':'service:current-localization-review',
        'approved_at':(datetime.now(timezone.utc)-timedelta(minutes=1)).isoformat(),
        'expires_at':(datetime.now(timezone.utc)+timedelta(days=1)).isoformat(),
        'fields':[{
            'sku':'1001', 'field':'details',
            'source_hash':localization_source_hash(records[0]),
            'field_source_hash':localization_field_source_hash(records[0], 'details'),
            'current_target_hash':value_hash(records[0]['details_zh']),
            'source_es':records[0]['details_es'], 'current_zh':records[0]['details_zh'],
            'action':'QUARANTINE', 'reason':'来源物态异常，保留待核状态。',
            'allowed_repair_codes':['SOURCE_DETAILS_TYPED_VALUE_INVALID'],
        }],
    }
    return cfg, source, rows, plan


def _export(fixture, **overrides):
    cfg, _, _, plan = fixture
    options = dict(language='zh', export_date=DATE, no_images=True, release_mode=MODE,
                   annotation_plan=deepcopy(plan))
    options.update(overrides)
    return service.export_catalog(cfg, **options)


def _read(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = workbook['商品全量']
        assert [cell.value for cell in sheet[1]] == HEADERS
        return [dict(zip(HEADERS, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    finally:
        workbook.close()


def test_annotated_service_publishes_explicit_notes_and_keeps_strict_result(current_source):
    result = _export(current_source)
    manifest = json.loads(Path(result['manifest']).read_text(encoding='utf-8'))
    rows = _read(result['output'])
    assert Path(result['output']).parent.name == 'annotated'
    assert '来源待核备注版' in Path(result['output']).name
    assert manifest['profile_id'].endswith('_annotated')
    assert manifest['release_mode'] == 'annotated_release'
    audit = manifest['annotated_delivery_audit']
    assert audit['research_verified'] is False
    assert audit['annotated_delivery_ready'] is True
    assert audit['strict_source_audit']['status'] == 'FAIL'
    assert audit['full_rows_hash'] == audit['published_full_rows_hash']
    assert rows[0]['产品详情'] == '待复核：来源字段存在异常，详见备注'
    assert 'Sustancia' in rows[0]['备注'] and 'Válido' in rows[0]['备注']
    assert rows[1] == current_source[2][1]
    assert [(r['编号'],r['折后价'],r['原价'],r['图片链接'],r['商品链接']) for r in rows] == [
        (r['编号'],r['折后价'],r['原价'],r['图片链接'],r['商品链接']) for r in current_source[2]
    ]


def test_normal_preview_unchanged_and_strict_release_still_blocks(current_source):
    cfg, _, original, _ = current_source
    result = service.export_catalog(cfg, language='zh', export_date=DATE, no_images=True)
    assert Path(result['output']).parent.name == 'preview'
    assert _read(result['output']) == original
    manifest = json.loads(Path(result['manifest']).read_text(encoding='utf-8'))
    assert 'annotated_delivery_audit' not in manifest
    with pytest.raises(service.ExportValidationError, match='EXPORT_REPAIR_GATE_FAILED'):
        service.export_catalog(cfg, language='zh', export_date=DATE, no_images=True,
                               release_mode=service.PRODUCTION_RELEASE)


@pytest.mark.parametrize('overrides', [
    {'annotation_plan':None}, {'language':'es'}, {'selection_id':'not-a-real-selection'},
])
def test_annotated_mode_rejects_missing_plan_es_or_selection(current_source, overrides):
    with pytest.raises((service.ExportValidationError, ValueError)):
        _export(current_source, **overrides)
    assert not list(Path(current_source[0]['paths']['exports']).rglob('*.xlsx'))


@pytest.mark.parametrize('key,value', [('approved_by',''), ('source_hash','wrong-source-hash')])
def test_plan_must_be_approved_and_bound_to_exact_source(current_source, key, value):
    plan = deepcopy(current_source[3]);plan[key] = value
    with pytest.raises((service.ExportValidationError, ValueError)):
        _export(current_source, annotation_plan=plan)


def test_known_exact_same_day_bundle_is_reused(current_source):
    first = _export(current_source)
    paths = [Path(first['output']), Path(first['manifest']), Path(first['repair_report'])]
    before = {str(p):(p.read_bytes(), p.stat().st_mtime_ns) for p in paths}
    second = _export(current_source)
    assert second['idempotent_reuse'] is True
    assert {str(p):(p.read_bytes(), p.stat().st_mtime_ns) for p in paths} == before


def test_unknown_same_day_bundle_is_never_overwritten(current_source):
    first = _export(current_source)
    manifest = Path(first['manifest']);payload = json.loads(manifest.read_text(encoding='utf-8'))
    payload['release_mode'] = 'untrusted-other-artifact'
    manifest.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
    paths = [Path(first['output']), manifest, Path(first['repair_report'])]
    before = {str(p):p.read_bytes() for p in paths}
    with pytest.raises(service.ExportValidationError):
        _export(current_source)
    assert {str(p):p.read_bytes() for p in paths} == before


@pytest.mark.parametrize('missing_key', ['manifest','repair_report'])
def test_incomplete_existing_bundle_is_never_replaced(current_source, missing_key):
    first = _export(current_source)
    Path(first[missing_key]).unlink()
    paths = [Path(first['output']),Path(first['manifest']),Path(first['repair_report'])]
    before = {str(p):(p.read_bytes() if p.exists() else None) for p in paths}
    with pytest.raises(service.ExportValidationError):
        _export(current_source)
    assert {str(p):(p.read_bytes() if p.exists() else None) for p in paths} == before


@pytest.mark.parametrize('mutation', ['plan','strict_audit'])
def test_existing_manifest_cannot_forge_plan_or_original_strict_audit(current_source, mutation):
    first = _export(current_source)
    path = Path(first['manifest']);manifest = json.loads(path.read_text(encoding='utf-8'))
    audit = manifest['annotated_delivery_audit']
    if mutation == 'plan':
        audit['annotation_plan']['approved_by'] = 'human:forged-approver'
    else:
        audit['strict_source_audit'] = {'status':'PASS','ok':True,'counts':{},'issues':[]}
    path.write_text(json.dumps(manifest,ensure_ascii=False),encoding='utf-8')
    before = {str(p):p.read_bytes() for p in [Path(first['output']),path,Path(first['repair_report'])]}
    with pytest.raises(service.ExportValidationError):
        _export(current_source)
    assert {p:Path(p).read_bytes() for p in before} == before


def test_existing_workbook_notes_and_hashes_cannot_be_forged_together(current_source):
    from action_tracker.exporting.annotated import full_delivery_rows_hash
    first = _export(current_source);output=Path(first['output']);manifest_path=Path(first['manifest'])
    workbook=openpyxl.load_workbook(output);workbook['商品全量']['N2']='伪造说明';workbook.save(output);workbook.close()
    manifest=json.loads(manifest_path.read_text(encoding='utf-8'));forged=full_delivery_rows_hash(_read(output),HEADERS)
    manifest['output_sha256']=hashlib.sha256(output.read_bytes()).hexdigest()
    for key in ['full_rows_hash','published_full_rows_hash','final_rows_hash']:
        manifest['annotated_delivery_audit'][key]=forged
    manifest_path.write_text(json.dumps(manifest,ensure_ascii=False),encoding='utf-8')
    before={str(p):p.read_bytes() for p in [output,manifest_path,Path(first['repair_report'])]}
    with pytest.raises(service.ExportValidationError):
        _export(current_source)
    assert {p:Path(p).read_bytes() for p in before}==before


def test_annotation_notes_are_wrapped_and_contribute_to_row_height(current_source):
    result = _export(current_source)
    workbook = openpyxl.load_workbook(result['output'])
    try:
        sheet = workbook['商品全量']
        assert sheet['N2'].alignment.wrap_text is True
        assert '来源事实未确认' in sheet['N2'].value
        assert sheet.row_dimensions[2].height > sheet.row_dimensions[3].height
        assert 50 <= sheet.row_dimensions[2].height <= 405
        assert not sheet._images
    finally:
        workbook.close()


def test_notes_tampering_after_writer_is_detected_before_publication(current_source, monkeypatch):
    actual_writer = service.write_catalog_xlsx

    def tamper(path, **kwargs):
        stats = actual_writer(path, **kwargs)
        workbook = openpyxl.load_workbook(path)
        workbook['商品全量']['N2'] = '篡改后的备注'
        workbook.save(path);workbook.close()
        return stats

    monkeypatch.setattr(service, 'write_catalog_xlsx', tamper)
    with pytest.raises(service.ExportValidationError):
        _export(current_source)
    assert not list(Path(current_source[0]['paths']['exports']).rglob('*.xlsx'))


def test_annotated_writer_runs_under_shared_daily_lock(current_source, monkeypatch):
    lock=Path(current_source[0]['paths']['state'])/'daily-run.lock'
    actual_writer=service.write_catalog_xlsx
    observed=[]

    def inspect_lock(path, **kwargs):
        assert lock.is_file()
        payload=json.loads(lock.read_text(encoding='utf-8'))
        assert payload['run_id']=='annotated-catalog-'+DATE
        assert payload['storage_token']
        observed.append(True)
        return actual_writer(path, **kwargs)

    monkeypatch.setattr(service,'write_catalog_xlsx',inspect_lock)
    _export(current_source)
    assert observed==[True]
    assert not lock.exists()


def test_unknown_existing_daily_lock_is_not_reclaimed_or_deleted(current_source, monkeypatch):
    state=Path(current_source[0]['paths']['state']);state.mkdir(parents=True)
    lock=state/'daily-run.lock';payload=b'{"pid":999999,"run_id":"unknown-owner","evidence":"retain"}'
    lock.write_bytes(payload)

    def forbidden_writer(*args, **kwargs):
        raise AssertionError('writer must not run while another owner holds the lock')

    monkeypatch.setattr(service,'write_catalog_xlsx',forbidden_writer)
    with pytest.raises(service.ExportValidationError,match='RUN_ALREADY_ACTIVE'):
        _export(current_source)
    assert lock.read_bytes()==payload
    assert not list(Path(current_source[0]['paths']['exports']).rglob('*.xlsx'))


def test_failed_writer_releases_only_its_own_lock(current_source, monkeypatch):
    lock=Path(current_source[0]['paths']['state'])/'daily-run.lock'

    def fail_writer(*args, **kwargs):
        assert lock.is_file()
        raise OSError('injected writer failure')

    monkeypatch.setattr(service,'write_catalog_xlsx',fail_writer)
    with pytest.raises(OSError,match='injected writer failure'):
        _export(current_source)
    assert not lock.exists()
    assert not list(Path(current_source[0]['paths']['exports']).rglob('*.xlsx'))


@pytest.mark.parametrize('mutation', ['duplicate','reorder'])
def test_workbook_verifier_rejects_extra_duplicate_or_changed_order(tmp_path, mutation):
    rows = []
    for sku in ['1001','1002']:
        row = dict.fromkeys(HEADERS);row.update(编号=sku, 标题='商品', 折后价=2.5)
        rows.append(row)
    path = tmp_path/'verification.xlsx'
    service.write_catalog_xlsx(path, headers=HEADERS, rows=rows, workbook_format={})
    workbook = openpyxl.load_workbook(path);sheet = workbook['商品全量']
    if mutation == 'duplicate':
        sheet.append([cell.value for cell in sheet[2]])
    else:
        sheet['B2']='1002';sheet['B3']='1001'
    workbook.save(path);workbook.close()
    with pytest.raises(service.ExportValidationError):
        service._verify_written_workbook(path, headers=HEADERS,
            expected_skus={'1001','1002'}, expected_sku_order=['1001','1002'])
