"""CI_SAFE: trusted generated CSV/XLSX shapes, existing single writer only."""
import copy
import csv
import hashlib
import sqlite3
from dataclasses import replace
from datetime import datetime, timezone

import openpyxl
import pytest

from action_tracker.database.connection import connect
from action_tracker.database.production import CommitBundle, ProductionWriter, ProductionDatabaseError
from action_tracker.database.repository import ProductionRepository
from action_tracker.data_quality.historical.variance import typed_hash
from action_tracker.localization.history_recovery import build_historical_import_bundle
from action_tracker.operations.registry import committed_registry_records
from action_tracker.exporting.service import _resolve_sqlite_current_source, ExportValidationError
from action_tracker.workflow_v2.fact_adapter import load_committed_daily_source


def _fixture(tmp_path, *, xlsx=False):
    path = tmp_path / 'isolated.sqlite3'
    writer = ProductionWriter(path, role='PRIMARY')
    statuses = ('CURRENT', 'MISSING', 'OFFLINE', 'HISTORICAL', 'UNKNOWN', 'HISTORICAL_ONLY')
    seed = [{'sku': str(100 + i), 'canonical_id': 'ACT' + str(100+i), 'status': state,
        'name_es': 'Original', 'current_price': 3.5, 'consecutive_missing': i,
        'first_seen_at': '2020-01-01', 'last_seen_at': '2020-02-01', 'last_checked_at': '2020-02-02'} for i, state in enumerate(statuses)]
    head = writer.commit(CommitBundle(run_id='seed', observation_date='2020-02-02', qa_state='PASS', current_products=tuple(seed),
        localization_updates=tuple({'sku': r['sku'], 'language': lang, 'name': 'Original' if lang == 'es' else '历史商品',
            'review_status': 'VERIFIED' if lang == 'es' else 'APPROVED', 'approved_by': 'human:fixture' if lang == 'zh' else None,
            'approved_at': '2020-01-01T00:00:00Z' if lang == 'zh' else None} for r in seed for lang in ('es', 'zh'))))
    records = [{'sku': str(100+i), 'name_es': 'Producto archivado', 'spec_es': '2 unidades',
        'current_price': '99.99'} for i in range(7)]
    asset = tmp_path / ('archive.xlsx' if xlsx else 'archive.csv')
    if xlsx:
        book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Archivo'
        sheet.append(list(records[0])); [sheet.append(list(r.values())) for r in records]; book.save(asset); book.close()
    else:
        with asset.open('w', encoding='utf-8', newline='') as stream:
            w = csv.DictWriter(stream, fieldnames=list(records[0])); w.writeheader(); w.writerows(records)
    source = {'path': str(asset), 'asset_sha256': hashlib.sha256(asset.read_bytes()).hexdigest(),
        'kind': 'VERIFIED_XLSX' if xlsx else 'SNAPSHOT_CSV', 'sheet': 'Archivo' if xlsx else None,
        'sku_header': 'sku', 'fields': {'name': 'name_es', 'spec': 'spec_es'}, 'language': 'es',
        'evidence_level': 'A', 'source_version': 'generated-archive-v1', 'observed_at': '2019-06-01',
        'observed_at_evidence': 'Synthetic original capture metadata, not inferred from path',
        'reviewed_by': 'human:fixture', 'reviewed_at': datetime.now(timezone.utc).isoformat(),
        'review_evidence': 'Generated fixture; source SKU/columns/hash checked'}
    candidates = [{'sku': r['sku'], 'canonical_id': 'ACT'+(r['sku'].zfill(7) if r['sku']=='106' else r['sku']), 'record_hash': typed_hash(r), 'source': source} for r in records]
    cfg = {'project_root': tmp_path, 'storage': {'mode': 'SQLITE_PRIMARY', 'db_path': path}, 'paths': {'state': tmp_path / 'state'}}
    return path, writer, head, asset, candidates, cfg


def _build(fixture, candidates=None, **kwargs):
    path, _, head, _, original, _ = fixture
    return build_historical_import_bundle(path, original if candidates is None else candidates,
        run_id=kwargs.pop('run_id', 'history'), base_commit_id=kwargs.pop('base_commit_id', head), run_date='2026-10-30', **kwargs)


@pytest.mark.parametrize('xlsx', [False, True])
def test_history_import_preserves_status_dates_prices_approved_zh_and_registry_scope(tmp_path, xlsx):
    fixture = _fixture(tmp_path, xlsx=xlsx)
    path, writer, _, _, candidates, cfg = fixture
    with connect(path) as db:
        before = [tuple(row) for row in db.execute('SELECT * FROM products ORDER BY official_sku')]
        approvals = [tuple(row) for row in db.execute("SELECT official_sku,name,review_status,approved_by,approved_at FROM product_localizations WHERE language='zh' ORDER BY official_sku")]
        schema = [tuple(row) for row in db.execute("SELECT name,sql FROM sqlite_master WHERE type='table' ORDER BY name")]
    bundle, report = _build(fixture)
    assert report['new_identities'] == 1 and report['filled_fields'] == 8
    assert len([r for r in report['blocked'] if r['reason'] == 'EXISTING_FACT_PRESERVED_SOURCE_VARIANCE']) == 6
    assert not bundle.observations and not bundle.lifecycle_updates and not bundle.price_events and not bundle.event_events
    head = writer.commit(bundle)
    assert writer.commit(bundle) == head
    with connect(path) as db:
        assert [tuple(row) for row in db.execute("SELECT * FROM products WHERE official_sku<>'106' ORDER BY official_sku")] == before
        assert [tuple(row) for row in db.execute("SELECT official_sku,name,review_status,approved_by,approved_at FROM product_localizations WHERE language='zh' ORDER BY official_sku")] == approvals
        row = dict(db.execute("SELECT * FROM products WHERE official_sku='106'").fetchone())
        assert row['status'] == 'HISTORICAL_ONLY' and row['current_price'] is None and row['product_url'] is None
        assert row['first_seen_at'] is None and row['last_seen_at'] is None and row['last_checked_at'] is None
        assert db.execute('SELECT COUNT(*) FROM observations').fetchone()[0] == 0
        assert db.execute('SELECT COUNT(*) FROM price_history').fetchone()[0] == 0
        assert db.execute('SELECT COUNT(*) FROM event_history').fetchone()[0] == 0
        assert [tuple(row) for row in db.execute("SELECT name,sql FROM sqlite_master WHERE type='table' ORDER BY name")] == schema
        assert db.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
        assert not db.execute('PRAGMA foreign_key_check').fetchall()
    rows = committed_registry_records(cfg, source_run_id='history', fact_commit_id=head, business_date='2026-10-30')
    assert len(rows) == 7 and all(row['_registry_observed_at'] == 'UNKNOWN' for row in rows)
    assert all(row['_registry_allowed_fields'] == ['spec_es'] for row in rows if row['sku'] != '106')
    repo = ProductionRepository(path)
    assert [r['sku'] for r in repo.load_current_export_records()] == ['100']
    all_rows = {r['sku']: r for r in repo.load_current_export_records(include_non_current=True, include_field_lineage=True)}
    assert all_rows['106']['field_source_lineage']['spec_es']['field_source_observed_at'] == '2019-06-01'
    assert all_rows['106']['field_source_lineage']['spec_es']['field_source_business_date'] == '2019-06-01'
    assert repo.latest_commit_info()['commit_scope'] == 'HISTORICAL_FACTS_ONLY'
    assert _resolve_sqlite_current_source(cfg, export_date='2026-10-30', requested_run_id=None) is None
    with pytest.raises(ValueError, match='COMPLETE_PRESENCE'):
        load_committed_daily_source(cfg, business_date='2026-10-30')
    # Reusing the exact source in another task must not duplicate fact versions.
    again, again_report = _build(fixture, run_id='history-repeat', base_commit_id=head)
    assert again_report['source_versions'] == 0 and again_report['filled_fields'] == 0 and again_report['new_identities'] == 0
    from action_tracker.localization.registry.repository import LocalizationRegistry
    registry = LocalizationRegistry(path, role='PRIMARY')
    try:
        first = registry.ingest_records(rows, source_run_id='history', observed_at='2026-10-30')
        with connect(path) as db:
            queue_before = db.execute('SELECT COUNT(*) FROM translation_queue').fetchone()[0]
        second = registry.ingest_records(rows, source_run_id='history', observed_at='2026-10-30')
        with connect(path) as db:
            assert db.execute('SELECT COUNT(*) FROM translation_queue').fetchone()[0] == queue_before
    finally:
        registry.close()
    with connect(path) as db:
        assert {r[0] for r in db.execute("SELECT observed_at FROM translation_source_versions WHERE source_run_id='history'")} == {'UNKNOWN'}


def test_history_import_fails_closed_on_identity_asset_hash_record_and_source_review(tmp_path):
    fixture = _fixture(tmp_path)
    for change in ('record', 'hash', 'sku', 'canonical', 'new_canonical', 'reviewer', 'language', 'level', 'price_field', 'url'):
        candidates = copy.deepcopy(fixture[4])
        row = candidates[0]
        if change == 'record': row['record_hash'] = '0'*64
        elif change == 'hash': row['source']['asset_sha256'] = '0'*64
        elif change == 'sku': row['sku'] = '999'
        elif change == 'canonical': row['canonical_id'] = 'ACT101'
        elif change == 'new_canonical': candidates[-1]['canonical_id'] = 'ACT106'
        elif change == 'reviewer': row['source']['reviewed_by'] = 'service:invented'
        elif change == 'language': row['source']['language'] = 'zh'
        elif change == 'level': row['source']['evidence_level'] = 'UNKNOWN'
        elif change == 'price_field': row['source']['fields']['current_price'] = 'current_price'
        elif change == 'url': row['source']['url_header'] = 'name_es'
        with pytest.raises(ValueError): _build(fixture, candidates)


def test_history_transaction_binds_preview_and_rejects_mutation_and_changed_source(tmp_path):
    fixture = _fixture(tmp_path)
    path, writer, head, asset, _, _ = fixture
    bundle, _ = _build(fixture)
    for key in ('current_products', 'localization_updates', 'price_events', 'event_events', 'observations', 'lifecycle_updates'):
        with pytest.raises(ValueError, match='SCOPE_OR_FACT_CHANGED'):
            writer.commit(replace(bundle, **{key: (*getattr(bundle, key), {'sku': '999'})}))
    original = asset.read_bytes(); asset.write_bytes(original + b'changed')
    with pytest.raises(ValueError, match='ASSET_HASH_CHANGED'): writer.commit(bundle)
    asset.write_bytes(original)
    writer.commit(CommitBundle(run_id='other', observation_date='2026-10-30', qa_state='PASS', base_commit_id=head))
    with pytest.raises(ProductionDatabaseError, match='BASELINE_CHANGED'): writer.commit(bundle)


def test_history_conflict_keeps_both_and_unknown_date_stays_unknown(tmp_path):
    fixture = _fixture(tmp_path)
    path, writer, head, _, candidates, _ = fixture
    one = copy.deepcopy(candidates[-1]); one['source']['observed_at'] = None
    two = copy.deepcopy(one)
    asset = tmp_path / 'other.csv'
    record = {'sku': '106', 'name_es': 'Producto archivado', 'spec_es': '3 unidades', 'current_price': '99.99'}
    with asset.open('w', encoding='utf-8', newline='') as stream:
        w = csv.DictWriter(stream, fieldnames=list(record)); w.writeheader(); w.writerow(record)
    two['source']['path'] = str(asset); two['source']['asset_sha256'] = hashlib.sha256(asset.read_bytes()).hexdigest()
    two['source']['source_version'] = 'another-verified-version'; two['record_hash'] = typed_hash(record)
    bundle, report = _build(fixture, [one, two])
    assert report['filled_fields'] == 1 and any(r['reason'] == 'SOURCE_VERSION_CONFLICT_KEEP_BOTH' for r in report['blocked'])
    writer.commit(bundle)
    with connect(path) as db:
        assert db.execute("SELECT spec FROM product_localizations WHERE official_sku='106' AND language='es'").fetchone()[0] is None
        assert {row[0] for row in db.execute("SELECT normalized_value FROM source_fact_versions WHERE official_sku='106' AND field_name='spec'")} == {'2 unidades', '3 unidades'}
    row = next(r for r in ProductionRepository(path).load_current_export_records(include_non_current=True, include_field_lineage=True) if r['sku']=='106')
    assert row['field_source_lineage']['name_es']['field_source_observed_at'] is None
    assert row['field_source_lineage']['name_es']['field_source_time_precision'] == 'UNKNOWN'


def test_product_title_fallback_protects_title_without_changing_empty_es_name(tmp_path):
    fixture = _fixture(tmp_path)
    path, writer, _, _, candidates, _ = fixture
    with connect(path) as db:
        db.execute("UPDATE product_localizations SET name=NULL WHERE official_sku='100' AND language='es'")
    bundle, _ = _build(fixture, candidates[:1])
    writer.commit(bundle)
    with connect(path) as db:
        assert db.execute("SELECT name,spec FROM product_localizations WHERE official_sku='100' AND language='es'").fetchone()[:] == (None, '2 unidades')
        assert db.execute("SELECT name_es FROM products WHERE official_sku='100'").fetchone()[0] == 'Original'


def test_history_overlay_uses_prior_complete_presence_for_formal_current_source(tmp_path):
    fixture = _fixture(tmp_path)
    path, writer, head, _, candidates, cfg = fixture
    complete = writer.commit(CommitBundle(run_id='complete-presence', observation_date='2026-10-30', qa_state='PASS', base_commit_id=head,
        observations=({'sku': '100', 'run_id': 'complete-presence', 'observation_date': '2026-10-30',
            'presence_state': 'PRESENT', 'observation_complete': True, 'absence_capable': True, 'current_price': 3.5},)))
    before = _resolve_sqlite_current_source(cfg, export_date='2026-10-30', requested_run_id=None)
    bundle, _ = _build(fixture, base_commit_id=complete)
    imported = writer.commit(bundle)
    after = _resolve_sqlite_current_source(cfg, export_date='2026-10-30', requested_run_id=None)
    assert after.presence_source_commit_id == before.presence_source_commit_id == complete
    assert after.source_commit_id == imported
    assert [(r['sku'], r['current_price'], r['product_url'], r['image_url']) for r in before.records] == [(r['sku'], r['current_price'], r['product_url'], r['image_url']) for r in after.records]
    assert [r['sku'] for r in after.records] == ['100']


def test_fact_review_service_authority_is_external_and_exact(tmp_path):
    fixture = _fixture(tmp_path)
    path, _, head, _, candidates, _ = fixture
    candidates = copy.deepcopy(candidates[:1]); candidates[0]['source']['reviewed_by'] = 'service:authorized-fixture'
    with pytest.raises(ValueError): _build(fixture, candidates)
    bundle, _ = _build(fixture, candidates, authorized_service_actors=('service:authorized-fixture',))
    with pytest.raises(ValueError): fixture[1].commit(bundle)
    writer = ProductionWriter(path, role='PRIMARY', authorized_fact_review_service_actors=('service:authorized-fixture',))
    writer.commit(bundle)
    with pytest.raises(ProductionDatabaseError, match='ACTOR_CONFIG_INVALID'):
        ProductionWriter(path, role='PRIMARY', authorized_fact_review_service_actors=('service:*',))


def test_reviewed_xlsx_duplicate_headers_cannot_silently_select_last_value(tmp_path):
    fixture = _fixture(tmp_path, xlsx=True)
    path, _, _, asset, candidates, _ = fixture
    workbook = openpyxl.load_workbook(asset)
    workbook.active.cell(1, 4, 'spec_es')
    workbook.save(asset); workbook.close()
    candidates = copy.deepcopy(candidates)
    for candidate in candidates:
        candidate['source']['asset_sha256'] = hashlib.sha256(asset.read_bytes()).hexdigest()
    with pytest.raises(ValueError, match='DUPLICATE_HEADERS'):
        _build(fixture, candidates)
