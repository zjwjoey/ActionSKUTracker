from openpyxl import Workbook
from openpyxl.styles import Font
import sqlite3
import pytest
from action_tracker.database.repository import _price_projection
from action_tracker.excel.writer import _update_or_append_current


@pytest.mark.parametrize('evidence,expected', [
    ('{}', (3, 5, '3 €/ud')),
    ('{', (3, 5, '3 €/ud')),
    ('null', (3, 5, '3 €/ud')),
    ('{"price_observation_invalid":[]}', (3, 5, '3 €/ud')),
    ('{"price_field_observation_invalid":"unknown"}', (3, 5, '3 €/ud')),
    ('{"price_field_observation_invalid":{"1001":"unknown"}}', (3, 5, '3 €/ud')),
    ('{"price_observation_invalid":{"1001":"1"}}', (3, 5, '3 €/ud')),
    ('{"price_observation_invalid":{"1001":true}}', (None, None, None)),
    ('{"price_field_observation_invalid":{"1001":{"original_price":true}}}', (3, None, '3 €/ud')),
    ('{"price_field_observation_invalid":{"1001":{"unit_price":true}}}', (3, 5, None)),
    ('{"price_observation_invalid":{"1001":false,"1001":true}}', (3, 5, '3 €/ud')),
    ('{"price_observation_invalid":{"1001":true,"1001":false}}', (None, None, None)),
    ('{"price_field_observation_invalid":{"1001":"unknown","1001":{"unit_price":true}}}', (3, 5, '3 €/ud')),
    ('{"price_field_observation_invalid":{"1001":{"original_price":false},"1001":{"original_price":true}}}', (3, 5, '3 €/ud')),
])
def test_price_marker_projection_preserves_types_independence_and_first_key(evidence, expected):
    db = _price_marker_fixture(evidence)
    try:
        assert _read_marker_prices(db) == [expected]
    finally:
        db.close()


def test_price_marker_from_old_run_cannot_leak_into_new_observation():
    db = _price_marker_fixture('{"price_observation_invalid":{"1001":true}}')
    try:
        db.execute("INSERT INTO commit_batches VALUES('repair','repair','2026-10-28','COMMITTED')")
        db.execute("INSERT INTO run_evidence VALUES('repair','{}')")
        assert _read_marker_prices(db) == [(None, None, None)]
        db.execute("INSERT INTO commit_batches VALUES('new','new','2026-10-27','COMMITTED')")
        db.execute("INSERT INTO run_evidence VALUES('new','{}')")
        db.execute("INSERT INTO observations VALUES('1001','new','2026-10-27','PRESENT')")
        assert _read_marker_prices(db) == [(3, 5, '3 €/ud')]
    finally:
        db.close()


def _price_marker_fixture(evidence):
    db = sqlite3.connect(':memory:')
    db.executescript("""CREATE TABLE products(official_sku TEXT PRIMARY KEY,current_price REAL,original_price REAL,unit_price_raw TEXT);
        CREATE TABLE observations(official_sku TEXT,run_id TEXT,observation_date TEXT,presence_state TEXT);
        CREATE TABLE commit_batches(commit_id TEXT,run_id TEXT,committed_at TEXT,status TEXT);
        CREATE TABLE run_evidence(run_id TEXT PRIMARY KEY,evidence_json TEXT);
        INSERT INTO products VALUES('1001',3,5,'3 €/ud');
        INSERT INTO observations VALUES('1001','old','2026-10-26','PRESENT');
        INSERT INTO commit_batches VALUES('old','old','2026-10-26','COMMITTED');""")
    db.execute('INSERT INTO run_evidence VALUES(?,?)', ('old', evidence))
    return db


def _read_marker_prices(db):
    from action_tracker.database.repository import PRICE_REVIEW_JOIN, current_price_expression
    expressions = ','.join(current_price_expression(c) for c in ('current_price','original_price','unit_price_raw'))
    return db.execute('SELECT ' + expressions + ' FROM products p ' + PRICE_REVIEW_JOIN).fetchall()


def test_last_change_is_distinct_from_current_business_period():
    history = [('s', '2026-01-09', None, .29, 'INITIAL'),
               ('s', '2026-04-05', .29, .25, 'DOWN')]
    r = _price_projection(.25, '', '', False, history, '2026-10-11')
    assert r['last_price'] == .29
    assert r['price_direction'] == 'DOWN'
    assert r['price_change'] == 'UNCHANGED'
    assert r['change_amount'] == -.04
    assert r['change_percent'] == -.137931
    assert r['last_change_date'] == '2026-04-05'
    assert (r['price_min'], r['price_max']) == (.25, .29)
    assert r['discount'] is None


def test_extrema_include_initial_old_and_current_price():
    r = _price_projection(4, None, None, False,
                          [('s', '2026-10-11', 6, 5, 'DOWN')], '2026-10-11')
    assert (r['price_min'], r['price_max']) == (4, 6)
    assert r['price_change'] == 'DOWN'
    empty = _price_projection(2, '', '', False, [], '2026-10-11')
    assert (empty['price_min'], empty['price_max']) == (2, 2)
    assert empty['last_price'] is None
    initial = _price_projection(2, None, '', False,
                               [('s', '2026-10-11', None, 2, 'INITIAL')], '2026-10-11')
    assert initial['price_change'] == 'NEW'
    assert initial['price_direction'] is None


def test_discount_uses_current_badge_or_price_not_old_excel_discount():
    assert _price_projection(8, 10, '', True, [], '')['discount'] == -.2
    assert _price_projection(8, 10, '-21%', True, [], '')['discount'] == -.21
    assert _price_projection(8, '', '-21%', False, [], '')['discount'] is None
    assert _price_projection(8, 8, '', True, [], '')['discount'] is None
    assert _price_projection(8, None, '', True, [], '')['discount'] is None


def test_current_writer_uses_headers_preserves_rows_and_aligns_languages():
    wb = Workbook(); wb.remove(wb.active)
    mappings = [({'SKU':'sku','Name':'name','Canonical_ID':'canonical_id','Price':'price'},
                 ['SKU','Name','Canonical_ID','Price']),
                ({'Canonical_ID':'canonical_id','SKU':'sku','Name':'name','Price':'price'},
                 ['Canonical_ID','SKU','Name','Price'])]
    records = {'1':{'sku':'1','canonical_id':'ACT1','name':'one','price':5},
               '2':{'sku':'2','canonical_id':'ACT2','name':'two','price':6}}
    for i,(mapping,headers) in enumerate(mappings):
        ws=wb.create_sheet(str(i)); ws.append(headers)
        for sku in ['2','1','3']:
            old={'sku':sku,'canonical_id':'ACT'+sku,'name':'old','price':99}
            ws.append([old[mapping[h]] for h in headers])
        ws.cell(3,headers.index('SKU')+1).font=Font(bold=True)
        assert _update_or_append_current(wb,ws.title,mapping,records,set(mapping.values())) == 2
        assert ws.max_row == 3
        assert [ws.cell(j,headers.index('SKU')+1).value for j in [2,3]] == ['1','2']
        assert ws.cell(2,headers.index('SKU')+1).font.bold
        assert [ws.cell(j,headers.index('Price')+1).value for j in [2,3]] == [5,6]
        assert _update_or_append_current(wb,ws.title,mapping,records,set(mapping.values())) == 2
        assert ws.max_row == 3


def test_current_writer_formats_fractional_discounts_and_prices_in_both_orders():
    wb=Workbook();ws=wb.active
    mapping={'SKU':'sku','Canonical_ID':'canonical_id','Price':'current_price','Change':'change_percent','Discount':'discount'}
    ws.append(list(mapping));ws.append(['1','ACT1',99,.99,.99])
    records={'1':{'sku':'1','canonical_id':'ACT1','current_price':8,'change_percent':-.2,'discount':-.21}}
    _update_or_append_current(wb,ws.title,mapping,records,set(mapping.values()))
    assert ws.cell(2,3).value == 8
    assert ws.cell(2,3).number_format == '€#,##0.00'
    assert ws.cell(2,4).value == -.2 and ws.cell(2,4).number_format == '0.0%'
    assert ws.cell(2,5).value == -.21 and ws.cell(2,5).number_format == '0.0%'
