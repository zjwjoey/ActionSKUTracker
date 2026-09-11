"""Read-only post-repair audit for SQLite PRIMARY and its Master projection."""
from __future__ import annotations

import json
import re
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from action_tracker.config import load_settings
from action_tracker.database.integration import database_path
from action_tracker.services.hashing import localization_source_hash


FIELDS = ("name", "cat1", "cat2", "spec", "description", "details")
SKU_RE = re.compile(r"/p/(\d+)(?:/|$)")


def audit(db_path: Path, master_path: Path) -> dict[str, object]:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        current = [str(row[0]) for row in conn.execute("SELECT official_sku FROM products WHERE status='CURRENT'")]
        result: dict[str, object] = {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "database": str(db_path),
            "master": str(master_path),
            "current_sku_count": len(current),
            "current_duplicate_sku_count": len(current) - len(set(current)),
            "required_blank_fields": {},
            "localization_fields_value_mismatches": {},
            "canonical_provenance_value_mismatches": {},
            "detail_pollution": {},
            "price_logic_error_count": 0,
            "sku_url_mismatch_count": 0,
            "sku_detail_mismatch_count": 0,
            "source_hash_mismatch_count": 0,
            "category_backlog_count": 0,
            "translation_queue_count": 0,
            "translation_queue_pending_count": 0,
            "fact_version_count": 0,
            "patch_count": 0,
            "master_row_count": None,
            "master_sku_mismatch_count": None,
        }
        for language in ("es", "zh"):
            blanks: dict[str, int] = {}
            for field in FIELDS:
                n = conn.execute(
                    f"SELECT COUNT(*) FROM product_localizations l JOIN products p ON p.official_sku=l.official_sku "
                    f"WHERE p.status='CURRENT' AND l.language=? AND (l.{field} IS NULL OR trim(l.{field})='')",
                    (language,),
                ).fetchone()[0]
                if n:
                    blanks[field] = int(n)
            result["required_blank_fields"][language] = blanks
            for table_name, result_key in (("localization_fields", "localization_fields_value_mismatches"),
                                           ("localization_field_provenance", "canonical_provenance_value_mismatches")):
                for field in FIELDS:
                    n = conn.execute(
                        f"SELECT COUNT(*) FROM {table_name} f JOIN product_localizations l "
                        f"ON l.official_sku=f.official_sku AND l.language=f.language "
                        f"JOIN products p ON p.official_sku=l.official_sku "
                        f"WHERE p.status='CURRENT' AND f.language=? AND f.field_name=? "
                        f"AND coalesce(f.value,'')<>coalesce(l.{field},'')",
                        (language, field),
                    ).fetchone()[0]
                    if n:
                        result[result_key][f"{language}.{field}"] = int(n)

        pollution = {
            "double_colon": "::",
            "null": "null",
            "undefined": "undefined",
            "html_tag": "<",
            "leer_mas": "leer más",
        }
        for label, needle in pollution.items():
            n = conn.execute(
                "SELECT COUNT(*) FROM product_localizations l JOIN products p ON p.official_sku=l.official_sku "
                "WHERE p.status='CURRENT' AND l.language='es' AND lower(coalesce(l.description,'') || ' ' || coalesce(l.details,'')) LIKE ?",
                (f"%{needle.lower()}%",),
            ).fetchone()[0]
            result["detail_pollution"][label] = int(n)

        price_errors = conn.execute(
            "SELECT COUNT(*) FROM products WHERE status='CURRENT' AND original_price IS NOT NULL "
            "AND current_price IS NOT NULL AND original_price<=current_price"
        ).fetchone()[0]
        result["price_logic_error_count"] = int(price_errors)
        for row in conn.execute("SELECT official_sku,product_url FROM products WHERE status='CURRENT'"):
            match = SKU_RE.search(str(row[1] or ""))
            if not match or match.group(1) != str(row[0]):
                result["sku_url_mismatch_count"] += 1
        for row in conn.execute(
            "SELECT p.official_sku,p.source_hash,es.name,es.cat1,es.cat2,es.spec,es.description,es.details,zh.source_hash "
            "FROM products p JOIN product_localizations es ON es.official_sku=p.official_sku AND es.language='es' "
            "LEFT JOIN product_localizations zh ON zh.official_sku=p.official_sku AND zh.language='zh' "
            "WHERE p.status='CURRENT'"
        ):
            expected = localization_source_hash({
                "name_es": row[2], "cat1_es": row[3], "cat2_es": row[4], "spec_es": row[5],
                "desc_es": row[6], "details_es": row[7],
            })
            if str(row[1] or "") != str(expected or "") or str(row[8] or "") != str(expected or ""):
                result["source_hash_mismatch_count"] += 1
        for row in conn.execute(
            "SELECT p.official_sku,l.details FROM products p JOIN product_localizations l "
            "ON l.official_sku=p.official_sku AND l.language='es' WHERE p.status='CURRENT'"
        ):
            matches = re.findall(r"(?:N(?:ú|u)mero del art(?:í|i)culo|Número de artículo)\s*:\s*(\d+)", str(row[1] or ""), re.IGNORECASE)
            if matches and any(value != str(row[0]) for value in matches):
                result["sku_detail_mismatch_count"] += 1
        for table, key in (("category_backlog", "category_backlog_count"),
                           ("product_fact_versions", "fact_version_count"),
                           ("localization_patches", "patch_count")):
            result[key] = int(conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0])
        result["translation_queue_count"] = int(conn.execute("SELECT COUNT(*) FROM translation_queue").fetchone()[0])
        result["translation_queue_pending_count"] = int(conn.execute("SELECT COUNT(*) FROM translation_queue WHERE status='PENDING'").fetchone()[0])

        if master_path.exists():
            import openpyxl
            workbook = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
            try:
                sheet = workbook["02_SKU_ES_CURRENT"]
                headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
                idx = headers.index("SKU")
                rows = [str(row[idx].value).strip() for row in sheet.iter_rows(min_row=2) if row[idx].value not in (None, "")]
                result["master_row_count"] = len(rows)
                result["master_sku_mismatch_count"] = len(set(rows) ^ set(current))
            finally:
                workbook.close()
        result["qc_result"] = "PASS" if all([
            result["current_sku_count"] == 5536,
            result["current_duplicate_sku_count"] == 0,
            not any(result["required_blank_fields"].values()),
            not result["localization_fields_value_mismatches"],
            not result["canonical_provenance_value_mismatches"],
            result["detail_pollution"]["double_colon"] == 0,
            result["detail_pollution"]["null"] == 0,
            result["detail_pollution"]["undefined"] == 0,
            result["detail_pollution"]["html_tag"] == 0,
            result["price_logic_error_count"] == 0,
            result["sku_url_mismatch_count"] == 0,
            result["sku_detail_mismatch_count"] == 0,
            result["source_hash_mismatch_count"] == 0,
            result["master_sku_mismatch_count"] == 0,
        ]) else "FAIL"
        return result
    finally:
        conn.close()


def main() -> None:
    cfg = load_settings()
    result = audit(database_path(cfg), Path(cfg["paths"]["master"]))
    out = Path(cfg["paths"].get("logs", cfg["paths"]["temp"])) / f"master_integrity_audit_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({**result, "report": str(out)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
