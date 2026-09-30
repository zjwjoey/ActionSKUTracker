"""Apply the owner-authorized 2026-09-21 localization repair batch.

This is deliberately a one-purpose, transaction-backed apply tool.  It makes
the repaired Spanish snapshot authoritative, applies only fields present in
the reviewed Qwen-MT/Codex artifact, and rebuilds the legacy Excel Master from
SQLite PRIMARY.  It does not fetch, translate, or infer any new product fact.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
import sqlite3
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from action_tracker.config import load_settings
from action_tracker.database.connection import connect
from action_tracker.database.integration import database_path, regenerate_compatibility_exports
from action_tracker.database.production import backup_database
from action_tracker.database.provenance import sync_localization_field_provenance
from action_tracker.dictionary import (
    CATEGORY_DICTIONARY_HEADERS,
    PRODUCT_DICTIONARY_HEADERS,
    load_dictionary_rows,
    product_source_hash,
    write_dictionary_csv,
)
from action_tracker.services.hashing import content_hash, localization_source_hash


BATCH_ID = "20260921_qwen_mt_flash_untranslated_v1"
POLICY_VERSION = "ACTION_LOCALIZATION_REVIEW_CONTRACT_V2"
SOURCE_XLSX = ROOT / "runtime/exports/20260921Action商品全量_西班牙语版_不带图_程序修复后_表格修复版.xlsx"
REVIEW_DIR = ROOT / "runtime/translation" / BATCH_ID / "review"
REVIEW_CSV = REVIEW_DIR / "reviewed_fields.csv"
CATEGORY_CSV = REVIEW_DIR / "category_mapping_recommendations.csv"

FIELD_SOURCE = {
    "name": "标题", "cat1": "分类1", "cat2": "分类2", "spec": "规格",
    "description": "描述", "details": "产品详情",
}
LOCALIZATION_COLUMNS = {field: field for field in FIELD_SOURCE}
PRODUCT_DICTIONARY_FIELDS = {
    "name": "name_zh_standard", "cat1": "cat1_zh", "cat2": "cat2_zh", "spec": "spec_zh_standard",
}
ALL_FIELDS = tuple(FIELD_SOURCE)


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def digest(value: Any) -> str:
    return hashlib.sha256(str(value or "").encode("utf-8")).hexdigest()


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def snapshot_hash(source: dict[str, str]) -> str:
    # This matches the candidate generator's frozen batch hash exactly.
    candidate_source = {column: source[field] for field, column in FIELD_SOURCE.items()}
    return digest(json.dumps(candidate_source, ensure_ascii=False, sort_keys=True))


def read_source() -> dict[str, dict[str, str]]:
    wb = openpyxl.load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    try:
        ws = wb["商品全量"]
        headers = [text(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        index = {value: pos for pos, value in enumerate(headers)}
        required = ["编号", *FIELD_SOURCE.values()]
        missing = [value for value in required if value not in index]
        if missing:
            raise RuntimeError(f"SOURCE_COLUMNS_MISSING:{','.join(missing)}")
        result: dict[str, dict[str, str]] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            sku = text(row[index["编号"]])
            if not sku:
                continue
            if sku in result:
                raise RuntimeError(f"SOURCE_DUPLICATE_SKU:{sku}")
            result[sku] = {field: text(row[index[column]]) for field, column in FIELD_SOURCE.items()}
        return result
    finally:
        wb.close()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return [{key: text(value) for key, value in row.items()} for row in csv.DictReader(handle)]


def validated_reviews(source: dict[str, dict[str, str]]) -> tuple[list[dict[str, str]], dict[tuple[str, str], str]]:
    rows = read_csv(REVIEW_CSV)
    recommendations = read_csv(CATEGORY_CSV)
    approved_categories = {
        (row["cat1_es"], row["cat2_es"]): row["recommended_cat2_zh"] for row in recommendations
    }
    seen: set[tuple[str, str]] = set()
    for row in rows:
        key = (row["sku"], row["field_name"])
        if key in seen:
            raise RuntimeError(f"REVIEW_DUPLICATE_FIELD:{key}")
        seen.add(key)
        if row["field_name"] not in FIELD_SOURCE or row["sku"] not in source:
            raise RuntimeError(f"REVIEW_IDENTITY_INVALID:{key}")
        official = source[row["sku"]]
        if row["source_value"] != official[row["field_name"]]:
            raise RuntimeError(f"REVIEW_SOURCE_VALUE_STALE:{key}")
        if row["source_hash"] != snapshot_hash(official):
            raise RuntimeError(f"REVIEW_SOURCE_HASH_STALE:{key}")
        if row["reviewed_hash"] != digest(row["reviewed_value"]):
            raise RuntimeError(f"REVIEWED_VALUE_HASH_INVALID:{key}")
        if row["decision"] not in {"KEEP", "CORRECTED", "REVIEW_REQUIRED", "NO_SOURCE"}:
            raise RuntimeError(f"REVIEW_DECISION_INVALID:{key}")
        if row["decision"] == "REVIEW_REQUIRED":
            category_key = (official["cat1"], official["cat2"])
            if row["field_name"] != "cat2" or category_key not in approved_categories:
                raise RuntimeError(f"UNAPPROVED_REVIEW_REQUIRED:{key}")
            if row["reviewed_value"] != approved_categories[category_key]:
                raise RuntimeError(f"CATEGORY_VALUE_MISMATCH:{key}")
    pending = [row for row in rows if row["decision"] == "REVIEW_REQUIRED"]
    if len(rows) != 2535 or len(pending) != 130 or len(approved_categories) != 43:
        raise RuntimeError(f"FROZEN_SCOPE_MISMATCH:rows={len(rows)},pending={len(pending)},categories={len(approved_categories)}")
    return rows, approved_categories


def db_rows(db: sqlite3.Connection, language: str) -> dict[str, dict[str, Any]]:
    values = db.execute(
        "SELECT official_sku,name,cat1,cat2,spec,description,details,source,review_status,"
        "name_source,cat1_source,cat2_source,spec_source,description_source,details_source,"
        "freshness_status FROM product_localizations WHERE language=?", (language,)
    ).fetchall()
    return {
        text(row[0]): {
            **dict(zip(("name", "cat1", "cat2", "spec", "description", "details"), (text(v) for v in row[1:7]))),
            "source": row[7], "review_status": row[8],
            **dict(zip(("name_source", "cat1_source", "cat2_source", "spec_source", "description_source", "details_source", "freshness_status"), row[9:])),
        }
        for row in values
    }


def category_dictionary_preview(approved: dict[tuple[str, str], str]) -> tuple[list[dict[str, str]], int]:
    path = ROOT / "runtime/dictionary/category_dictionary.csv"
    rows = load_dictionary_rows(path, headers=CATEGORY_DICTIONARY_HEADERS, key_fields=("cat1_es", "cat2_es"))
    index = {(row["cat1_es"], row["cat2_es"]): row for row in rows}
    changed = 0
    for key, value in approved.items():
        row = index.get(key)
        if row is None:
            raise RuntimeError(f"CATEGORY_DICTIONARY_KEY_MISSING:{key}")
        if row["cat2_zh"] not in {"", value}:
            raise RuntimeError(f"CATEGORY_DICTIONARY_CONFLICT:{key}:{row['cat2_zh']}!={value}")
        if row["cat2_zh"] != value:
            row["cat2_zh"] = value
            row["review_status"] = "HUMAN_REVIEWED"
            row["notes"] = (row.get("notes", "").strip() + "；" if row.get("notes", "").strip() else "") + "2026-09-21 Owner-authorized category mapping"
            changed += 1
    return rows, changed


def product_dictionary_preview(source: dict[str, dict[str, str]], reviews: list[dict[str, str]]) -> tuple[list[dict[str, str]], int]:
    path = ROOT / "runtime/dictionary/product_dictionary.csv"
    rows = load_dictionary_rows(path, headers=PRODUCT_DICTIONARY_HEADERS, key_fields=("sku",))
    index = {row["sku"]: row for row in rows}
    by_sku: dict[str, dict[str, str]] = {}
    for review in reviews:
        if review["field_name"] in PRODUCT_DICTIONARY_FIELDS:
            by_sku.setdefault(review["sku"], {})[review["field_name"]] = review["reviewed_value"]
    changed = 0
    stamp = datetime.now(timezone.utc).date().isoformat()
    for sku, values in by_sku.items():
        row = index.get(sku)
        if row is None:
            # The dictionary is a growing long-lived asset.  A current SKU
            # absent from an older dictionary is not a reason to discard an
            # owner-approved field repair; create a minimal, traceable row.
            row = {header: "" for header in PRODUCT_DICTIONARY_HEADERS}
            row.update({"sku": sku, "canonical_id": f"ACT{sku}", "locked": "0"})
            rows.append(row)
            index[sku] = row
        official = source[sku]
        row_changed = False
        for product_field, source_field in (("name_es_raw", "name"), ("cat1_es", "cat1"), ("cat2_es", "cat2"), ("spec_es_raw", "spec")):
            if row[product_field] != official[source_field]:
                row[product_field] = official[source_field]
                row_changed = True
        for field, value in values.items():
            target = PRODUCT_DICTIONARY_FIELDS[field]
            if row[target] != value:
                row[target] = value
                row_changed = True
        if row_changed:
            row["source_hash"] = product_source_hash(row)
            row["translation_status"] = "HUMAN_APPROVED"
            row["review_status"] = "HUMAN_REVIEWED"
            row["updated_at"] = stamp
            row["notes"] = (row.get("notes", "").strip() + "；" if row.get("notes", "").strip() else "") + "2026-09-21 reviewed localization apply"
            changed += 1
    return rows, changed


def backup_file(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)


def restore_database(backup: Path, target: Path) -> None:
    source = sqlite3.connect(backup)
    destination = sqlite3.connect(target)
    try:
        source.backup(destination)
        destination.commit()
    finally:
        destination.close()
        source.close()


def insert_patch(db: sqlite3.Connection, *, sku: str, field: str, before: str, after: str, source_hash: str, run_id: str, now: str) -> None:
    latest = db.execute(
        "SELECT patch_id,revision FROM localization_patches WHERE official_sku=? AND language='zh' AND field_name=? ORDER BY revision DESC LIMIT 1",
        (sku, field),
    ).fetchone()
    parent, revision = (latest[0], int(latest[1]) + 1) if latest else (None, 1)
    patch_id = hashlib.sha256(f"{run_id}|{sku}|{field}|{digest(before)}|{digest(after)}".encode()).hexdigest()
    db.execute(
        "INSERT INTO localization_patches(patch_id,parent_patch_id,official_sku,language,field_name,old_value,new_value,source_hash,reason,created_by,created_at,revision) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
        (patch_id, parent, sku, "zh", field, before, after, source_hash, "2026-09-21 Owner-authorized reviewed localization repair", "project-owner", now, revision),
    )
    for event in ("PATCH_CREATED", "PATCH_APPROVED", "PATCH_APPLIED"):
        event_id = hashlib.sha256(f"{patch_id}|{event}".encode()).hexdigest()
        db.execute(
            "INSERT INTO localization_patch_events(event_id,patch_id,event_type,actor,reason,event_json,occurred_at) VALUES(?,?,?,?,?,?,?)",
            (event_id, patch_id, event, "project-owner", "2026-09-21 explicit repair authorization", json.dumps({"apply_run_id": run_id, "sku": sku, "field": field}, ensure_ascii=False), now),
        )


def apply_database(db_path: Path, source: dict[str, dict[str, str]], reviews: list[dict[str, str]], run_id: str) -> dict[str, int]:
    now = datetime.now(timezone.utc).isoformat()
    counts = {"spanish_fact_fields": 0, "zh_review_fields": 0, "zh_value_updates": 0, "patches": 0}
    with connect(db_path) as db:
        db.row_factory = sqlite3.Row
        db.execute("BEGIN IMMEDIATE")
        role = db.execute("SELECT value FROM schema_metadata WHERE key='database_role'").fetchone()
        if role is None or role[0] != "PRIMARY":
            raise RuntimeError("PRIMARY_DATABASE_REQUIRED")
        es = db_rows(db, "es")
        zh = db_rows(db, "zh")
        # SQLite retains historical localizations; the source workbook is the
        # 5,510-SKU current observation subset.  All current source SKUs must
        # exist in both language projections, while historical rows remain
        # intentionally untouched.
        missing_es = sorted(set(source) - set(es))
        missing_zh = sorted(set(source) - set(zh))
        if missing_es or missing_zh:
            raise RuntimeError(
                f"CURRENT_SOURCE_SKU_MISSING:es={missing_es[:10]},zh={missing_zh[:10]}"
            )
        for sku, official in source.items():
            existing = es[sku]
            changed_source_fields = [field for field in ALL_FIELDS if existing[field] != official[field]]
            final_hash = localization_source_hash({
                "name_es": official["name"], "cat1_es": official["cat1"], "cat2_es": official["cat2"],
                "spec_es": official["spec"], "desc_es": official["description"], "details_es": official["details"],
            })
            product = db.execute("SELECT product_url,image_url FROM products WHERE official_sku=?", (sku,)).fetchone()
            if product is None:
                raise RuntimeError(f"PRODUCT_ROW_MISSING:{sku}")
            if changed_source_fields:
                db.execute(
                    "UPDATE product_localizations SET name=?,cat1=?,cat2=?,spec=?,description=?,details=?,source_hash=?,updated_at=?,applied_commit_id=? WHERE official_sku=? AND language='es'",
                    (official["name"], official["cat1"], official["cat2"], official["spec"], official["description"], official["details"], final_hash, now, run_id, sku),
                )
                counts["spanish_fact_fields"] += len(changed_source_fields)
            else:
                db.execute("UPDATE product_localizations SET source_hash=? WHERE official_sku=? AND language='es'", (final_hash, sku))
            db.execute(
                "UPDATE products SET name_es=?,source_hash=?,content_hash=?,updated_at=? WHERE official_sku=?",
                (official["name"], final_hash, content_hash({
                    "name_es": official["name"], "cat1_es": official["cat1"], "cat2_es": official["cat2"],
                    "spec_es": official["spec"], "desc_es": official["description"], "details_es": official["details"],
                    "product_url": product[0], "image_url": product[1],
                }), now, sku),
            )
            db.execute("UPDATE product_localizations SET source_hash=? WHERE official_sku=? AND language='zh'", (final_hash, sku))
            db.execute("UPDATE localization_fields SET source_hash=? WHERE official_sku=?", (final_hash, sku))
            db.execute("UPDATE localization_field_provenance SET source_hash=? WHERE official_sku=?", (final_hash, sku))
            sync_localization_field_provenance(
                db, {"sku": sku, "language": "es", **official, "source": existing["source"],
                     "review_status": existing["review_status"], "source_hash": final_hash, "applied_commit_id": run_id},
                commit_id=run_id, now=now,
            )

        for review in reviews:
            sku, field, value = review["sku"], review["field_name"], review["reviewed_value"]
            official = source[sku]
            final_hash = localization_source_hash({
                "name_es": official["name"], "cat1_es": official["cat1"], "cat2_es": official["cat2"],
                "spec_es": official["spec"], "desc_es": official["description"], "details_es": official["details"],
            })
            before = zh[sku][field]
            column = LOCALIZATION_COLUMNS[field]
            source_type = "category_dictionary" if field in {"cat1", "cat2"} else "qwen_mt_flash_reviewed"
            db.execute(
                f"UPDATE product_localizations SET {column}=?,{field}_source=?,updated_at=?,last_commit_id=?,approved_by=?,approved_at=?,applied_commit_id=? WHERE official_sku=? AND language='zh'",
                (value, source_type, now, run_id, "project-owner", now, run_id, sku),
            )
            sync_localization_field_provenance(
                db, {"sku": sku, "language": "zh", field: value, f"{field}_source": source_type,
                     f"{field}_review_status": "HUMAN_APPROVED", f"{field}_source_hash": final_hash,
                     f"{field}_updated_at": now, f"{field}_applied_commit_id": run_id,
                     f"{field}_approved_by": "project-owner", f"{field}_approved_at": now,
                     f"{field}_freshness_status": "FRESH"},
                commit_id=run_id, now=now,
            )
            if field == "name":
                db.execute("UPDATE products SET name_zh=?,updated_at=? WHERE official_sku=?", (value, now, sku))
            if before != value:
                insert_patch(db, sku=sku, field=field, before=before, after=value, source_hash=final_hash, run_id=run_id, now=now)
                counts["zh_value_updates"] += 1
                counts["patches"] += 1
            counts["zh_review_fields"] += 1
        db.commit()
    return counts


def post_verify(db_path: Path, source: dict[str, dict[str, str]], reviews: list[dict[str, str]]) -> dict[str, Any]:
    with connect(db_path) as db:
        es = db_rows(db, "es")
        zh = db_rows(db, "zh")
        source_mismatch = [sku for sku in source if any(es[sku][field] != source[sku][field] for field in ALL_FIELDS)]
        review_mismatch = [f"{row['sku']}|{row['field_name']}" for row in reviews if zh[row["sku"]][row["field_name"]] != row["reviewed_value"]]
        product_name_mismatch = [row["sku"] for row in reviews if row["field_name"] == "name" and text(db.execute("SELECT name_zh FROM products WHERE official_sku=?", (row["sku"],)).fetchone()[0]) != row["reviewed_value"]]
        integrity = db.execute("PRAGMA integrity_check").fetchone()[0]
    return {
        "status": "PASS" if not source_mismatch and not review_mismatch and not product_name_mismatch and integrity == "ok" else "FAIL",
        "source_fact_mismatch_count": len(source_mismatch), "reviewed_value_mismatch_count": len(review_mismatch),
        "product_name_mismatch_count": len(product_name_mismatch), "integrity_check": integrity,
        "source_fact_mismatch_sample": source_mismatch[:20], "review_mismatch_sample": review_mismatch[:20],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="perform the authorized write; default is a dry-run")
    args = parser.parse_args()
    cfg = load_settings()
    source = read_source()
    reviews, approved_categories = validated_reviews(source)
    category_rows, category_changes = category_dictionary_preview(approved_categories)
    product_rows, product_changes = product_dictionary_preview(source, reviews)
    db_path = database_path(cfg)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_id = f"localization-repair-{stamp}"
    out = ROOT / "runtime/localization_apply" / f"20260921_{stamp}"
    preview = {
        "batch_id": BATCH_ID, "policy_version": POLICY_VERSION, "run_id": run_id,
        "mode": "APPLY" if args.apply else "DRY_RUN", "source_sku_count": len(source),
        "review_field_count": len(reviews), "category_mapping_count": len(approved_categories),
        "category_dictionary_changes": category_changes, "product_dictionary_changes": product_changes,
        "source_snapshot_sha256": file_hash(SOURCE_XLSX), "review_csv_sha256": file_hash(REVIEW_CSV),
        "category_csv_sha256": file_hash(CATEGORY_CSV), "database_sha256_before": file_hash(db_path),
        "master_sha256_before": file_hash(Path(cfg["paths"]["master"])),
        "explicit_owner_authorization": "2026-09-21 用户指令：有问题的一起修复掉；西语元数据不对时可重新补全再翻译",
    }
    out.mkdir(parents=True, exist_ok=True)
    (out / "apply_preview.json").write_text(json.dumps(preview, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if not args.apply:
        print(json.dumps({**preview, "status": "READY_FOR_EXPLICIT_APPLY"}, ensure_ascii=False, indent=2))
        return 0

    master = Path(cfg["paths"]["master"])
    state = Path(cfg["paths"]["state"])
    backups = {
        "database": out / "before_action_tracker.db", "master": out / "before_Action_Master.xlsx",
        "category_dictionary": out / "before_category_dictionary.csv", "product_dictionary": out / "before_product_dictionary.csv",
        "known_skus": out / "before_known_skus.csv", "offline_skus": out / "before_offline_skus.csv",
    }
    backup_database(db_path, backups["database"])
    for key, path in (("master", master), ("category_dictionary", ROOT / "runtime/dictionary/category_dictionary.csv"),
                      ("product_dictionary", ROOT / "runtime/dictionary/product_dictionary.csv"),
                      ("known_skus", state / "known_skus.csv"), ("offline_skus", state / "offline_skus.csv")):
        backup_file(path, backups[key])
    pre_hashes = {key: file_hash(path) for key, path in (("database", db_path), ("master", master), ("category_dictionary", ROOT / "runtime/dictionary/category_dictionary.csv"), ("product_dictionary", ROOT / "runtime/dictionary/product_dictionary.csv"), ("known_skus", state / "known_skus.csv"), ("offline_skus", state / "offline_skus.csv"))}
    try:
        database_counts = apply_database(db_path, source, reviews, run_id)
        write_dictionary_csv(ROOT / "runtime/dictionary/category_dictionary.csv", category_rows, CATEGORY_DICTIONARY_HEADERS, key_fields=("cat1_es", "cat2_es"))
        write_dictionary_csv(ROOT / "runtime/dictionary/product_dictionary.csv", product_rows, PRODUCT_DICTIONARY_HEADERS, key_fields=("sku",))
        from action_tracker.database.repository import ProductionRepository
        head = ProductionRepository(db_path).current_head()
        if not head:
            raise RuntimeError("SQLITE_PRIMARY_HEAD_MISSING")
        sync = regenerate_compatibility_exports(cfg, head)
        verification = post_verify(db_path, source, reviews)
        if verification["status"] != "PASS":
            raise RuntimeError("POST_APPLY_VERIFICATION_FAILED:" + json.dumps(verification, ensure_ascii=False))
    except Exception:
        # Roll back all writable targets controlled by this tool.  The SQLite
        # backup API avoids a torn WAL restoration.
        restore_database(backups["database"], db_path)
        for key, target in (("master", master), ("category_dictionary", ROOT / "runtime/dictionary/category_dictionary.csv"),
                            ("product_dictionary", ROOT / "runtime/dictionary/product_dictionary.csv"),
                            ("known_skus", state / "known_skus.csv"), ("offline_skus", state / "offline_skus.csv")):
            shutil.copy2(backups[key], target)
        raise
    after_hashes = {key: file_hash(path) for key, path in (("database", db_path), ("master", master), ("category_dictionary", ROOT / "runtime/dictionary/category_dictionary.csv"), ("product_dictionary", ROOT / "runtime/dictionary/product_dictionary.csv"), ("known_skus", state / "known_skus.csv"), ("offline_skus", state / "offline_skus.csv"))}
    report = {**preview, "status": "SUCCESS", "database_counts": database_counts, "compatibility_sync": sync,
              "verification": verification, "backups": {key: str(path) for key, path in backups.items()},
              "hashes_before": pre_hashes, "hashes_after": after_hashes,
              "browser_evidence": {"sku": "3222971", "url": "https://www.action.com/es-es/p/3222971/", "official_breadcrumb": "Multimedia > Cascos de teléfono", "conclusion": "保留官网一级/二级类目，不因商品身份看似异常而改写西语来源"}}
    (out / "apply_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
