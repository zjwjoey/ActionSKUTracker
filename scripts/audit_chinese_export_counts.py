from pathlib import Path
import re
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
path = root / "runtime" / "exports" / "20260907Action商品全量_中文版_不带图.xlsx"
ws = load_workbook(path, read_only=True, data_only=True).active
rows = list(ws.iter_rows(values_only=True))
h = {str(v): i for i, v in enumerate(rows[0])}
def col(*names):
    return next(i for k, i in h.items() if any(n in k for n in names))
cols = {"name": col("标题", "品名"), "cat1": col("分类1"), "cat2": col("分类2"), "spec": col("规格"), "desc": col("描述"), "details": col("产品详情")}
stats = {}
for key, idx in cols.items():
    vals = [str(r[idx] or '').strip() for r in rows[1:]]
    stats[key] = {"blank": sum(not v for v in vals), "no_hanzi": sum(bool(v) and not re.search(r"[\u4e00-\u9fff]", v) for v in vals)}
any_blank = sum(any(not str(r[i] or '').strip() for i in cols.values()) for r in rows[1:])
any_no_hanzi = sum(any(str(r[i] or '').strip() and not re.search(r"[\u4e00-\u9fff]", str(r[i] or '')) for i in cols.values()) for r in rows[1:])
print({"rows": len(rows)-1, "any_blank": any_blank, "any_no_hanzi": any_no_hanzi, "fields": stats})
