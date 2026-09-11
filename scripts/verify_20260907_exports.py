from pathlib import Path
import json
from openpyxl import load_workbook
from PIL import Image

root = Path(__file__).resolve().parents[1]
for name in ("20260907Action商品全量_西班牙语版_不带图.xlsx", "20260907Action商品全量_西班牙语版_带图.xlsx"):
    path = root / "runtime" / "exports" / name
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    idx = {str(v): i for i, v in enumerate(headers)}
    body = rows[1:]
    print(name, {"rows": len(body), "blank_cat2": sum(not str(r[idx.get("二级类目（西语）", -1)] or "").strip() for r in body),
                 "image_http": sum(str(r[idx.get("图片链接", -1)] or "").startswith("http") for r in body),
                 "product_http": sum(str(r[idx.get("商品链接", -1)] or "").startswith("http") for r in body)})
    wb.close()
sample = root / "runtime" / "images" / "derivatives" / "excel_250" / "2520562.png"
with Image.open(sample) as im:
    print("image_sample", im.size, im.mode, im.getpixel((0, 0)))
