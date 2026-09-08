"""Build a read-only, auditable category repair candidate.

This command never writes the production SQLite database, Master projection,
state files, or runtime artifacts.  It reads the active production projection
in read-only mode and emits a candidate workbook plus patch/audit evidence in
the data-closure worktree.
"""
from __future__ import annotations

import csv
import hashlib
import json
import shutil
import sqlite3
import argparse
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]

# These are official product-page breadcrumb facts already checked separately;
# they are proposed evidence only and are not applied to PRIMARY here.
OFFICIAL_CAT2 = {
    "3217313": {
        "cat1_es": "Mascotas", "cat2_es": "Juguetes para mascotas",
        "url": "https://www.action.com/es-es/p/3217313/juguetes-para-perro/",
    },
    "3221995": {
        "cat1_es": "Juguetes", "cat2_es": "Juegos simbólicos y de construcción",
        "url": "https://www.action.com/es-es/p/3221995/set-de-juego-spidey/",
    },
    "3225631": {
        "cat1_es": "Cocina", "cat2_es": "Organizadores de despensa",
        "url": "https://www.action.com/es-es/p/3225631/botella-aislante-tal-ranger-pro/",
    },
    "3227020": {
        "cat1_es": "Cuidado personal", "cat2_es": "Salud",
        "url": "https://www.action.com/es-es/p/3227020/apositos-first-aid-sensitive/",
    },
}

# Suggestions are deliberately emitted into a review queue only.  They are
# not written to the dictionary and cannot be applied by this command.
SUGGESTED_ZH_CAT2 = {
    ("Cuidado personal", "Depilación y afeitado"): "脱毛与剃须",
    ("Moda", "Relojes y joyas"): "手表与珠宝",
    ("Comer y beber", "Bebidas"): "饮料",
    ("Mascotas", "Alimentación para animales"): "宠物食品",
    ("Cuidado personal", "Maquillaje"): "彩妆",
    ("Comer y beber", "Galletas"): "饼干",
    ("Mascotas", "Perro"): "狗狗用品",
    ("Cuidado personal", "Cuidado facial"): "面部护理",
    ("Vivienda", "Accesorios de baño"): "浴室配件",
    ("Jardín", "Decoración para el jardín"): "花园装饰",
    ("Multimedia", "Audio"): "音频",
    ("Jardín", "Iluminación exterior"): "户外照明",
    ("Oficina y papelería", "Cartuchos de tinta"): "墨盒",
    ("Multimedia", "Accesorios multimedia"): "多媒体配件",
    ("Multimedia", "Pilas"): "电池",
    ("Vivienda", "Muebles"): "家具",
    ("Cuidado personal", "Cuidados para el bebé"): "婴儿护理",
    ("Comer y beber", "Alimentación"): "食品",
    ("Multimedia", "Cables y divisores"): "线材与分配器",
    ("Vivienda", "Lámparas"): "灯具",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_dictionary(path: Path) -> dict[tuple[str, str], dict[str, str]]:
    # The active workspace may contain locally reviewed mappings that are not
    # yet part of the remote main branch.  Read them as candidate evidence;
    # never copy them into PRIMARY or the closure branch automatically.
    path = Path(path)
    out: dict[tuple[str, str], dict[str, str]] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            key = (str(row.get("cat1_es") or "").strip(), str(row.get("cat2_es") or "").strip())
            if key[0] and key[1]:
                out[key] = row
    return out


def load_current(db_path: Path) -> list[dict[str, object]]:
    uri = f"file:{Path(db_path).absolute().as_posix()}?mode=ro"
    query = """
    SELECT p.official_sku, p.status,
           es.cat1 AS cat1_es, es.cat2 AS cat2_es,
           zh.cat1 AS cat1_zh, zh.cat2 AS cat2_zh
    FROM products p
    LEFT JOIN product_localizations es
      ON es.official_sku=p.official_sku AND es.language='es'
    LEFT JOIN product_localizations zh
      ON zh.official_sku=p.official_sku AND zh.language='zh'
    WHERE p.status='CURRENT'
    ORDER BY p.official_sku
    """
    with sqlite3.connect(uri, uri=True) as db:
        db.row_factory = sqlite3.Row
        return [dict(row) for row in db.execute(query)]


def build_patch(rows: list[dict[str, object]], category_map: dict[tuple[str, str], dict[str, str]]) -> list[dict[str, object]]:
    patches: list[dict[str, object]] = []
    for row in rows:
        sku = str(row["official_sku"])
        es_cat1 = str(row.get("cat1_es") or "").strip()
        es_cat2 = str(row.get("cat2_es") or "").strip()
        zh_cat2 = str(row.get("cat2_zh") or "").strip()
        if not zh_cat2 and es_cat2:
            candidate = category_map.get((es_cat1, es_cat2))
            if candidate and str(candidate.get("cat2_zh") or "").strip():
                patches.append({
                    "sku": sku, "language": "zh", "field": "cat2",
                    "old_value": "", "new_value": str(candidate["cat2_zh"]).strip(),
                    "source": "category_dictionary",
                    "review_status": str(candidate.get("review_status") or "REVIEW_REQUIRED"),
                    "evidence": f"category_dictionary:{es_cat1}|{es_cat2}",
                    "apply_status": "CANDIDATE_ONLY",
                })
            else:
                patches.append({
                    "sku": sku, "language": "zh", "field": "cat2",
                    "old_value": "", "new_value": "", "source": "category_dictionary",
                    "review_status": "REVIEW_REQUIRED",
                    "evidence": f"missing_mapping:{es_cat1}|{es_cat2}",
                    "apply_status": "BLOCKED",
                })
        if not es_cat2:
            evidence = OFFICIAL_CAT2.get(sku)
            patches.append({
                "sku": sku, "language": "es", "field": "cat2",
                "old_value": "", "new_value": evidence["cat2_es"] if evidence else "",
                "source": "official_product_breadcrumb" if evidence else "missing_official_evidence",
                "review_status": "VERIFIED" if evidence else "REVIEW_REQUIRED",
                "evidence": evidence["url"] if evidence else "",
                "apply_status": "CANDIDATE_ONLY" if evidence else "BLOCKED",
            })
    return patches


def write_candidate(patches: list[dict[str, object]], master_path: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    candidate = output_dir / "Action_Master_data_repair_candidate_20260908.xlsx"
    shutil.copy2(master_path, candidate)
    wb = openpyxl.load_workbook(candidate)
    try:
        for sheet, language in (("01_SKU_ZH_CURRENT", "zh"), ("02_SKU_ES_CURRENT", "es")):
            ws = wb[sheet]
            headers = [cell.value for cell in ws[1]]
            sku_col = headers.index("SKU") + 1
            cat2_name = "二级类目（中文）" if language == "zh" else "二级类目（西语）"
            cat2_col = headers.index(cat2_name) + 1
            by_sku = {str(ws.cell(r, sku_col).value or "").strip(): r for r in range(2, ws.max_row + 1)}
            for item in patches:
                if item["language"] != language or item["apply_status"] != "CANDIDATE_ONLY":
                    continue
                row_no = by_sku.get(str(item["sku"]))
                if row_no and item["new_value"]:
                    ws.cell(row_no, cat2_col).value = item["new_value"]
        wb.save(candidate)
    finally:
        wb.close()
    return candidate


def write_review_queue(rows: list[dict[str, object]], patches: list[dict[str, object]], output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    queue = output_dir / "category_mapping_review_queue.csv"
    counts: dict[tuple[str, str], int] = {}
    for item in patches:
        if item["apply_status"] != "BLOCKED" or item["language"] != "zh":
            continue
        prefix, pair = str(item["evidence"]).split(":", 1)
        cat1, cat2 = pair.split("|", 1)
        counts[(cat1, cat2)] = counts.get((cat1, cat2), 0) + 1
    with queue.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=(
            "cat1_es", "cat2_es", "sku_count", "suggested_cat2_zh",
            "review_status", "reason",
        ))
        writer.writeheader()
        for (cat1, cat2), count in sorted(counts.items()):
            writer.writerow({
                "cat1_es": cat1, "cat2_es": cat2, "sku_count": count,
                "suggested_cat2_zh": SUGGESTED_ZH_CAT2.get((cat1, cat2), ""),
                "review_status": "REVIEW_REQUIRED",
                "reason": "当前正式字典无 cat2_zh；候选译名仅供人工确认，不得自动 Apply",
            })
    return queue


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-db", type=Path, required=True)
    parser.add_argument("--source-master", type=Path, required=True)
    parser.add_argument("--dictionary-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    rows = load_current(args.source_db)
    category_map = load_dictionary(args.dictionary_root / "category_dictionary.csv")
    patches = build_patch(rows, category_map)
    candidate = write_candidate(patches, args.source_master, args.output_dir)
    review_queue = write_review_queue(rows, patches, args.output_dir)
    patch_path = args.output_dir / "category_repair_candidates.json"
    patch_path.write_text(json.dumps(patches, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "production_write": False,
        "database_mode": "READ_ONLY",
        "source_database": str(args.source_db),
        "source_database_sha256": sha256(args.source_db),
        "source_master": str(args.source_master),
        "source_master_sha256": sha256(args.source_master),
        "candidate_dictionary": str(args.dictionary_root / "category_dictionary.csv"),
        "candidate_dictionary_sha256": sha256(args.dictionary_root / "category_dictionary.csv"),
        "current_sku_count": len(rows),
        "candidate_count": len(patches),
        "zh_cat2_candidates": sum(p["language"] == "zh" for p in patches),
        "es_cat2_candidates": sum(p["language"] == "es" for p in patches),
        "verified_candidates": sum(p["apply_status"] == "CANDIDATE_ONLY" for p in patches),
        "blocked_candidates": sum(p["apply_status"] == "BLOCKED" for p in patches),
        "candidate_workbook": str(candidate),
        "patch_file": str(patch_path),
        "review_queue": str(review_queue),
        "blocked_pair_count": len({p["evidence"] for p in patches if p["apply_status"] == "BLOCKED"}),
    }
    (args.output_dir / "audit.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
