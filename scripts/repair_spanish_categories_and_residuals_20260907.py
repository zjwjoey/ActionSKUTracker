"""Repair verified Spanish export categories and the last text artefacts.

Only changes backed by the official Action product-page breadcrumb are applied.
No prices, titles, specifications, URLs or product facts are regenerated.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / "src"))

from action_tracker.exporting.excel_writer import write_catalog_xlsx
from openpyxl import load_workbook


SOURCE = Path(
    r"F:\ActionSKUTracker\runtime\exports\20260907Action商品全量_西班牙语版_不带图_最终修复版_格式清理版.xlsx"
)
OUTPUT = SOURCE.with_name(
    "20260907Action商品全量_西班牙语版_不带图_最终修复版_分类清理版.xlsx"
)
AUDIT = OUTPUT.with_suffix(".audit.json")
HEADERS = [
    "图片", "编号", "标题", "分类1", "分类2", "规格", "折后价", "原价",
    "单价", "描述", "产品详情", "图片链接", "商品链接", "备注",
]

# Product-page breadcrumb evidence captured through the browser plugin on 2026-09-07.
CATEGORY_FIXES = {
    "2535685": ("Moda", "Ropa"),
    "3210419": ("Oficina y papelería", "Accesorios de oficina"),
    "2577107": ("Bricolaje", "Coche"),
    "3222519": ("Viajes", "Artículos de acampada"),
    "3007683": ("Oficina y papelería", "Accesorios de oficina"),
    "3216603": ("Oficina y papelería", "Accesorios de oficina"),
    "3223884": ("Oficina y papelería", "Accesorios de oficina"),
    "3222499": ("Bricolaje", "Herramientas"),
    "3206019": ("Hobby", "Artículos de fiesta"),
    "3205740": ("Viajes", "Accesorios de viaje"),
    "3205891": ("Viajes", "Accesorios de viaje"),
    "3209038": ("Viajes", "Accesorios de viaje"),
    "3011954": ("Oficina y papelería", "Artículos de papel"),
    "3207974": ("Oficina y papelería", "Artículos de papel"),
    "3208746": ("Oficina y papelería", "Material escolar"),
    "3222425": ("Oficina y papelería", "Material escolar"),
}

DESCRIPTION_ARTIFACT_SKUS = {"2510061", "2580205"}
READ_MORE_SKUS = {"2576247", "3217077"}


def clean_description(value: object, sku: str) -> str:
    text = "" if value is None else str(value)
    if sku in READ_MORE_SKUS:
        text = re.sub(r"(?im)^\s*Leer más\s*$", "", text)
    if sku in DESCRIPTION_ARTIFACT_SKUS:
        # These two rows contain literal HTML-style `>` separators left by the
        # source renderer rather than meaningful product text.
        text = re.sub(r"\s*>+\s*", "\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    category_changes = []
    description_changes = []

    for source_row in ws.iter_rows(min_row=2):
        row = {h: source_row[i].value for h, i in idx.items()}
        sku = str(row["编号"]).strip()

        if sku in CATEGORY_FIXES:
            expected = CATEGORY_FIXES[sku]
            before = (row["分类1"], row["分类2"])
            if before != expected:
                row["分类1"], row["分类2"] = expected
                category_changes.append({
                    "sku": sku,
                    "before": {"分类1": before[0], "分类2": before[1]},
                    "after": {"分类1": expected[0], "分类2": expected[1]},
                })

        before_desc = row.get("描述") or ""
        after_desc = clean_description(before_desc, sku)
        if after_desc != before_desc:
            row["描述"] = after_desc
            description_changes.append({"sku": sku, "before": before_desc, "after": after_desc})

        rows.append(row)

    profile = {
        "sheet_name": "商品全量",
        "freeze_panes": "A2",
        "auto_filter": True,
        "header": {"bold": True, "fill": "1F4E78", "font_color": "FFFFFF"},
        "body": {
            "wrap_text_columns": ["标题", "分类1", "分类2", "规格", "描述", "产品详情", "备注"],
            "max_row_height": 405,
        },
        "price": {"number_format": "€#,##0.00"},
    }
    write_catalog_xlsx(OUTPUT, headers=HEADERS, rows=rows, workbook_format=profile)

    # Re-open the produced workbook so the audit describes the artifact that
    # was actually written, not only the in-memory transformation.
    check_wb = load_workbook(OUTPUT, read_only=False, data_only=True)
    check_ws = check_wb.active
    check_rows = list(check_ws.iter_rows(min_row=2, values_only=True))
    check_skus = [str(r[1]).strip() for r in check_rows]
    all_text = "\n".join(
        f"{r[9] or ''}\n{r[10] or ''}" for r in check_rows
    )
    qa = {
        "sku_count": len(check_rows),
        "unique_sku_count": len(set(check_skus)),
        "blank_sku": sum(not s or s == "None" for s in check_skus),
        "blank_title": sum(not str(r[2] or "").strip() for r in check_rows),
        "blank_category1": sum(not str(r[3] or "").strip() for r in check_rows),
        "blank_category2": sum(not str(r[4] or "").strip() for r in check_rows),
        "blank_description": sum(not str(r[9] or "").strip() for r in check_rows),
        "blank_detail": sum(not str(r[10] or "").strip() for r in check_rows),
        "html_tag_rows": sum(bool(re.search(r"<[^>]+>", f"{r[9] or ''} {r[10] or ''}")) for r in check_rows),
        "greater_artifact_rows": sum(">" in str(r[9] or "") for r in check_rows),
        "read_more_rows": sum(bool(re.search(r"(?im)^\s*Leer más\s*$", str(r[9] or ""))) for r in check_rows),
        "description_heading_rows": sum(bool(re.search(r"(?im)^\s*Descripción\s*(?::\s*)?$", str(r[9] or ""))) for r in check_rows),
        "null_undefined_rows": sum(bool(re.search(r"(?i)\b(?:null|undefined)\b", f"{r[9] or ''} {r[10] or ''}")) for r in check_rows),
        "replacement_character_rows": sum("\ufffd" in f"{r[9] or ''} {r[10] or ''}" for r in check_rows),
        "freeze_panes": check_ws.freeze_panes,
        "auto_filter": check_ws.auto_filter.ref,
    }
    check_wb.close()

    audit = {
        "source": str(SOURCE),
        "output": str(OUTPUT),
        "sku_count": len(rows),
        "category_fix_targets": sorted(CATEGORY_FIXES),
        "category_changes_applied": category_changes,
        "description_cleanup_targets": sorted(DESCRIPTION_ARTIFACT_SKUS | READ_MORE_SKUS),
        "description_changes_applied": description_changes,
        "qa": qa,
        "evidence": "Official Action product-page main breadcrumb verified through browser plugin on 2026-09-07.",
    }
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False))


if __name__ == "__main__":
    main()
