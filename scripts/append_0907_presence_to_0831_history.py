"""Append the 2026-09-07 presence column without touching 2026-08-31 data."""
from __future__ import annotations

import datetime as dt
import hashlib
import json
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


INPUT = Path(r"F:\ActionSKUTracker\runtime\exports\20260831_Action商品上下架明细.xlsx")
MASTER = Path(r"F:\ActionSKUTracker\runtime\master\Action_Master.xlsx")
OUTPUT = Path(r"F:\ActionSKUTracker\runtime\exports\20260907_Action商品上下架明细.xlsx")
MIRROR = Path(r"F:\按日期整理\20260907\20260907_Action商品上下架明细.xlsx")
AUDIT = Path(r"F:\ActionSKUTracker\runtime\exports\20260907_Action商品上下架明细.audit.json")
NEW_DATE = "26.09.07"


def load_current_skus() -> set[str]:
    workbook = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)
    try:
        sheet = workbook["01_SKU_ZH_CURRENT"]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        index = {str(value): pos for pos, value in enumerate(headers)}
        if not {"SKU", "当前状态"}.issubset(index):
            raise ValueError("MASTER_SCHEMA_MISMATCH")
        return {
            str(row[index["SKU"]]).strip()
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if str(row[index["当前状态"]] or "").strip().upper() == "CURRENT"
            and str(row[index["SKU"]] or "").strip()
        }
    finally:
        workbook.close()


def write(path: Path, headers: list[Any], rows: list[list[Any]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "商品上下架明细"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    widths = {"A": 8, "B": 14, "C": 30, "D": 65, "E": 65}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for col in range(6, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 11
        for row_no in range(2, sheet.max_row + 1):
            sheet.cell(row=row_no, column=col).number_format = "0"
    for row_no in range(2, sheet.max_row + 1):
        for col in (3, 4, 5):
            sheet.cell(row=row_no, column=col).alignment = Alignment(vertical="top", wrap_text=True)
        for col in (4, 5):
            value = sheet.cell(row=row_no, column=col).value
            if isinstance(value, str) and value.startswith(("http://", "https://")):
                sheet.cell(row=row_no, column=col).hyperlink = value
                sheet.cell(row=row_no, column=col).style = "Hyperlink"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.stem}.tmp.xlsx")
    workbook.save(temporary)
    workbook.close()
    temporary.replace(path)


def main() -> None:
    if not INPUT.exists():
        raise FileNotFoundError(INPUT)
    workbook = openpyxl.load_workbook(INPUT, read_only=True, data_only=True)
    try:
        sheet = workbook["商品上下架明细"] if "商品上下架明细" in workbook.sheetnames else workbook.active
        headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        rows = [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)]
    finally:
        workbook.close()
    if "26.08.31" not in headers or NEW_DATE in headers:
        raise ValueError("HISTORY_DATE_COLUMNS_INVALID")
    sku_col = headers.index("编号")
    old_col = headers.index("26.08.31")
    old_values = [row[old_col] for row in rows]
    if len({str(row[sku_col]).strip() for row in rows}) != len(rows):
        raise ValueError("HISTORY_DUPLICATE_SKU")
    current = load_current_skus()
    headers.append(NEW_DATE)
    for row in rows:
        row.append(1 if str(row[sku_col]).strip() in current else 0)
    write(OUTPUT, headers, rows)
    write(MIRROR, headers, rows)
    check = openpyxl.load_workbook(OUTPUT, read_only=True, data_only=True)
    try:
        out_sheet = check["商品上下架明细"]
        out_headers = list(next(out_sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        out_rows = [list(row) for row in out_sheet.iter_rows(min_row=2, values_only=True)]
    finally:
        check.close()
    new_col = out_headers.index(NEW_DATE)
    after_old_values = [row[old_col] for row in out_rows]
    audit = {
        "source_file": str(INPUT),
        "master_file": str(MASTER),
        "output_file": str(OUTPUT),
        "mirror_output_file": str(MIRROR),
        "latest_date": NEW_DATE,
        "sku_count": len(out_rows),
        "current_sku_count_from_master": len(current),
        "presence_one_count": sum(row[new_col] == 1 for row in out_rows),
        "presence_zero_count": sum(row[new_col] == 0 for row in out_rows),
        "date_columns": out_headers[5:],
        "april_date_included": "26.04.05" in out_headers,
        "aug31_values_preserved": old_values == after_old_values,
        "aug31_value_hash_before": hashlib.sha256(repr(old_values).encode()).hexdigest(),
        "aug31_value_hash_after": hashlib.sha256(repr(after_old_values).encode()).hexdigest(),
        "duplicate_sku_count": len(out_rows) - len({str(row[sku_col]).strip() for row in out_rows}),
        "blank_sku_count": sum(not str(row[sku_col] or "").strip() for row in out_rows),
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "qc_result": "PASS" if (
            len(out_rows) == len(current) + sum(
                1 for row in out_rows if str(row[sku_col]).strip() not in current
            )
            and old_values == after_old_values
            and "26.04.05" not in out_headers
            and all(row[new_col] in (0, 1) for row in out_rows)
        ) else "FAIL",
    }
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False))


if __name__ == "__main__":
    main()
