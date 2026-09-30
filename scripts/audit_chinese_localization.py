"""Read-only full-workbook Spanish-to-Chinese localization audit.

The scan emits candidate findings only. It never edits either workbook or
applies proposed terminology repairs.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[1]
APPROVED_TERM_CHECKER_PATH = ROOT / "src/action_tracker/translation/approved_terms.py"
APPROVED_TERM_REVIEW_POLICY_PATH = ROOT / "config/stage5/approved_term_review_policy.json"
APPROVED_CATEGORY_STATUSES = frozenset({
    "APPROVED", "CONFIRMED", "HUMAN_APPROVED", "HUMAN_REVIEWED", "LOCKED",
})
APPROVED_TERM_STATUSES = frozenset({
    "APPROVED", "CONFIRMED", "HUMAN_APPROVED", "HUMAN_REVIEWED", "LOCKED", "SEED_REVIEWED",
})
sys.path.insert(0, str(ROOT / "src"))

import openpyxl  # noqa: E402

from action_tracker.stage5.source_candidate_v2 import (  # noqa: E402
    source_consistency_evidence, source_consistency_rules_manifest,
)
from action_tracker.products.details_parser import parse_details  # noqa: E402
from action_tracker.translation.detail_terminology import (  # noqa: E402
    detail_rule_context_matches, detail_value_rule_matches_source,
    load_detail_terminology_rules, repair_detail_candidate,
)
from action_tracker.translation.description_fidelity import (  # noqa: E402
    description_compression_finding, load_description_fidelity_policy,
)
from action_tracker.translation.approved_terms import (  # noqa: E402
    inspect_approved_term_candidate, inspect_approved_term_cross_field_movement,
    load_approved_term_review_policy,
)
from action_tracker.translation.model_guard import validate_model_output  # noqa: E402
from action_tracker.translation.source_fact_repair import repair_model_output as inspect_source_facts  # noqa: E402
from action_tracker.translation.title_policy import load_title_display_policy  # noqa: E402
from action_tracker.translation.regression_cases import (  # noqa: E402
    load_regression_cases, summarize_occurrences,
)
from action_tracker.services.normalization import parse_price  # noqa: E402


FIELD_ALIASES = {
    "sku": ("编号", "SKU", "sku", "official_sku"),
    "name": ("标题", "标题（西语）", "标题(西语)", "中文品名", "品名", "名称", "Título", "Nombre", "name", "name_es", "name_zh"),
    "cat1": ("分类1", "分类1（西语）", "分类1(西语)", "一级分类", "一级类目", "Categoría 1", "Category 1", "cat1", "cat1_es", "cat1_zh"),
    "cat2": ("分类2", "分类2（西语）", "分类2(西语)", "二级分类", "二级类目", "Categoría 2", "Category 2", "cat2", "cat2_es", "cat2_zh"),
    "spec": ("规格", "规格（西语）", "规格(西语)", "中文规格", "Especificación", "Specification", "spec", "spec_es", "spec_zh"),
    "description": ("描述", "描述（西语）", "描述(西语)", "Descripción", "description", "desc_es", "desc_zh", "description_es", "description_zh"),
    "details": ("产品详情", "产品详情（西语）", "产品详情(西语)", "Detalles", "Details", "details", "details_es", "details_zh"),
}
FACT_FIELD_ALIASES = {
    "current_price": ("折后价", "Precio rebajado", "Precio de oferta", "Precio actual", "Current price"),
    "original_price": ("原价", "Precio anterior", "Precio original", "Original price"),
    "image_url": ("图片链接", "URL imagen", "Enlace imagen", "Image URL", "Image link"),
    "product_url": ("商品链接", "URL producto", "Enlace producto", "Product URL", "Product link"),
}
REMARK_FIELD_ALIASES = ("备注", "Observaciones", "Remarks", "Notas")
RAW_TAG_FIELD_ALIASES = ("raw_tags", "raw_badges", "Etiquetas oficiales", "Official labels")
BUSINESS_TAG_FIELD_ALIASES = {
    "new": ("is_new_badge", "action_new_badge"),
    "promotion": ("promotion", "promotion_active"),
    "sustainable": ("sustainable", "sustainable_badge"),
}
GUARD_CODES = frozenset({
    "NUMERIC_DROPPED", "NUMERIC_HALLUCINATED", "UNIT_DROPPED", "UNIT_HALLUCINATED",
    "TECH_TOKEN_DROPPED", "TECH_TOKEN_HALLUCINATED", "CERTIFICATION_DROPPED",
    "CERTIFICATION_HALLUCINATED", "INTERNAL_QA_NOTE_LEAKED",
    "NEGATION_DROPPED", "NEGATION_HALLUCINATED", "UNSUPPORTED_NEGATIVE_ATTRIBUTE",
    "SOURCE_EMPTY_NONEMPTY", "EMPTY_REQUIRED_FIELD", "INVALID_CATEGORY",
    "SOURCE_FACT_REVIEW_REQUIRED",
})
_NUMBER = re.compile(r"\d")


def _norm(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold()
    text = "".join(ch for ch in unicodedata.normalize("NFD", text) if unicodedata.category(ch) != "Mn")
    return " ".join(text.split())


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _stable_issue_value(value: Any) -> str:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return _norm(value)


def _dedupe_issues(issues: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Collapse repeated findings without discarding variant evidence.

    A finding identity is the SKU, field, code and normalized source fact. A
    different candidate or evidence payload is retained as a bounded variant
    list on the surviving finding instead of inflating the issue count.
    """
    grouped: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    duplicate_count = 0
    for issue in issues:
        source = _stable_issue_value(issue.get("source"))
        identity_payload = "|".join((
            str(issue.get("sku") or "").strip(),
            str(issue.get("field") or "").strip(),
            str(issue.get("code") or "").strip(),
            source,
        ))
        finding_id = "finding-" + hashlib.sha256(identity_payload.encode("utf-8")).hexdigest()[:24]
        if finding_id not in grouped:
            kept = dict(issue)
            kept["finding_id"] = finding_id
            kept["duplicate_count"] = 0
            grouped[finding_id] = kept
            order.append(finding_id)
            continue
        duplicate_count += 1
        kept = grouped[finding_id]
        kept["duplicate_count"] = int(kept.get("duplicate_count") or 0) + 1
        for key in ("candidate", "evidence", "proposed"):
            value = issue.get(key)
            if value in (None, "") or value == kept.get(key):
                continue
            variants = kept.setdefault(f"{key}_variants", [])
            if value not in variants and len(variants) < 20:
                variants.append(value)
    return [grouped[key] for key in order], {
        "raw_issue_count": len(issues),
        "unique_issue_count": len(grouped),
        "duplicate_issue_count": duplicate_count,
    }


def _read_sheet(path: Path, sheet_name: str) -> tuple[list[str], dict[str, dict[str, Any]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"SHEET_NOT_FOUND:{path.name}:{sheet_name}")
        rows = workbook[sheet_name].iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        columns: dict[str, int] = {}
        for index, name in enumerate(headers):
            normalized = _norm(name)
            if not normalized:
                continue
            if normalized in columns:
                first = columns[normalized] + 1
                raise ValueError(
                    f"DUPLICATE_NORMALIZED_HEADER:{path.name}:{normalized}:columns={first},{index + 1}"
                )
            columns[normalized] = index
        resolved: dict[str, int] = {}

        def resolve_alias(field: str, aliases: tuple[str, ...]) -> int | None:
            # Aliases can differ only by case (for example ``SKU`` and
            # ``sku``). Collapse those before ambiguity checking so one actual
            # column is not mistaken for two competing matches.
            matches: dict[int, list[str]] = defaultdict(list)
            for alias in aliases:
                normalized = _norm(alias)
                if normalized in columns:
                    index = columns[normalized]
                    matches[index].append(alias)
            if len(matches) > 1:
                names = ",".join(
                    f"{'/'.join(names)}@{index + 1}" for index, names in sorted(matches.items())
                )
                raise ValueError(f"AMBIGUOUS_FIELD_COLUMNS:{path.name}:{field}:{names}")
            return next(iter(matches)) if matches else None

        for field, aliases in FIELD_ALIASES.items():
            index = resolve_alias(field, aliases)
            if index is not None:
                resolved[field] = index
        for field, aliases in FACT_FIELD_ALIASES.items():
            index = resolve_alias(field, aliases)
            if index is not None:
                resolved[field] = index
        for field, aliases in BUSINESS_TAG_FIELD_ALIASES.items():
            index = resolve_alias(f"business_tag_{field}", aliases)
            if index is not None:
                resolved[f"business_tag_{field}"] = index
        for field, aliases in (("remarks", REMARK_FIELD_ALIASES), ("raw_tags", RAW_TAG_FIELD_ALIASES)):
            index = resolve_alias(field, aliases)
            if index is not None:
                resolved[field] = index
        missing = sorted(set(FIELD_ALIASES) - set(resolved))
        if missing:
            raise ValueError(f"REQUIRED_COLUMNS_MISSING:{path.name}:{','.join(missing)}")
        output: dict[str, dict[str, Any]] = {}
        for line_number, row in enumerate(rows, start=2):
            sku = _cell_text(row[resolved["sku"]]) if resolved["sku"] < len(row) else ""
            if not sku:
                continue
            if sku in output:
                raise ValueError(f"DUPLICATE_SKU:{path.name}:{sku}:row={line_number}")
            output[sku] = {}
            for field, index in resolved.items():
                if field == "sku":
                    continue
                value = row[index] if index < len(row) else ""
                output[sku][field] = value if field.startswith("business_tag_") else _cell_text(value)
        return headers, output
    finally:
        workbook.close()


def _source_record(row: dict[str, str]) -> dict[str, str]:
    return {
        "name": row["name"], "cat1": row["cat1"], "cat2": row["cat2"],
        "spec": row["spec"], "description": row["description"], "details": row["details"],
    }


def _compare_fact_value(field: str, source: dict[str, Any], target: dict[str, Any]) -> str | None:
    """Return a concrete mismatch code without guessing malformed facts."""
    source_value = source.get(field)
    target_value = target.get(field)
    source_text = str(source_value or "").strip()
    target_text = str(target_value or "").strip()
    if field in {"current_price", "product_url"} and (not source_text or not target_text):
        return "FACT_VALUE_MISSING"
    if field in {"product_url", "image_url"}:
        for value in (source_text, target_text):
            if value:
                parsed = urlparse(value)
                if parsed.scheme.casefold() not in {"http", "https"} or not parsed.netloc:
                    return "FACT_VALUE_INVALID"
    if field in {"current_price", "original_price"}:
        source_price = parse_price(source_value)
        target_price = parse_price(target_value)
        if (source_text and source_price is None) or (target_text and target_price is None):
            return "FACT_VALUE_UNPARSEABLE"
        if field == "original_price":
            source_current = parse_price(source.get("current_price"))
            target_current = parse_price(target.get("current_price"))
            if source_current is None or target_current is None:
                return "FACT_VALUE_UNPARSEABLE"
            source_price = source_price if source_price is not None and source_price > source_current else None
            target_price = target_price if target_price is not None and target_price > target_current else None
        if source_price is None and target_price is None:
            return None
        if source_price is None or target_price is None:
            return "FACT_MISMATCH"
        if field == "current_price" and (source_price <= 0 or target_price <= 0):
            return "FACT_VALUE_INVALID"
        return None if abs(source_price - target_price) <= 0.005 else "FACT_MISMATCH"
    return None if source_text == target_text else "FACT_MISMATCH"


def _official_labels(row: dict[str, Any], *, chinese: bool) -> list[str] | None:
    """Extract source official labels or their explicitly retained export evidence.

    ``None`` means the workbook has no usable label source column, which is
    different from an explicitly empty label set.
    """
    raw_tags = str(row.get("raw_tags") or "").strip()
    remarks = str(row.get("remarks") or "").strip()
    if raw_tags:
        payload = raw_tags
    else:
        marker = "官网官方标签：" if chinese else "Etiquetas oficiales:"
        match = re.search(rf"(?:^|[；;])\s*{re.escape(marker)}\s*([^；;]*)", remarks, re.IGNORECASE)
        if match is None:
            if "remarks" in row or "raw_tags" in row:
                return []
            return None
        payload = match.group(1).strip()
    labels = []
    for item in payload.split("|"):
        label = item.strip()
        if chinese:
            # The Chinese export retains the Spanish official label verbatim,
            # followed by a Chinese explanation in full-width parentheses.
            label = re.sub(r"\s*（[^（）]*）\s*$", "", label).strip()
        if label and not re.fullmatch(r"[−–-]?\s*\d+(?:[.,]\d+)?\s*%", label):
            labels.append(label)
    return labels


def _parse_bool_flag(value: object) -> bool | None:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in {0, 1}:
        return bool(value)
    normalized = _norm(value)
    if normalized in {"1", "true", "yes", "si", "activo", "active"}:
        return True
    if normalized in {"0", "false", "no", "inactivo", "inactive"}:
        return False
    return None


def _structured_business_tag_flags(
    row: dict[str, Any],
) -> tuple[dict[str, bool] | None, tuple[str, ...], bool]:
    """Return structured flags, missing/invalid fields, and whether any were supplied."""
    fields = tuple(f"business_tag_{key}" for key in BUSINESS_TAG_FIELD_ALIASES)
    supplied = any(field in row for field in fields)
    if not supplied:
        return None, (), False
    invalid = []
    flags = {}
    for key, field in zip(BUSINESS_TAG_FIELD_ALIASES, fields, strict=True):
        if field not in row:
            invalid.append(field)
            continue
        parsed = _parse_bool_flag(row[field])
        if parsed is None:
            invalid.append(field)
            continue
        flags[key] = parsed
    return (flags if not invalid else None), tuple(invalid), True


def _business_tag_flags_from_remarks(row: dict[str, Any], *, chinese: bool) -> dict[str, bool] | None:
    """Read legacy display tags before the official-label segment."""
    if "remarks" not in row:
        return None
    remarks = str(row.get("remarks") or "").strip()
    marker = "官网官方标签：" if chinese else "Etiquetas oficiales:"
    prefix = re.split(re.escape(marker), remarks, maxsplit=1, flags=re.IGNORECASE)[0]
    values = {_norm(part) for part in re.split(r"[；;]", prefix) if part.strip()}
    source_terms = ("新品", "促销", "可持续") if chinese else ("nuevo", "promocion", "sostenible")
    return dict(zip(("new", "promotion", "sustainable"), (term in values for term in source_terms), strict=True))


def _business_tag_flags(row: dict[str, Any], *, chinese: bool) -> tuple[dict[str, bool] | None, str, tuple[str, ...]]:
    """Prefer structured source-of-truth flags; retain remarks fallback for old workbooks."""
    # The Chinese workbook is the rendered deliverable, so check what it
    # actually displays in Remarks even if an auxiliary boolean column exists.
    if chinese and "remarks" in row:
        return _business_tag_flags_from_remarks(row, chinese=True), "REMARKS_EXPORT", ()
    structured, invalid, supplied = _structured_business_tag_flags(row)
    if supplied:
        return structured, "STRUCTURED_FIELDS", invalid
    return _business_tag_flags_from_remarks(row, chinese=chinese), "REMARKS_FALLBACK", ()


def _audit_business_tags(
    shared_skus: set[str], source_rows: dict[str, dict[str, Any]],
    target_rows: dict[str, dict[str, Any]], issues: list[dict[str, Any]],
) -> dict[str, Any]:
    """Compare generic business badges separately from preserved official labels."""
    checked = 0
    unavailable = 0
    source_basis: Counter[str] = Counter()
    target_basis: Counter[str] = Counter()
    source_names = {"new": "Nuevo", "promotion": "Promoción", "sustainable": "Sostenible"}
    target_names = {"new": "新品", "promotion": "促销", "sustainable": "可持续"}
    for sku in sorted(shared_skus):
        source_flags, source_kind, source_invalid = _business_tag_flags(source_rows[sku], chinese=False)
        target_flags, target_kind, target_invalid = _business_tag_flags(target_rows[sku], chinese=True)
        source_basis[source_kind] += 1
        target_basis[target_kind] += 1
        if source_invalid:
            issues.append({
                "sku": sku, "field": "business_tags", "code": "BUSINESS_TAG_SOURCE_INVALID",
                "evidence": json.dumps({"invalid_fields": source_invalid}, ensure_ascii=False),
            })
            unavailable += 1
            continue
        if target_invalid:
            issues.append({
                "sku": sku, "field": "business_tags", "code": "BUSINESS_TAG_TARGET_INVALID",
                "evidence": json.dumps({"invalid_target_fields": target_invalid}, ensure_ascii=False),
            })
            unavailable += 1
            continue
        if source_flags is None or target_flags is None:
            unavailable += 1
            continue
        checked += 1
        for key in source_names:
            if source_flags[key] == target_flags[key]:
                continue
            code = "BUSINESS_TAG_DROPPED" if source_flags[key] else "BUSINESS_TAG_HALLUCINATED"
            issues.append({
                "sku": sku, "field": "remarks", "code": code,
                "source": source_names[key] if source_flags[key] else "",
                "candidate": target_names[key] if target_flags[key] else "",
                "tag_type": key,
            })
    return {
        "status": "CHECKED" if checked else "NOT_PROVIDED",
        "checked_skus": checked,
        "unavailable_skus": unavailable,
        "source_evidence": dict(source_basis),
        "target_evidence": dict(target_basis),
    }


def _audit_official_labels(
    shared_skus: set[str], source_rows: dict[str, dict[str, Any]],
    target_rows: dict[str, dict[str, Any]], issues: list[dict[str, Any]],
) -> dict[str, Any]:
    """Check that source official labels survive in the Chinese export remarks."""
    checked = 0
    unavailable = 0
    for sku in sorted(shared_skus):
        source_labels = _official_labels(source_rows[sku], chinese=False)
        target_labels = _official_labels(target_rows[sku], chinese=True)
        if source_labels is None or target_labels is None:
            unavailable += 1
            continue
        checked += 1
        source_norm = [_norm(value) for value in source_labels]
        target_norm = [_norm(value) for value in target_labels]
        if source_norm and not target_norm:
            code = "OFFICIAL_LABEL_DROPPED"
        elif target_norm and not source_norm:
            code = "OFFICIAL_LABEL_HALLUCINATED"
        elif source_norm != target_norm:
            code = "OFFICIAL_LABEL_MISMATCH"
        else:
            continue
        issues.append({
            "sku": sku, "field": "remarks", "code": code,
            "source": source_labels, "candidate": target_labels,
        })
    return {
        "status": "CHECKED" if checked else "NOT_PROVIDED",
        "checked_skus": checked,
        "unavailable_skus": unavailable,
    }


def _audit_term_dictionary(
    source_rows: dict[str, dict[str, Any]], target_rows: dict[str, dict[str, Any]],
    path: Path, issues: list[dict[str, Any]],
) -> dict[str, Any]:
    """Measure approved glossary coverage and enforce only explicit prohibitions."""
    required = {"term_es", "term_zh", "term_type", "forbidden_zh", "review_status"}
    invalid_approved_rows = 0
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if not required.issubset(set(reader.fieldnames or ())):
                issues.append({"sku": "", "field": "term_dictionary", "code": "TERM_DICTIONARY_INVALID"})
                return {"status": "INVALID_SCHEMA", "sha256": _sha256_file(path), "approved_terms": 0}
            approved_rows = [
                row for row in reader
                if _norm(row.get("review_status")).upper() in APPROVED_TERM_STATUSES
            ]
            terms = []
            for row in approved_rows:
                if not all(_norm(row.get(field)) for field in ("term_es", "term_zh", "term_type")):
                    invalid_approved_rows += 1
                    continue
                terms.append(row)
    except FileNotFoundError:
        issues.append({"sku": "", "field": "term_dictionary", "code": "TERM_DICTIONARY_MISSING"})
        return {"status": "MISSING", "sha256": None, "approved_terms": 0}
    except (OSError, UnicodeError, csv.Error):
        issues.append({"sku": "", "field": "term_dictionary", "code": "TERM_DICTIONARY_INVALID"})
        return {"status": "READ_ERROR", "sha256": None, "approved_terms": 0}

    if invalid_approved_rows:
        issues.append({
            "sku": "", "field": "term_dictionary", "code": "TERM_DICTIONARY_INVALID_ROW",
            "row_count": invalid_approved_rows,
        })
    if not terms:
        issues.append({"sku": "", "field": "term_dictionary", "code": "TERM_DICTIONARY_NO_APPROVED_ROWS"})
    grouped_terms: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in terms:
        grouped_terms[_norm(row["term_es"])].append(row)
    unique_terms: list[dict[str, str]] = []
    ambiguous_terms = 0
    for normalized_term, rows in sorted(grouped_terms.items()):
        signatures = {
            (
                _norm(row.get("term_zh")), _norm(row.get("term_type")),
                tuple(sorted(_norm(value) for value in str(row.get("forbidden_zh") or "").split("|") if _norm(value))),
            )
            for row in rows
        }
        if len(signatures) > 1:
            ambiguous_terms += 1
            issues.append({
                "sku": "", "field": "term_dictionary", "code": "TERM_DICTIONARY_AMBIGUOUS",
                "source_term": normalized_term,
                "approved_targets": sorted({str(row.get("term_zh") or "").strip() for row in rows}),
            })
            continue
        unique_terms.append(rows[0])
    terms = unique_terms
    review_policy = load_approved_term_review_policy(APPROVED_TERM_REVIEW_POLICY_PATH)
    fields = ("name", "spec", "description", "details")
    unit_types = {"unit", "quantity"}
    counts: dict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    examples: dict[tuple[str, str], set[str]] = defaultdict(set)
    review_only_suggestions = 0
    for row in terms:
        source_term = _norm(row["term_es"])
        term_type = _norm(row.get("term_type")).casefold()
        allowed_fields = ("spec", "description", "details") if term_type in unit_types else fields
        source_pattern = re.compile(rf"(?<![a-z0-9]){re.escape(source_term)}(?![a-z0-9])")
        canonical = _norm(row["term_zh"])
        for sku in sorted(set(source_rows) & set(target_rows)):
            for field in allowed_fields:
                if not source_pattern.search(_norm(source_rows[sku].get(field))):
                    continue
                candidate = _norm(target_rows[sku].get(field))
                counts[(row["term_es"], field)]["source_occurrences"] += 1
                if canonical and canonical in candidate:
                    counts[(row["term_es"], field)]["canonical_target_present"] += 1
                if len(examples[(row["term_es"], field)]) < 10:
                    examples[(row["term_es"], field)].add(sku)
                findings = inspect_approved_term_candidate(
                    field, source_rows[sku].get(field), target_rows[sku].get(field), (row,),
                    review_policy=review_policy,
                )
                for finding in findings:
                    issues.append({"sku": sku, **finding})
                    if finding.get("code") == "APPROVED_TERM_CANONICAL_ABSENT_REVIEW":
                        review_only_suggestions += 1
    for sku in sorted(set(source_rows) & set(target_rows)):
        source_fields = {field: source_rows[sku].get(field, "") for field in fields}
        target_fields = {field: target_rows[sku].get(field, "") for field in fields}
        for finding in inspect_approved_term_cross_field_movement(
            source_fields, target_fields, terms, review_policy=review_policy,
        ):
            issues.append({"sku": sku, **finding})
            if finding.get("review_only"):
                review_only_suggestions += 1
    coverage = [
        {
            "source_term": term, "field": field,
            "source_occurrences": count["source_occurrences"],
            "canonical_target_present": count["canonical_target_present"],
            "canonical_presence_ratio": count["canonical_target_present"] / count["source_occurrences"],
            "sku_examples": sorted(examples[(term, field)]),
        }
        for (term, field), count in sorted(counts.items())
    ]
    return {
        "status": "LOADED", "sha256": _sha256_file(path), "approved_terms": len(terms),
        "ambiguous_normalized_terms": ambiguous_terms,
        "source_term_field_pairs_observed": len(coverage),
        "review_only_suggestions": review_only_suggestions,
        "canonical_presence_is_not_semantic_acceptance": True,
        "review_policy_id": review_policy.get("policy_id"),
        "review_policy_sha256": _sha256_file(APPROVED_TERM_REVIEW_POLICY_PATH),
        "coverage": coverage[:500],
    }


def _term_rule_review_rows(issues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Aggregate approved-glossary findings to term/field review units.

    Per-SKU evidence remains in QA_LOG.csv. This compact queue helps a reviewer
    decide whether to accept a synonym, correct a rule, or mark a false signal;
    observed translations never become approved automatically.
    """
    review_codes = {
        "APPROVED_TERM_CANONICAL_ABSENT_REVIEW",
        "APPROVED_TERM_CROSS_FIELD_REVIEW",
        "APPROVED_TERM_UNSUPPORTED_TARGET_REVIEW",
        "TERM_FORBIDDEN_TRANSLATION",
    }
    grouped: dict[tuple[str, str, str, str, str, str, str], dict[str, Any]] = {}
    for issue in issues:
        code = str(issue.get("code") or "")
        if code not in review_codes:
            continue
        key = (
            code, _norm(issue.get("source_term")),
            _norm(issue.get("term_type")), _norm(issue.get("source_field")),
            _norm(issue.get("field")), _norm(issue.get("approved_translation")),
            _norm(issue.get("forbidden_candidate")),
        )
        group = grouped.setdefault(key, {
            "review_kind": code,
            "source_term": str(issue.get("source_term") or "").strip(),
            "term_type": str(issue.get("term_type") or "").strip(),
            "source_field": str(issue.get("source_field") or "").strip(),
            "target_field": str(issue.get("field") or "").strip(),
            "approved_translation": str(issue.get("approved_translation") or "").strip(),
            "forbidden_candidate": str(issue.get("forbidden_candidate") or "").strip(),
            "occurrences": 0,
            "skus": set(),
            "candidates": Counter(),
            "examples": {},
            "review_only": bool(issue.get("review_only", code != "TERM_FORBIDDEN_TRANSLATION")),
        })
        group["occurrences"] += 1
        sku = str(issue.get("sku") or "").strip()
        if sku:
            group["skus"].add(sku)
        candidate = str(issue.get("candidate") or "").strip()
        if candidate:
            group["candidates"][candidate] += 1
        if sku and sku not in group["examples"]:
            group["examples"][sku] = {
                "sku": sku,
                "source_value": str(issue.get("source_value") or "")[:240],
                "candidate": candidate[:240],
            }

    rows = []
    for _, group in sorted(grouped.items()):
        candidates = group["candidates"]
        candidate_examples = [
            {"candidate": candidate, "occurrences": count}
            for candidate, count in sorted(candidates.items(), key=lambda item: (-item[1], item[0]))[:10]
        ]
        example_rows = [group["examples"][sku] for sku in sorted(group["examples"])[:10]]
        rows.append({
            "review_kind": group["review_kind"],
            "source_term": group["source_term"],
            "term_type": group["term_type"],
            "source_field": group["source_field"],
            "target_field": group["target_field"],
            "approved_translation": group["approved_translation"],
            "forbidden_candidate": group["forbidden_candidate"],
            "occurrences": group["occurrences"],
            "sku_count": len(group["skus"]),
            "sku_examples": json.dumps(sorted(group["skus"])[:20], ensure_ascii=False),
            "candidate_variant_count": len(candidates),
            "candidate_examples": json.dumps(candidate_examples, ensure_ascii=False),
            "candidate_examples_truncated": len(candidates) > len(candidate_examples),
            "evidence_examples": json.dumps(example_rows, ensure_ascii=False),
            "review_only": group["review_only"],
            "candidate_is_approved": False,
        })
    return rows


def _load_approved_category_map(path: Path) -> dict[str, Any]:
    """Load only explicitly approved category dictionary rows; never infer from frequency."""
    cat1_values: dict[str, set[str]] = defaultdict(set)
    cat2_values: dict[tuple[str, str], set[str]] = defaultdict(set)
    row_counts = Counter()
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            required = {"cat1_es", "cat2_es", "cat1_zh", "cat2_zh", "review_status"}
            if not required.issubset(set(reader.fieldnames or ())):
                return {"status": "INVALID_SCHEMA", "cat1_values": {}, "cat2_values": {},
                        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "approved_row_counts": dict(row_counts)}
            for row in reader:
                status = _norm(row.get("review_status")).upper()
                if status not in APPROVED_CATEGORY_STATUSES and status != "CAT1_CONFIRMED":
                    continue
                cat1_es, cat1_zh = _norm(row.get("cat1_es")), str(row.get("cat1_zh") or "").strip()
                cat2_es, cat2_zh = _norm(row.get("cat2_es")), str(row.get("cat2_zh") or "").strip()
                if cat1_es and cat1_zh and (status in APPROVED_CATEGORY_STATUSES or status == "CAT1_CONFIRMED"):
                    cat1_values[cat1_es].add(cat1_zh)
                    row_counts["cat1"] += 1
                if cat1_es and cat2_es and cat2_zh and status in APPROVED_CATEGORY_STATUSES:
                    cat2_values[(cat1_es, cat2_es)].add(cat2_zh)
                    row_counts["cat2"] += 1
    except FileNotFoundError:
        return {"status": "MISSING", "cat1_values": {}, "cat2_values": {},
                "sha256": None, "approved_row_counts": dict(row_counts)}
    except (OSError, UnicodeError, csv.Error):
        return {"status": "READ_ERROR", "cat1_values": {}, "cat2_values": {},
                "sha256": None, "approved_row_counts": dict(row_counts)}
    return {
        "status": "LOADED",
        "cat1_values": dict(cat1_values),
        "cat2_values": dict(cat2_values),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "approved_row_counts": dict(row_counts),
    }


def _load_confirmed_brand_phrases(path: Path) -> dict[str, Any]:
    """Load confirmed brands for source-token disambiguation only.

    The values never authorize brands in Chinese display titles; that policy
    remains NO_BRAND_TITLE_V1. A phrase can suppress a model-token finding
    only when it occurs verbatim in the Spanish source title.
    """
    accepted_statuses = {"APPROVED", "CONFIRMED", "HUMAN_APPROVED", "HUMAN_REVIEWED", "LOCKED"}
    phrases: set[str] = set()
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            required = {"canonical_name", "aliases_es", "review_status"}
            if not required.issubset(set(reader.fieldnames or ())):
                return {
                    "status": "INVALID_SCHEMA", "phrases": (),
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                }
            for row in reader:
                if _norm(row.get("review_status")).upper() not in accepted_statuses:
                    continue
                canonical = str(row.get("canonical_name") or "").strip()
                aliases = [value.strip() for value in str(row.get("aliases_es") or "").split("|") if value.strip()]
                phrases.update(value for value in (canonical, *aliases) if value)
    except FileNotFoundError:
        return {"status": "MISSING", "phrases": (), "sha256": None}
    except (OSError, UnicodeError, csv.Error):
        return {"status": "READ_ERROR", "phrases": (), "sha256": None}
    return {
        "status": "LOADED",
        "phrases": tuple(sorted(phrases, key=lambda value: (value.casefold(), value))),
        "phrase_count": len(phrases),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def _audit_category_rows(
    shared_skus: set[str], source_rows: dict[str, dict[str, Any]],
    target_rows: dict[str, dict[str, Any]], category_map: dict[str, Any],
    issues: list[dict[str, Any]],
) -> dict[str, Any]:
    if category_map["status"] != "LOADED":
        code = "CATEGORY_DICTIONARY_MISSING" if category_map["status"] == "MISSING" else "CATEGORY_DICTIONARY_INVALID"
        issues.append({"sku": "", "field": "category_dictionary", "code": code,
                       "evidence": category_map["status"]})
        return {
            "status": category_map["status"], "sha256": category_map.get("sha256"),
            "approved_cat1_keys": 0, "approved_cat2_pairs": 0,
            "checked_cat1_rows": 0, "checked_cat2_rows": 0,
        }

    cat1_values = category_map["cat1_values"]
    cat2_values = category_map["cat2_values"]
    checked = Counter()
    for sku in sorted(shared_skus):
        source, target = source_rows[sku], target_rows[sku]
        source_cat1 = _norm(source.get("cat1"))
        source_cat2 = _norm(source.get("cat2"))
        for field, key, mappings in (
            ("cat1", source_cat1, cat1_values),
            ("cat2", (source_cat1, source_cat2), cat2_values),
        ):
            if field == "cat1" and not source_cat1:
                continue
            if field == "cat2" and (not source_cat1 or not source_cat2):
                continue
            expected_values = mappings.get(key, set())
            actual = str(target.get(field) or "").strip()
            checked[field] += 1
            if not expected_values:
                issues.append({
                    "sku": sku, "field": field, "code": "CATEGORY_MAPPING_UNAVAILABLE",
                    "source": source.get(field, ""), "candidate": actual,
                    "source_cat1": source.get("cat1", ""),
                    "candidate_cat1": target.get("cat1", ""),
                })
            elif len(expected_values) > 1:
                issues.append({
                    "sku": sku, "field": field, "code": "CATEGORY_MAPPING_AMBIGUOUS",
                    "source": source.get(field, ""), "candidate": actual,
                    "approved_values": sorted(expected_values),
                    "source_cat1": source.get("cat1", ""),
                    "candidate_cat1": target.get("cat1", ""),
                })
            else:
                expected = next(iter(expected_values))
                if _norm(actual) != _norm(expected):
                    issues.append({
                        "sku": sku, "field": field, "code": "CATEGORY_MAPPING_MISMATCH",
                        "source": source.get(field, ""), "expected": expected, "candidate": actual,
                        "source_cat1": source.get("cat1", ""),
                        "candidate_cat1": target.get("cat1", ""),
                    })
    return {
        "status": "LOADED",
        "sha256": category_map.get("sha256"),
        "approved_cat1_keys": len(cat1_values),
        "approved_cat2_pairs": len(cat2_values),
        "checked_cat1_rows": checked["cat1"],
        "checked_cat2_rows": checked["cat2"],
        "ambiguous_cat1_keys": sum(len(values) > 1 for values in cat1_values.values()),
        "ambiguous_cat2_pairs": sum(len(values) > 1 for values in cat2_values.values()),
    }


def _detail_rule_coverage(
    source_rows: dict[str, dict[str, str]],
    target_rows: dict[str, dict[str, str]],
    rules: dict[str, Any],
) -> dict[str, Any]:
    """Report where the deterministic detail lexicon covers observed source terms.

    Coverage gaps are discovery signals, not translation failures: this report
    never edits a workbook and does not infer a translation for an unruled term.
    """
    key_occurrences: Counter[str] = Counter()
    key_occurrences_with_rule: Counter[str] = Counter()
    key_examples: dict[str, set[str]] = defaultdict(set)
    key_uncovered_examples: dict[str, set[str]] = defaultdict(set)
    key_targets: dict[str, Counter[str]] = defaultdict(Counter)
    key_target_skus: dict[tuple[str, str], set[str]] = defaultdict(set)
    value_occurrences: Counter[tuple[str, str]] = Counter()
    value_occurrences_with_rule: Counter[tuple[str, str]] = Counter()
    value_targets: dict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    value_target_pairs: dict[tuple[str, str], Counter[tuple[str, str]]] = defaultdict(Counter)
    value_target_pair_skus: dict[tuple[tuple[str, str], str, str], set[str]] = defaultdict(set)
    value_target_skus: dict[tuple[tuple[str, str], str], set[str]] = defaultdict(set)
    value_examples: dict[tuple[str, str], set[str]] = defaultdict(set)
    value_uncovered_examples: dict[tuple[str, str], set[str]] = defaultdict(set)
    key_rules = {_norm(key) for key in (rules.get("key_translations") or {})}
    contextual_key_rules = list(rules.get("contextual_key_translations") or ())
    numeric_rules = {
        _norm(item.get("source_key")): item
        for item in rules.get("numeric_suffix_rules", ()) if item.get("source_key")
    }
    value_rules = list(rules.get("value_translations") or ())
    shared = set(source_rows) & set(target_rows)
    skipped_pair_alignment = 0
    alignment_review_rows: list[dict[str, Any]] = []

    for sku in shared:
        source_pairs = parse_details(source_rows[sku]["details"])
        target_pairs = parse_details(target_rows[sku]["details"])
        if len(source_pairs) != len(target_pairs):
            skipped_pair_alignment += 1
            alignment_review_rows.append({
                "review_kind": "PAIR_ALIGNMENT",
                "source_key_normalized": "",
                "source_value_normalized": "",
                "occurrences": 1,
                "rule_covered_occurrences": 0,
                "uncovered_occurrences": 1,
                "sku_count": 1,
                "sku_examples": json.dumps([sku], ensure_ascii=False),
                "observed_target_candidates": "[]",
                "single_observed_candidate": "",
                "review_reasons": "SOURCE_TARGET_PAIR_COUNT_MISMATCH",
                "candidate_is_approved": False,
                "sku": sku,
                "source_details": source_rows[sku]["details"],
                "target_details": target_rows[sku]["details"],
                "source_pair_count": len(source_pairs),
                "target_pair_count": len(target_pairs),
            })
            continue
        context = {"name_es": source_rows[sku].get("name", "")}
        for source_pair, target_pair in zip(source_pairs, target_pairs, strict=True):
            key = _norm(source_pair.key_es)
            target_key = str(target_pair.key_es or "").strip()
            key_occurrences[key] += 1
            key_examples[key].add(sku)
            key_targets[key][target_key] += 1
            key_target_skus[(key, _norm(target_key))].add(sku)
            key_covered = key in key_rules or key in numeric_rules or any(
                _norm(rule.get("source_key")) == key
                and detail_rule_context_matches(rule, context)
                for rule in contextual_key_rules
            )
            if key_covered:
                key_occurrences_with_rule[key] += 1
            else:
                key_uncovered_examples[key].add(sku)
            source_value = _norm(source_pair.value_es)
            pair = (key, source_value)
            value_occurrences[pair] += 1
            target_value = str(target_pair.value_es or "").strip()
            value_targets[pair][target_value] += 1
            value_target_pairs[pair][(target_key, target_value)] += 1
            value_target_pair_skus[(pair, _norm(target_key), _norm(target_value))].add(sku)
            value_target_skus[(pair, _norm(target_value))].add(sku)
            value_examples[pair].add(sku)
            value_covered = any(
                detail_value_rule_matches_source(
                    rule, source_pair.key_es, source_pair.value_es, context=context,
                )
                for rule in value_rules
            )
            numeric_rule = numeric_rules.get(key)
            if numeric_rule and re.fullmatch(r"\d+(?:[.,]\d+)?", source_pair.value_es.strip()):
                value_covered = True
            if value_covered:
                value_occurrences_with_rule[pair] += 1
            else:
                value_uncovered_examples[pair].add(sku)

    covered_key_occurrences = sum(key_occurrences_with_rule.values())
    covered_value_occurrences = sum(value_occurrences_with_rule.values())
    covered_value_pairs = {
        pair for pair, count in value_occurrences.items()
        if value_occurrences_with_rule[pair] == count
    }

    unmapped_keys = [
        {
            "source_key": key,
            "occurrences": count,
            "covered_occurrences": key_occurrences_with_rule[key],
            "uncovered_occurrences": count - key_occurrences_with_rule[key],
            "sku_count": len(key_examples[key]),
            "sku_examples": sorted(key_examples[key])[:12],
            "uncovered_sku_examples": sorted(key_uncovered_examples[key])[:12],
            "observed_chinese_keys": key_targets[key].most_common(8),
        }
        for key, count in key_occurrences.most_common()
        if key_occurrences_with_rule[key] < count
    ]
    unmapped_values = [
        {
            "source_key": key,
            "source_value": value,
            "occurrences": count,
            "covered_occurrences": value_occurrences_with_rule[(key, value)],
            "uncovered_occurrences": count - value_occurrences_with_rule[(key, value)],
            "sku_examples": sorted(value_examples[(key, value)])[:12],
            "uncovered_sku_examples": sorted(value_uncovered_examples[(key, value)])[:12],
            "observed_chinese_values": value_targets[(key, value)].most_common(8),
        }
        for (key, value), count in value_occurrences.most_common()
        if (key, value) not in covered_value_pairs
    ]
    translation_variants = []
    for (key, value), count in value_occurrences.most_common():
        target_groups: dict[str, list[str]] = defaultdict(list)
        for target_value in value_targets[(key, value)]:
            target_groups[_norm(target_value)].append(target_value)
        if len(target_groups) <= 1:
            continue
        variant_rows = []
        for normalized_target, forms in sorted(target_groups.items()):
            variant_rows.append({
                "normalized_target": normalized_target,
                "observed_forms": sorted(forms),
                "occurrences": sum(value_targets[(key, value)][form] for form in forms),
                "sku_examples": sorted(value_target_skus[((key, value), normalized_target)])[:12],
            })
        covered = value_occurrences_with_rule[(key, value)]
        translation_variants.append({
            "source_key": key,
            "source_value": value,
            "occurrences": count,
            "rule_covered_occurrences": covered,
            "rule_coverage_status": (
                "FULL_RULE_COVERAGE" if covered == count else
                "PARTIAL_RULE_COVERAGE" if covered else "NO_RULE_COVERAGE"
            ),
            "translation_variant_count": len(target_groups),
            "observed_chinese_values": value_targets[(key, value)].most_common(8),
            "variants": variant_rows,
            "review_only": True,
            "variation_is_not_proof_of_error": True,
        })

    value_variants_by_pair = {
        (item["source_key"], item["source_value"]): item
        for item in translation_variants
    }
    review_queue_rows: list[dict[str, Any]] = []
    for key, count in key_occurrences.most_common():
        target_groups: dict[str, list[str]] = defaultdict(list)
        for target_key in key_targets[key]:
            target_groups[_norm(target_key)].append(target_key)
        uncovered = count - key_occurrences_with_rule[key]
        has_variants = len(target_groups) > 1
        if not uncovered and not has_variants:
            continue
        candidates = [
            {
                "target_key_forms": sorted(forms),
                "occurrences": sum(key_targets[key][form] for form in forms),
                "sku_examples": sorted(key_target_skus[(key, normalized)])[:20],
            }
            for normalized, forms in sorted(target_groups.items())
        ]
        consensus = (
            candidates[0]["target_key_forms"][0]
            if len(candidates) == 1 and len(candidates[0]["target_key_forms"]) == 1 else ""
        )
        review_queue_rows.append({
            "review_kind": "SOURCE_KEY", "source_key_normalized": key,
            "source_value_normalized": "", "occurrences": count,
            "rule_covered_occurrences": key_occurrences_with_rule[key],
            "uncovered_occurrences": uncovered, "sku_count": len(key_examples[key]),
            "sku_examples": json.dumps(sorted(key_examples[key])[:20], ensure_ascii=False),
            "observed_target_candidates": json.dumps(candidates, ensure_ascii=False, sort_keys=True),
            "single_observed_candidate": consensus,
            "review_reasons": "|".join(filter(None, (
                "KEY_RULE_UNCOVERED" if uncovered else "",
                "TRANSLATION_VARIANTS" if has_variants else "",
            ))),
            "candidate_is_approved": False,
        })

    for (key, value), count in value_occurrences.most_common():
        variant = value_variants_by_pair.get((key, value))
        uncovered = count - value_occurrences_with_rule[(key, value)]
        pair_groups: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
        for target_key, target_value in value_target_pairs[(key, value)]:
            pair_groups[(_norm(target_key), _norm(target_value))].append((target_key, target_value))
        has_target_pair_variants = len(pair_groups) > 1
        if not uncovered and variant is None and not has_target_pair_variants:
            continue
        observed_candidates = []
        for normalized_pair, forms in sorted(pair_groups.items()):
            observed_candidates.append({
                "target_key_value_forms": [list(form) for form in sorted(forms)],
                "occurrences": sum(value_target_pairs[(key, value)][form] for form in forms),
                "sku_examples": sorted(value_target_pair_skus[(key, value, *normalized_pair)])[:20],
            })
        candidate = ""
        if len(observed_candidates) == 1 and len(observed_candidates[0]["target_key_value_forms"]) == 1:
            candidate = json.dumps(observed_candidates[0]["target_key_value_forms"][0], ensure_ascii=False)
        review_queue_rows.append({
            "review_kind": "SOURCE_VALUE", "source_key_normalized": key,
            "source_value_normalized": value, "occurrences": count,
            "rule_covered_occurrences": value_occurrences_with_rule[(key, value)],
            "uncovered_occurrences": uncovered,
            "sku_count": len(value_examples[(key, value)]),
            "sku_examples": json.dumps(sorted(value_examples[(key, value)])[:20], ensure_ascii=False),
            "observed_target_candidates": json.dumps(observed_candidates, ensure_ascii=False, sort_keys=True),
            "single_observed_candidate": candidate,
            "review_reasons": "|".join(filter(None, (
                "VALUE_RULE_UNCOVERED" if uncovered else "",
                "TRANSLATION_VARIANTS" if variant or has_target_pair_variants else "",
            ))),
            "candidate_is_approved": False,
        })

    review_queue_rows.extend(alignment_review_rows)
    review_queue_rows.sort(key=lambda item: (
        item["review_kind"], str(item.get("sku") or ""),
        item["source_key_normalized"], item["source_value_normalized"],
    ))
    key_total = sum(key_occurrences.values())
    value_total = len(value_occurrences)
    return {
        "review_only": True,
        "key_occurrences": key_total,
        "key_occurrences_with_rule": covered_key_occurrences,
        "key_occurrence_coverage": (
            covered_key_occurrences / key_total
            if key_total else 1.0
        ),
        "unmapped_source_key_count": len(unmapped_keys),
        "unmapped_source_keys": unmapped_keys[:200],
        "unique_source_key_value_pairs": value_total,
        "source_key_value_pairs_with_rule": len(covered_value_pairs),
        "source_key_value_pair_coverage": len(covered_value_pairs) / value_total if value_total else 1.0,
        "source_key_value_occurrences": sum(value_occurrences.values()),
        "source_key_value_occurrences_with_rule": covered_value_occurrences,
        "source_key_value_occurrence_coverage": (
            covered_value_occurrences / sum(value_occurrences.values())
            if value_occurrences else 1.0
        ),
        "unmapped_source_key_value_pair_count": len(unmapped_values),
        "unmapped_source_key_value_pairs": unmapped_values[:300],
        "source_key_value_pairs_with_translation_variants": len(translation_variants),
        "source_key_value_translation_variants": translation_variants[:300],
        "review_queue_row_count": len(review_queue_rows),
        "review_queue_rows": review_queue_rows,
        "skus_skipped_for_pair_alignment": skipped_pair_alignment,
        "coverage_is_not_semantic_acceptance": True,
    }


def _category_rule_review_rows(issues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Collapse per-SKU category findings into scoped, non-approved rule reviews."""
    grouped: dict[tuple[str, str, str, str, str], dict[str, Any]] = {}
    review_codes = {
        "CATEGORY_MAPPING_UNAVAILABLE", "CATEGORY_MAPPING_AMBIGUOUS",
        "CATEGORY_MAPPING_MISMATCH",
    }
    for issue in issues:
        code = str(issue.get("code") or "")
        if code not in review_codes:
            continue
        key = (
            str(issue.get("field") or ""), _norm(issue.get("source_cat1")),
            _norm(issue.get("source")), _norm(issue.get("candidate_cat1")), code,
        )
        group = grouped.setdefault(key, {
            "review_kind": code,
            "field": key[0],
            "source_cat1": str(issue.get("source_cat1") or "").strip(),
            "source_category": str(issue.get("source") or "").strip(),
            "observed_target_cat1": str(issue.get("candidate_cat1") or "").strip(),
            "occurrences": 0,
            "skus": set(),
            "observed_target_values": Counter(),
            "approved_values": set(),
        })
        group["occurrences"] += 1
        sku = str(issue.get("sku") or "").strip()
        if sku:
            group["skus"].add(sku)
        candidate = str(issue.get("candidate") or "").strip()
        if candidate:
            group["observed_target_values"][candidate] += 1
        group["approved_values"].update(
            str(value).strip()
            for value in issue.get("approved_values", ())
            if str(value).strip()
        )
        expected = str(issue.get("expected") or "").strip()
        if expected:
            group["approved_values"].add(expected)

    rows = []
    for _, group in sorted(grouped.items()):
        observed = group["observed_target_values"]
        rows.append({
            "review_kind": group["review_kind"],
            "field": group["field"],
            "source_cat1": group["source_cat1"],
            "source_category": group["source_category"],
            "observed_target_cat1": group["observed_target_cat1"],
            "occurrences": group["occurrences"],
            "sku_count": len(group["skus"]),
            "sku_examples": json.dumps(sorted(group["skus"])[:20], ensure_ascii=False),
            "observed_target_values": json.dumps(
                [{"value": value, "occurrences": count} for value, count in sorted(observed.items())],
                ensure_ascii=False,
            ),
            "single_observed_target_value": next(iter(observed)) if len(observed) == 1 else "",
            "approved_values": json.dumps(sorted(group["approved_values"]), ensure_ascii=False),
            "candidate_is_approved": False,
        })
    return rows


def _load_qa_policy(path: Path) -> dict[str, Any]:
    policy = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(policy, dict) or policy.get("policy_id") != "ACTION_CHINESE_GOLD_QA_V2":
        raise ValueError("GOLD_QA_POLICY_ID_INVALID")
    if not isinstance(policy.get("layers"), list):
        raise ValueError("GOLD_QA_POLICY_SCHEMA_INVALID")
    if any(not isinstance(layer, dict) for layer in policy["layers"]):
        raise ValueError("GOLD_QA_POLICY_LAYER_SCHEMA_INVALID")
    layer_ids = {str(layer.get("id") or "") for layer in policy["layers"]}
    if layer_ids != {"L1", "L2", "L3", "L4", "L5", "L6"} or len(layer_ids) != len(policy["layers"]):
        raise ValueError("GOLD_QA_POLICY_LAYERS_INVALID")
    gold_claim_allowed = policy.get("gold_claim_allowed")
    if not isinstance(gold_claim_allowed, bool):
        raise ValueError("GOLD_QA_POLICY_GOLD_CLAIM_FLAG_INVALID")
    if any(layer.get("coverage") not in {"PARTIAL", "COMPLETE"} for layer in policy["layers"]):
        raise ValueError("GOLD_QA_POLICY_COVERAGE_INVALID")
    if gold_claim_allowed and any(layer.get("coverage") != "COMPLETE" for layer in policy["layers"]):
        raise ValueError("GOLD_QA_POLICY_GOLD_COVERAGE_INCOMPLETE")
    for layer in policy["layers"]:
        covered = layer.get("covered_checks")
        uncovered = layer.get("not_covered")
        if (
            not isinstance(covered, list)
            or any(not isinstance(item, str) or not item.strip() for item in covered)
            or not isinstance(uncovered, list)
            or any(not isinstance(item, str) or not item.strip() for item in uncovered)
        ):
            raise ValueError(f"GOLD_QA_POLICY_LAYER_COVERAGE_SCHEMA_INVALID:{layer['id']}")
        if not covered:
            raise ValueError(f"GOLD_QA_POLICY_LAYER_HAS_NO_COVERED_CHECKS:{layer['id']}")
        if layer["coverage"] == "COMPLETE" and uncovered:
            raise ValueError(f"GOLD_QA_POLICY_COMPLETE_LAYER_HAS_UNCOVERED:{layer['id']}")
    if policy.get("unknown_issue_layer") != "UNCLASSIFIED":
        raise ValueError("GOLD_QA_POLICY_UNKNOWN_LAYER_INVALID")
    for mapping_name in ("issue_code_layers", "issue_code_prefix_layers", "issue_code_suffix_layers"):
        mapping = policy.get(mapping_name)
        if not isinstance(mapping, dict) or any(str(layer_id) not in layer_ids for layer_id in mapping.values()):
            raise ValueError(f"GOLD_QA_POLICY_MAPPING_INVALID:{mapping_name}")
    return policy


def _issue_layer(code: str, policy: dict[str, Any]) -> str:
    exact = policy["issue_code_layers"].get(code)
    if exact:
        return str(exact)
    for prefix, layer_id in policy["issue_code_prefix_layers"].items():
        if code.startswith(prefix):
            return str(layer_id)
    for suffix, layer_id in policy["issue_code_suffix_layers"].items():
        if code.endswith(suffix):
            return str(layer_id)
    return str(policy["unknown_issue_layer"])


def audit_workbooks(
    source_path: Path, target_path: Path, sheet_name: str, rules_path: Path,
    qa_policy_path: Path | None = None, category_dictionary_path: Path | None = None,
    term_dictionary_path: Path | None = None, brand_dictionary_path: Path | None = None,
) -> dict[str, Any]:
    source_headers, source_rows = _read_sheet(source_path, sheet_name)
    target_headers, target_rows = _read_sheet(target_path, sheet_name)
    rules = load_detail_terminology_rules(rules_path)
    description_policy_path = ROOT / "config/stage5/description_fidelity_policy.json"
    description_policy = load_description_fidelity_policy(description_policy_path)
    issues: list[dict[str, Any]] = []
    key_variants: dict[str, dict[str, Counter[str]]] = defaultdict(lambda: defaultdict(Counter))
    value_variants: dict[tuple[str, str], dict[tuple[str, str], Counter[str]]] = defaultdict(lambda: defaultdict(Counter))
    shared = set(source_rows) & set(target_rows)
    regression_records = []
    for sku in sorted(shared):
        source_row, target_row = source_rows[sku], target_rows[sku]
        record = {"sku": sku}
        for field in FIELD_ALIASES:
            if field == "sku":
                continue
            record[f"{field}_es"] = str(source_row.get(field) or "")
            record[f"{field}_zh"] = str(target_row.get(field) or "")
        regression_records.append(record)
    regression_cases_path = ROOT / "data/qa/localization_regressions_v1.jsonl"
    try:
        regression_cases = load_regression_cases(regression_cases_path)
        regression_summaries = summarize_occurrences(regression_records, regression_cases)
        historical_regression = {
            "status": "PASS" if all(item["regression_passed"] for item in regression_summaries) else "FAIL",
            "policy_version": "LOCALIZATION_REGRESSION_V1",
            "case_count": len(regression_cases),
            "summaries": regression_summaries,
            "wrong_target_count": sum(item["wrong_target_count"] for item in regression_summaries),
            "expected_target_count": sum(item["expected_target_count"] for item in regression_summaries),
            "source_sku_count": len(source_rows),
            "target_sku_count": len(target_rows),
            "matched_sku_count": len(shared),
            "source_only_sku_count": len(set(source_rows) - shared),
            "target_only_sku_count": len(set(target_rows) - shared),
            "read_only": True,
            "master_writes": 0,
            "production_apply": False,
        }
        if historical_regression["wrong_target_count"]:
            issues.extend({
                "sku": "",
                "field": item["field"],
                "code": "HISTORICAL_REGRESSION_WRONG_TARGET",
                "review_only": False,
                "evidence": json.dumps(item, ensure_ascii=False),
            } for item in regression_summaries if item["wrong_target_count"])
    except (OSError, ValueError) as exc:
        historical_regression = {
            "status": "UNAVAILABLE",
            "policy_version": "LOCALIZATION_REGRESSION_V1",
            "error": str(exc),
            "read_only": True,
            "master_writes": 0,
            "production_apply": False,
        }
        issues.append({
            "sku": "", "field": "regression", "code": "HISTORICAL_REGRESSION_UNAVAILABLE",
            "review_only": False, "evidence": str(exc),
        })
    category_dictionary_path = category_dictionary_path or ROOT / "data/dictionary/category_dictionary.csv"
    term_dictionary_path = term_dictionary_path or ROOT / "data/dictionary/term_dictionary.csv"
    brand_dictionary_path = brand_dictionary_path or ROOT / "data/dictionary/brand_dictionary.csv"
    approved_categories = _load_approved_category_map(category_dictionary_path)
    confirmed_brands = _load_confirmed_brand_phrases(brand_dictionary_path)
    source_fact_fields = {
        field for field, aliases in FACT_FIELD_ALIASES.items()
        if any(_norm(alias) in {_norm(header) for header in source_headers} for alias in aliases)
    }
    target_fact_fields = {
        field for field, aliases in FACT_FIELD_ALIASES.items()
        if any(_norm(alias) in {_norm(header) for header in target_headers} for alias in aliases)
    }
    fact_field_coverage: dict[str, str] = {}
    if not source_rows:
        issues.append({"sku": "", "field": "SKU", "code": "SOURCE_WORKBOOK_EMPTY"})
    if not target_rows:
        issues.append({"sku": "", "field": "SKU", "code": "TARGET_WORKBOOK_EMPTY"})
    for field in FACT_FIELD_ALIASES:
        source_has, target_has = field in source_fact_fields, field in target_fact_fields
        if source_has and target_has:
            fact_field_coverage[field] = "CHECKED"
            for sku in sorted(shared):
                source_row, target_row = source_rows[sku], target_rows[sku]
                code = _compare_fact_value(field, source_row, target_row)
                if code:
                    issues.append({
                        "sku": sku, "field": field, "code": code,
                        "source": source_row.get(field, ""), "candidate": target_row.get(field, ""),
                    })
        elif source_has != target_has:
            fact_field_coverage[field] = "SCHEMA_MISMATCH"
            issues.append({
                "sku": "", "field": field, "code": "FACT_COLUMN_SCOPE_MISMATCH",
                "evidence": json.dumps({"source_has_column": source_has, "target_has_column": target_has}),
            })
        else:
            fact_field_coverage[field] = "NOT_PROVIDED"

    if set(source_rows) != set(target_rows):
        issues.append({
            "sku": "", "field": "SKU", "code": "SKU_SET_MISMATCH",
            "evidence": json.dumps({"source_only": sorted(set(source_rows) - set(target_rows)),
                                     "target_only": sorted(set(target_rows) - set(source_rows))}, ensure_ascii=False),
        })
    if list(source_rows) != list(target_rows):
        first_difference = next(
            (index for index, pair in enumerate(zip(source_rows, target_rows)) if pair[0] != pair[1]),
            min(len(source_rows), len(target_rows)),
        )
        issues.append({"sku": "", "field": "SKU", "code": "SKU_ORDER_MISMATCH",
                       "evidence": f"first_difference_index={first_difference}"})

    for sku in sorted(shared):
        source = source_rows[sku]
        target = target_rows[sku]
        context = {"name_es": source["name"], "cat1_es": source["cat1"], "cat2_es": source["cat2"]}
        for field in ("name", "description"):
            source_semantics = {"name": source["name"], "description": source["description"]}
            _, semantic_findings = inspect_source_facts(
                source_semantics, {field: target[field]},
                confirmed_brand_phrases=confirmed_brands.get("phrases", ()),
            )
            if semantic_findings:
                issues.append({
                    "sku": sku, "field": field, "code": "SOURCE_FACT_REVIEW_REQUIRED",
                    "source": source[field], "candidate": target[field],
                    "evidence": json.dumps(semantic_findings, ensure_ascii=False),
                })
        for field in FIELD_ALIASES:
            if field == "sku":
                continue
            source_value, target_value = source[field], target[field]
            checks = validate_model_output(
                {field: source_value}, {field: target_value}, expected_fields=(field,),
            )
            for code in checks.field_reasons.get(field, ()):
                if code in GUARD_CODES:
                    issues.append({"sku": sku, "field": field, "code": code,
                                   "source": source_value, "candidate": target_value})

        repaired, rule_results = repair_detail_candidate(
            source["details"], target["details"], rules, context=context,
        )
        for result in rule_results:
            issues.append({"sku": sku, "field": "details", "code": result,
                           "source": source["details"], "candidate": target["details"],
                           "proposed": repaired})

        source_record = _source_record(source)
        source_evidence_rows = source_consistency_evidence(source_record)
        source_flags = sorted({str(item.get("code") or "") for item in source_evidence_rows if item.get("code")})
        for flag in source_flags:
            if flag.startswith("SOURCE_ANOMALY_") or flag.endswith("_CONFLICT"):
                matching_evidence = [item for item in source_evidence_rows if item.get("code") == flag]
                issues.append({"sku": sku, "field": "source", "code": flag,
                               "evidence": json.dumps(matching_evidence, ensure_ascii=False, sort_keys=True)
                               if matching_evidence else source["details"]})

        source_pairs = parse_details(source["details"])
        target_pairs = parse_details(target["details"])
        if len(source_pairs) == len(target_pairs):
            for source_pair, target_pair in zip(source_pairs, target_pairs, strict=True):
                key_variants[source_pair.normalized_key][target_pair.key_es][sku] += 1
                value_variants[(source_pair.normalized_key, source_pair.normalized_value)][
                    (target_pair.key_es, target_pair.value_es)
                ][sku] += 1
        else:
            issues.append({
                "sku": sku, "field": "details", "code": "DETAIL_PAIR_ALIGNMENT_REVIEW",
                "source": source["details"], "candidate": target["details"],
                "evidence": json.dumps({
                    "source_pair_count": len(source_pairs),
                    "target_pair_count": len(target_pairs),
                    "automatic_pair_comparison_skipped": True,
                }, ensure_ascii=False),
            })

        compression_finding = description_compression_finding(
            source["description"], target["description"], description_policy,
        )
        if compression_finding:
            issues.append({"sku": sku, "field": "description", "code": "DESCRIPTION_COMPRESSION_REVIEW",
                           "evidence": json.dumps(compression_finding, ensure_ascii=False)})

    category_mapping_report = _audit_category_rows(
        shared, source_rows, target_rows, approved_categories, issues,
    )
    term_dictionary_report = _audit_term_dictionary(
        source_rows, target_rows, term_dictionary_path, issues,
    )
    business_tag_report = _audit_business_tags(shared, source_rows, target_rows, issues)
    official_label_report = _audit_official_labels(shared, source_rows, target_rows, issues)

    for source_key, translations in key_variants.items():
        if len(translations) > 1:
            observed_translations = [
                {
                    "target_key": target_key,
                    "occurrences": sum(sku_counts.values()),
                    "sku_count": len(sku_counts),
                    "sku_examples": sorted(sku_counts)[:20],
                }
                for target_key, sku_counts in sorted(translations.items())
            ]
            sku_examples = sorted({sku for sku_counts in translations.values() for sku in sku_counts})[:20]
            issues.append({"sku": "", "field": "details", "code": "SOURCE_KEY_TRANSLATION_VARIANTS",
                           "source": source_key, "candidate": json.dumps(sorted(translations), ensure_ascii=False),
                           "evidence": json.dumps({
                               "sku_examples": sku_examples, "observed_translations": observed_translations,
                               "heuristic_only": True,
                           }, ensure_ascii=False)})
    for (source_key, source_value), translations in value_variants.items():
        if len(translations) > 1:
            observed_translations = [
                {
                    "target_key": target_key, "target_value": target_value,
                    "occurrences": sum(sku_counts.values()),
                    "sku_count": len(sku_counts),
                    "sku_examples": sorted(sku_counts)[:20],
                }
                for (target_key, target_value), sku_counts in sorted(translations.items())
            ]
            sku_examples = sorted({sku for sku_counts in translations.values() for sku in sku_counts})[:20]
            issues.append({"sku": "", "field": "details", "code": "SOURCE_VALUE_TRANSLATION_VARIANTS",
                           "source": f"{source_key}:{source_value}",
                           "candidate": json.dumps([list(item) for item in sorted(translations)], ensure_ascii=False),
                           "evidence": json.dumps({
                               "sku_examples": sku_examples, "observed_translations": observed_translations,
                               "heuristic_only": True,
                           }, ensure_ascii=False)})

    raw_issue_count = len(issues)
    issues, issue_dedupe = _dedupe_issues(issues)
    counts = Counter(issue["code"] for issue in issues)
    title_policy = load_title_display_policy(ROOT / "config/stage5/title_display_policy.json")
    qa_policy_path = qa_policy_path or ROOT / "config/stage5/chinese_gold_qa_policy.json"
    qa_policy = _load_qa_policy(qa_policy_path)
    layer_counts: Counter[str] = Counter()
    unclassified_codes: set[str] = set()
    for issue in issues:
        code = str(issue.get("code") or "")
        layer_id = _issue_layer(code, qa_policy)
        issue["qa_layer"] = layer_id
        layer_counts[layer_id] += 1
        if layer_id == qa_policy["unknown_issue_layer"]:
            unclassified_codes.add(code)
    layer_summary = [
        {
            **layer,
            "issue_count": layer_counts[str(layer["id"])],
            "status": "FINDINGS" if layer_counts[str(layer["id"])] else "NO_FINDINGS_IN_SCOPE",
        }
        for layer in qa_policy["layers"]
    ]
    all_layers_complete = all(layer.get("coverage") == "COMPLETE" for layer in qa_policy["layers"])
    full_gold_eligible = bool(
        qa_policy.get("gold_claim_allowed") and all_layers_complete
        and not issues and not unclassified_codes
    )
    return {
        "source_file": str(source_path), "target_file": str(target_path), "sheet": sheet_name,
        "source_sha256": _sha256_file(source_path), "target_sha256": _sha256_file(target_path),
        "source_headers": source_headers, "target_headers": target_headers,
        "source_sku_count": len(source_rows), "target_sku_count": len(target_rows),
        "matched_sku_count": len(shared), "issue_count": len(issues), "issue_counts": dict(sorted(counts.items())),
        "raw_issue_count": raw_issue_count,
        "duplicate_issue_count": issue_dedupe["duplicate_issue_count"],
        "unique_issue_count": issue_dedupe["unique_issue_count"],
        "l1_fact_field_coverage": fact_field_coverage,
        "category_mapping_audit": category_mapping_report,
        "confirmed_brand_dictionary_audit": {
            "status": confirmed_brands.get("status"),
            "sha256": confirmed_brands.get("sha256"),
            "phrase_count": confirmed_brands.get("phrase_count", 0),
            "scope": "MODEL_TOKEN_DISAMBIGUATION_ONLY_NO_BRAND_TITLE_POLICY_UNCHANGED",
        },
        "category_rule_review_rows": _category_rule_review_rows(issues),
        "term_rule_review_rows": _term_rule_review_rows(issues),
        "term_dictionary_audit": term_dictionary_report,
        "approved_term_checker_sha256": _sha256_file(APPROVED_TERM_CHECKER_PATH),
        "business_tag_audit": business_tag_report,
        "official_label_audit": official_label_report,
        "detail_terminology_version": rules.get("version"),
        "detail_terminology_rules_sha256": hashlib.sha256(rules_path.read_bytes()).hexdigest(),
        "description_fidelity_policy_id": description_policy.get("policy_id"),
        "description_fidelity_policy_sha256": hashlib.sha256(description_policy_path.read_bytes()).hexdigest(),
        "title_display_policy_id": title_policy.get("policy_id"),
        "title_display_policy_sha256": hashlib.sha256(
            (ROOT / "config/stage5/title_display_policy.json").read_bytes()
        ).hexdigest(),
        "gold_qa_policy_id": qa_policy.get("policy_id"),
        "gold_qa_policy_sha256": hashlib.sha256(qa_policy_path.read_bytes()).hexdigest(),
        "qa_layers": layer_summary,
        "unclassified_issue_codes": sorted(unclassified_codes),
        "unclassified_issue_count": layer_counts["UNCLASSIFIED"],
        "full_gold_eligible": full_gold_eligible,
        "full_gold_scope_status": "COMPLETE" if all_layers_complete else "PARTIAL",
        "source_consistency_rules": source_consistency_rules_manifest(),
        "historical_regression": historical_regression,
        "detail_rule_coverage": _detail_rule_coverage(source_rows, target_rows, rules),
        "writes_source_workbook": False, "writes_target_workbook": False,
        "issues": issues,
    }


def write_qa_log(path: Path, report: dict[str, Any]) -> Path:
    """Write a deterministic, separate QA finding ledger for a read-only audit.

    This artifact is review metadata, not a product export: it never modifies
    either workbook or writes QA/process notes into the business remarks field.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = (
        "qa_id", "sku", "field", "qa_layer", "code", "review_only",
        "source_sha256", "target_sha256", "source", "candidate", "evidence", "issue_json",
    )
    rows = []
    duplicate_issue_occurrences: Counter[str] = Counter()
    for issue in report.get("issues", ()):
        payload = json.dumps(issue, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        occurrence = duplicate_issue_occurrences[payload]
        duplicate_issue_occurrences[payload] += 1
        identity = "|".join((
            str(report.get("source_sha256") or ""),
            str(report.get("target_sha256") or ""), payload, str(occurrence),
        ))
        rows.append({
            "qa_id": "qa-" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24],
            "sku": issue.get("sku", ""), "field": issue.get("field", ""),
            "qa_layer": issue.get("qa_layer", "UNCLASSIFIED"), "code": issue.get("code", ""),
            "review_only": bool(issue.get("review_only", False)),
            "source_sha256": report.get("source_sha256", ""),
            "target_sha256": report.get("target_sha256", ""),
            "source": issue.get("source", ""), "candidate": issue.get("candidate", ""),
            "evidence": issue.get("evidence", ""), "issue_json": payload,
        })
    rows.sort(key=lambda item: (str(item["sku"]), str(item["field"]), str(item["code"]), item["qa_id"]))
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return path


def write_source_anomaly_log(path: Path, report: dict[str, Any]) -> Path:
    """Write source anomalies separately from translation/process findings.

    The Spanish source remains authoritative and unchanged.  This ledger is
    an audit sidecar so a Chinese export can filter or neutralize a malformed
    source field without losing the original evidence or its source/target
    snapshot binding.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = (
        "anomaly_id", "sku", "field", "code", "source_sha256", "target_sha256",
        "source", "candidate", "evidence", "status", "action",
    )
    rows: list[dict[str, Any]] = []
    for issue in report.get("issues", ()):
        code = str(issue.get("code") or "")
        if not (code.startswith("SOURCE_ANOMALY_") or code.endswith("_CONFLICT")):
            continue
        identity = {
            "sku": str(issue.get("sku") or ""),
            "field": str(issue.get("field") or "source"),
            "code": code,
            "source": str(issue.get("source") or ""),
            "evidence": issue.get("evidence") or "",
        }
        anomaly_id = "anomaly-" + hashlib.sha256(
            json.dumps(identity, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
        ).hexdigest()[:24]
        rows.append({
            "anomaly_id": anomaly_id,
            "sku": identity["sku"], "field": identity["field"], "code": code,
            "source_sha256": report.get("source_sha256", ""),
            "target_sha256": report.get("target_sha256", ""),
            "source": identity["source"],
            "candidate": str(issue.get("candidate") or ""),
            "evidence": json.dumps(identity["evidence"], ensure_ascii=False, sort_keys=True)
            if not isinstance(identity["evidence"], str) else identity["evidence"],
            "status": "OPEN_REVIEW", "action": "SOURCE_UNCHANGED_REVIEW_ONLY",
        })
    rows.sort(key=lambda row: (row["sku"], row["code"], row["field"], row["anomaly_id"]))
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return path


def write_detail_rule_review_queue(
    path: Path, rows: list[dict[str, Any]], *, source_sha256: str,
    target_sha256: str, rules_sha256: str,
) -> tuple[Path, Path]:
    """Write a complete, non-promoting review queue for detail terminology gaps."""
    path = Path(path)
    manifest_path = path.with_suffix(path.suffix + ".manifest.json")
    columns = (
        "review_id", "review_kind", "source_key_normalized", "source_value_normalized",
        "occurrences", "rule_covered_occurrences", "uncovered_occurrences", "sku_count",
        "sku_examples", "observed_target_candidates", "single_observed_candidate",
        "review_reasons", "candidate_is_approved", "sku", "source_details", "target_details",
        "source_pair_count", "target_pair_count", "source_sha256", "target_sha256",
        "detail_rules_sha256",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    ordered = sorted(rows, key=lambda item: (
        str(item.get("review_kind") or ""), str(item.get("source_key_normalized") or ""),
        str(item.get("source_value_normalized") or ""), str(item.get("sku") or ""),
    ))
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        for row in ordered:
            identity = "|".join((
                source_sha256, target_sha256, str(row.get("review_kind") or ""),
                str(row.get("source_key_normalized") or ""),
                str(row.get("source_value_normalized") or ""), str(row.get("sku") or ""),
            ))
            writer.writerow({
                **row,
                "review_id": "detail-" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24],
                "candidate_is_approved": False,
                "source_sha256": source_sha256,
                "target_sha256": target_sha256,
                "detail_rules_sha256": rules_sha256,
            })
    manifest = {
        "schema": "ACTION_DETAIL_RULE_REVIEW_QUEUE_V1",
        "artifact": path.name,
        "row_count": len(ordered),
        "sha256": _sha256_file(path),
        "source_sha256": source_sha256,
        "target_sha256": target_sha256,
        "detail_rules_sha256": rules_sha256,
        "review_only": True,
        "auto_apply": False,
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path, manifest_path


def write_category_rule_review_queue(
    path: Path, rows: list[dict[str, Any]], *, source_sha256: str,
    target_sha256: str, category_dictionary_sha256: str,
) -> tuple[Path, Path]:
    """Write scoped category mapping candidates without promoting them."""
    path = Path(path)
    manifest_path = path.with_suffix(path.suffix + ".manifest.json")
    columns = (
        "review_id", "review_kind", "field", "source_cat1", "source_category",
        "observed_target_cat1", "occurrences", "sku_count", "sku_examples",
        "observed_target_values", "single_observed_target_value", "approved_values",
        "candidate_is_approved", "source_sha256", "target_sha256",
        "category_dictionary_sha256",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    ordered = sorted(rows, key=lambda item: (
        str(item.get("field") or ""), str(item.get("source_cat1") or ""),
        str(item.get("source_category") or ""), str(item.get("observed_target_cat1") or ""),
        str(item.get("review_kind") or ""),
    ))
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        for row in ordered:
            identity = "|".join((
                source_sha256, target_sha256, category_dictionary_sha256,
                str(row.get("field") or ""), str(row.get("source_cat1") or ""),
                str(row.get("source_category") or ""),
                str(row.get("observed_target_cat1") or ""), str(row.get("review_kind") or ""),
            ))
            writer.writerow({
                **row,
                "review_id": "category-" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24],
                "candidate_is_approved": False,
                "source_sha256": source_sha256,
                "target_sha256": target_sha256,
                "category_dictionary_sha256": category_dictionary_sha256,
            })
    manifest = {
        "schema": "ACTION_CATEGORY_RULE_REVIEW_QUEUE_V1",
        "artifact": path.name,
        "row_count": len(ordered),
        "sha256": _sha256_file(path),
        "source_sha256": source_sha256,
        "target_sha256": target_sha256,
        "category_dictionary_sha256": category_dictionary_sha256,
        "review_only": True,
        "auto_apply": False,
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path, manifest_path


def write_term_rule_review_queue(
    path: Path, rows: list[dict[str, Any]], *, source_sha256: str,
    target_sha256: str, term_dictionary_sha256: str,
) -> tuple[Path, Path]:
    """Write a compact, source-bound glossary review queue; never approves candidates."""
    path = Path(path)
    manifest_path = path.with_suffix(path.suffix + ".manifest.json")
    columns = (
        "review_id", "review_kind", "source_term", "term_type", "source_field", "target_field",
        "approved_translation", "forbidden_candidate", "occurrences", "sku_count", "sku_examples",
        "candidate_variant_count", "candidate_examples", "candidate_examples_truncated",
        "evidence_examples", "review_only", "candidate_is_approved", "source_sha256",
        "target_sha256", "term_dictionary_sha256",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    ordered = sorted(rows, key=lambda item: (
        str(item.get("review_kind") or ""), _norm(item.get("source_term")),
        _norm(item.get("term_type")), _norm(item.get("source_field")),
        _norm(item.get("target_field")), _norm(item.get("approved_translation")),
        _norm(item.get("forbidden_candidate")),
    ))
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        for row in ordered:
            identity = "|".join((
                source_sha256, target_sha256, term_dictionary_sha256,
                str(row.get("review_kind") or ""), _norm(row.get("source_term")),
                _norm(row.get("term_type")), _norm(row.get("source_field")),
                _norm(row.get("target_field")), _norm(row.get("approved_translation")),
                _norm(row.get("forbidden_candidate")),
            ))
            writer.writerow({
                **row,
                "review_id": "term-" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24],
                "candidate_is_approved": False,
                "source_sha256": source_sha256,
                "target_sha256": target_sha256,
                "term_dictionary_sha256": term_dictionary_sha256,
            })
    manifest = {
        "schema": "ACTION_TERM_RULE_REVIEW_QUEUE_V1",
        "artifact": path.name,
        "row_count": len(ordered),
        "sha256": _sha256_file(path),
        "source_sha256": source_sha256,
        "target_sha256": target_sha256,
        "term_dictionary_sha256": term_dictionary_sha256,
        "review_only": True,
        "auto_apply": False,
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path, manifest_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path, help="Spanish source workbook")
    parser.add_argument("--target", required=True, type=Path, help="Chinese candidate workbook")
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--sheet", default="商品全量")
    parser.add_argument("--rules", type=Path, default=ROOT / "config/stage5/detail_terminology_rules.json")
    parser.add_argument("--qa-policy", type=Path, default=ROOT / "config/stage5/chinese_gold_qa_policy.json")
    parser.add_argument("--category-dictionary", type=Path, default=ROOT / "data/dictionary/category_dictionary.csv")
    parser.add_argument("--term-dictionary", type=Path, default=ROOT / "data/dictionary/term_dictionary.csv")
    parser.add_argument("--brand-dictionary", type=Path, default=ROOT / "data/dictionary/brand_dictionary.csv")
    args = parser.parse_args()
    report = audit_workbooks(
        args.source, args.target, args.sheet, args.rules, args.qa_policy, args.category_dictionary,
        args.term_dictionary, args.brand_dictionary,
    )
    args.output_dir.mkdir(parents=True, exist_ok=True)
    detail_queue_rows = report.get("detail_rule_coverage", {}).pop("review_queue_rows", [])
    detail_queue_path, detail_queue_manifest_path = write_detail_rule_review_queue(
        args.output_dir / "DETAIL_RULE_REVIEW.csv", detail_queue_rows,
        source_sha256=report["source_sha256"], target_sha256=report["target_sha256"],
        rules_sha256=report.get("detail_terminology_rules_sha256", ""),
    )
    category_queue_rows = report.pop("category_rule_review_rows", [])
    category_queue_path, category_queue_manifest_path = write_category_rule_review_queue(
        args.output_dir / "CATEGORY_RULE_REVIEW.csv", category_queue_rows,
        source_sha256=report["source_sha256"], target_sha256=report["target_sha256"],
        category_dictionary_sha256=report.get("category_mapping_audit", {}).get("sha256") or "",
    )
    term_queue_rows = report.pop("term_rule_review_rows", [])
    term_queue_path, term_queue_manifest_path = write_term_rule_review_queue(
        args.output_dir / "TERM_RULE_REVIEW.csv", term_queue_rows,
        source_sha256=report["source_sha256"], target_sha256=report["target_sha256"],
        term_dictionary_sha256=report.get("term_dictionary_audit", {}).get("sha256") or "",
    )
    report["detail_rule_review_queue_artifact"] = {
        "path": detail_queue_path.name,
        "sha256": _sha256_file(detail_queue_path),
        "row_count": len(detail_queue_rows),
        "manifest_path": detail_queue_manifest_path.name,
        "manifest_sha256": _sha256_file(detail_queue_manifest_path),
        "auto_apply": False,
    }
    report["category_rule_review_queue_artifact"] = {
        "path": category_queue_path.name,
        "sha256": _sha256_file(category_queue_path),
        "row_count": len(category_queue_rows),
        "manifest_path": category_queue_manifest_path.name,
        "manifest_sha256": _sha256_file(category_queue_manifest_path),
        "auto_apply": False,
    }
    report["term_rule_review_queue_artifact"] = {
        "path": term_queue_path.name,
        "sha256": _sha256_file(term_queue_path),
        "row_count": len(term_queue_rows),
        "manifest_path": term_queue_manifest_path.name,
        "manifest_sha256": _sha256_file(term_queue_manifest_path),
        "auto_apply": False,
    }
    report_path = args.output_dir / "localization_audit.json"
    qa_log_path = write_qa_log(args.output_dir / "QA_LOG.csv", report)
    report["qa_log_artifact"] = {
        "path": qa_log_path.name,
        "sha256": _sha256_file(qa_log_path),
        "row_count": len(report.get("issues", ())),
    }
    source_anomaly_path = write_source_anomaly_log(args.output_dir / "SOURCE_ANOMALY.csv", report)
    report["source_anomaly_artifact"] = {
        "path": source_anomaly_path.name,
        "sha256": _sha256_file(source_anomaly_path),
        "row_count": sum(
            1 for issue in report.get("issues", ())
            if str(issue.get("code") or "").startswith("SOURCE_ANOMALY_")
            or str(issue.get("code") or "").endswith("_CONFLICT")
        ),
        "source_unchanged": True,
        "auto_apply": False,
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    coverage = report.get("detail_rule_coverage", {})
    console_summary = {
        "source_file": report.get("source_file"),
        "target_file": report.get("target_file"),
        "source_sku_count": report.get("source_sku_count"),
        "target_sku_count": report.get("target_sku_count"),
        "matched_sku_count": report.get("matched_sku_count"),
        "issue_count": report.get("issue_count"),
        "issue_counts": report.get("issue_counts", {}),
        "full_gold_eligible": report.get("full_gold_eligible"),
        "detail_rule_coverage": {
            key: value for key, value in coverage.items()
            if key not in {
                "unmapped_source_keys", "unmapped_source_key_value_pairs",
                "source_key_value_translation_variants", "review_queue_rows",
            }
        },
        "detail_rule_review_queue_artifact": report.get("detail_rule_review_queue_artifact"),
        "category_rule_review_queue_artifact": report.get("category_rule_review_queue_artifact"),
        "term_rule_review_queue_artifact": report.get("term_rule_review_queue_artifact"),
        "qa_log_artifact": report.get("qa_log_artifact"),
        "source_anomaly_artifact": report.get("source_anomaly_artifact"),
    }
    # Windows consoles may use a legacy code page. Emit ASCII JSON on stdout;
    # the full UTF-8 artifacts retain their native Chinese text.
    print(json.dumps(console_summary, ensure_ascii=True, indent=2))
    print(f"report={report_path}")
    print(f"qa_log={qa_log_path}")
    return 0 if not report["issue_count"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
