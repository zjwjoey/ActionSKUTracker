"""Prepare a read-only, disjoint remediation candidate set for two-pass review.

This does not create Gold.  It merely binds each selected official Spanish
source row to its current Chinese projection, checks that the Spanish facts
still match the queue, and emits a resumable review input for the existing
DeepSeek + deterministic-guard workflow.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from action_tracker.services.hashing import localization_source_hash

RUN_DIR = ROOT / "runtime/training/qwen3_8b/20260911"
MASTER = ROOT / "runtime/master/Action_Master.xlsx"
FIELDS = ("name", "cat1", "cat2", "spec", "description", "details")
ES_COLUMNS = {
    "name": "西班牙语品名", "cat1": "一级类目（西语）", "cat2": "二级类目（西语）",
    "spec": "规格（西语）", "description": "描述（西语）", "details": "产品详情（西语）",
}
ZH_COLUMNS = {
    "name": "中文品名", "cat1": "一级类目（中文）", "cat2": "二级类目（中文）",
    "spec": "规格（中文）", "description": "中文描述", "details": "中文产品详情",
}
QUEUE_COLUMNS = {
    "name": "source_name_es", "spec": "source_spec_es",
    "description": "source_description_es", "details": "source_details_es",
}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_sheet_by_sku(sheet) -> dict[str, dict[str, str]]:
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value or "") for value in next(iterator)]
    rows: dict[str, dict[str, str]] = {}
    for values in iterator:
        row = {header: str(value or "").strip() for header, value in zip(headers, values)}
        sku = row.get("SKU", "")
        if sku:
            rows[sku] = row
    return rows


def source_payload(es_row: dict[str, str]) -> dict[str, str]:
    return {field: es_row.get(column, "").strip() for field, column in ES_COLUMNS.items()}


def target_payload(zh_row: dict[str, str]) -> dict[str, str]:
    return {field: zh_row.get(column, "").strip() for field, column in ZH_COLUMNS.items()}


def build_candidate(
    queue_row: dict[str, str], es_row: dict[str, str], zh_row: dict[str, str],
) -> tuple[dict[str, object] | None, str]:
    sku = str(queue_row.get("sku") or "").strip()
    if not sku or es_row.get("SKU") != sku or zh_row.get("SKU") != sku:
        return None, "SKU_MASTER_JOIN_MISMATCH"
    source = source_payload(es_row)
    if not all(source.values()):
        return None, "SOURCE_FIELD_EMPTY"
    queue_pairs = {field: str(queue_row.get(column) or "").strip() for field, column in QUEUE_COLUMNS.items()}
    if any(source[field] != queue_pairs[field] for field in QUEUE_COLUMNS):
        return None, "SOURCE_QUEUE_MISMATCH"
    expected_hash = localization_source_hash({
        "name_es": source["name"], "cat1_es": source["cat1"], "cat2_es": source["cat2"],
        "spec_es": source["spec"], "desc_es": source["description"], "details_es": source["details"],
    })
    record = {
        "messages": [
            {"role": "system", "content": "将 Action 西语商品字段标准化为简洁、忠实的中文。保持数字、单位、数量、尺寸和品牌/型号事实；不得臆造源数据没有的信息。输出固定六字段 JSON。"},
            {"role": "user", "content": json.dumps(source, ensure_ascii=False, sort_keys=True)},
            {"role": "assistant", "content": json.dumps(target_payload(zh_row), ensure_ascii=False, sort_keys=True)},
        ],
        "metadata": {
            "sku": sku,
            "source_hash": expected_hash,
            "label_tier": "REMEDIATION_CANDIDATE_NOT_GOLD",
            "review_status": "PENDING_TWO_PASS_SOURCE_FIDELITY_REVIEW",
            "failure_class": str(queue_row.get("failure_class") or ""),
            "target_field": str(queue_row.get("target_field") or ""),
            "family_proxy": str(queue_row.get("family_proxy") or ""),
            "source_status": "MASTER_OFFICIAL_FACT_MATCHED",
            "training_eligible": False,
        },
    }
    return record, "READY_FOR_TWO_PASS_REVIEW"


def main() -> int:
    parser = argparse.ArgumentParser(description="Build Stage-4 remediation review candidates")
    parser.add_argument("--queue", default=str(RUN_DIR / "stage4_targeted_remediation_review_queue_v3.csv"))
    parser.add_argument("--output", default=str(RUN_DIR / "qwen_incremental_remediation_candidate_50.jsonl"))
    parser.add_argument("--manifest", default=str(RUN_DIR / "qwen_incremental_remediation_candidate_50.manifest.json"))
    args = parser.parse_args()
    queue_path, output_path, manifest_path = (Path(args.queue), Path(args.output), Path(args.manifest))
    queue = list(csv.DictReader(queue_path.open(encoding="utf-8-sig", newline="")))
    workbook = load_workbook(MASTER, read_only=True, data_only=True)
    es_rows = read_sheet_by_sku(workbook["02_SKU_ES_CURRENT"])
    zh_rows = read_sheet_by_sku(workbook["01_SKU_ZH_CURRENT"])
    ready, blocked = [], []
    for row in queue:
        candidate, status = build_candidate(row, es_rows.get(row.get("sku", ""), {}), zh_rows.get(row.get("sku", ""), {}))
        if candidate is None:
            blocked.append({"sku": row.get("sku", ""), "status": status})
        else:
            ready.append(candidate)
    output_path.write_text(
        "\n".join(json.dumps(item, ensure_ascii=False, sort_keys=True) for item in ready) + ("\n" if ready else ""),
        encoding="utf-8",
    )
    manifest = {
        "status": "REMEDIATION_CANDIDATES_PREPARED",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "queue": {"path": str(queue_path), "sha256": sha(queue_path), "rows": len(queue)},
        "master": {"path": str(MASTER), "sha256": sha(MASTER)},
        "ready_for_two_pass_review": len(ready),
        "blocked_preflight": blocked,
        "training_eligible_rows": 0,
        "production_writes": False,
        "candidate_jsonl": {"path": str(output_path), "sha256": sha(output_path)},
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"ready": len(ready), "blocked": blocked, "output": str(output_path)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
