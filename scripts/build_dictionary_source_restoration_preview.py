"""Build a non-mutating source restoration preview from archived Spanish workbooks."""
from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from openpyxl import load_workbook


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def build(queue_path: Path, archive_dir: Path, output_path: Path) -> dict[str, int]:
    with queue_path.open("r", encoding="utf-8-sig", newline="") as handle:
        queue = list(csv.DictReader(handle))
    target = {row["sku"] for row in queue if row.get("sku")}
    found: dict[str, dict[str, str]] = {}
    files = sorted(archive_dir.rglob("*西语*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    for path in files:
        try:
            book = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        for sheet in book.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            try:
                header = next(iterator)
            except StopIteration:
                continue
            columns = {_text(value): index for index, value in enumerate(header) if value is not None}
            if "编号" not in columns or "标题" not in columns:
                continue
            for values in iterator:
                sku = _text(values[columns["编号"]])
                if sku not in target or sku in found:
                    continue
                row = {name: _text(values[columns[name]]) for name in ("标题", "分类1", "分类2", "规格", "描述", "产品详情") if name in columns}
                # Only accept a row as authoritative Spanish evidence when the
                # title and all non-empty source fields are free of CJK/encoding noise.
                joined = "\n".join(row.values())
                if re.search(r"[\u3400-\u9fff�]", joined):
                    continue
                found[sku] = {
                    "sku": sku,
                    "source_file": str(path),
                    "name_es_raw": row.get("标题", ""),
                    "cat1_es": row.get("分类1", ""),
                    "cat2_es": row.get("分类2", ""),
                    "spec_es_raw": row.get("规格", ""),
                    "description_es": row.get("描述", ""),
                    "details_es": row.get("产品详情", ""),
                    "restoration_status": "SOURCE_RESTORABLE",
                }
        book.close()
    rows = [found[key] for key in sorted(found)]
    columns = ["sku", "source_file", "name_es_raw", "cat1_es", "cat2_es", "spec_es_raw", "description_es", "details_es", "restoration_status"]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)
    return {"queue_sku_count": len(target), "restorable_sku_count": len(rows), "missing_sku_count": len(target - set(found)), "workbook_count": len(files)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--queue", default="runtime/dictionary/owner_review_20260915/current_dictionary_owner_review_queue.csv")
    parser.add_argument("--archive-dir", default=r"F:\按日期整理")
    parser.add_argument("--output", default="runtime/dictionary/owner_review_20260915/source_restoration_preview.csv")
    args = parser.parse_args()
    print(build(Path(args.queue), Path(args.archive_dir), Path(args.output)))
