"""Build a human-review-only targeted remediation queue.

Candidates come from the current Spanish Master projection, but no candidate is
promoted to Gold and no model/API is called.  Frozen 485 and every SKU already
in the combined train/validation/test splits are excluded.  Family isolation
is conservative: exact Spanish cat1/cat2 pairs used by the P0 rows are also
excluded, and the remaining family decision is left for human review.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import argparse
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"F:/ActionSKUTracker")
RUN_DIR = ROOT / "runtime/training/qwen3_8b/20260911"
MASTER = ROOT / "runtime/master/Action_Master.xlsx"
PREDICTIONS = RUN_DIR / "stage4_full_eval_owner_signed_20260912_predictions.jsonl"

NUM = re.compile(r"\d+(?:[.,]\d+)?")
TECH = re.compile(r"\b(?:IP\d{2}|USB(?:-[A-Z])?|LED|H[1-9]|P\d{1,2}/?\d?W|\d+(?:[.,]\d+)?\s?(?:W|V|kW|mAh|GB|TB|cm|mm|kg|g|L|ml))\b", re.I)
AMBIGUOUS = re.compile(r"\b(?:funda|alfombra|bolsa|pala|iluminador|póster|poster|taza|cesta|cubo|cepillo|bandeja|organizador|soporte)\b", re.I)
FAILURE_CLASSES = ("BRAND_FACT_LOSS", "NUMERIC_MISSING", "NUMERIC_HALLUCINATION", "TECH_TOKEN_LOSS", "PRODUCT_OBJECT_SEMANTIC_ERROR")


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def jsonl_skus(path: Path) -> set[str]:
    """Read SKU from either production-style or training-style JSONL rows."""
    if not path.exists():
        return set()
    values: set[str] = set()
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        sku = str(row.get("sku") or (row.get("metadata") or {}).get("sku") or "").strip()
        if sku:
            values.add(sku)
    return values


def has_complete_training_source(
    name: str, spec: str, description: str, details: str, family: tuple[str, str],
) -> bool:
    """A six-field SFT row cannot silently treat absent official fields as labels."""
    return all(str(value).strip() for value in (name, spec, description, details, family[0], family[1]))


def parse_class_quotas(value: str, default: int) -> dict[str, int]:
    """Parse a strict `CLASS=count,...` quota override for replacement pools."""
    quotas = {kind: default for kind in FAILURE_CLASSES}
    if not value.strip():
        return quotas
    for fragment in value.split(","):
        kind, separator, raw_count = fragment.strip().partition("=")
        if not separator or kind not in quotas or not raw_count.isdigit() or int(raw_count) <= 0:
            raise ValueError("INVALID_CLASS_QUOTAS")
        quotas[kind] = int(raw_count)
    return quotas


def load_excluded(
    date_dir: Path = RUN_DIR, predictions: Path = PREDICTIONS,
) -> tuple[set[str], set[str], set[tuple[str, str]]]:
    """Return frozen/evaluation and historical corpus exclusions.

    Every existing train, validation, candidate, or test set is excluded.  A
    SKU can only become targeted-remediation supervision if it has not already
    appeared in a frozen evaluation or historical model corpus.
    """
    # A daily directory is only the output location, not the complete corpus.
    # When called with the production 20260911 directory, scan its parent so
    # 20260908/20260910 train/validation/test artifacts cannot leak into a
    # later remediation queue.  Unit tests pass a temporary root directly.
    corpus_root = date_dir.parent if date_dir.name.isdigit() and date_dir.parent.name == "qwen3_8b" else date_dir
    frozen = jsonl_skus(predictions)
    for path in corpus_root.rglob("stage4_test_only_485.jsonl"):
        frozen.update(jsonl_skus(path))
    for path in corpus_root.rglob("qwen_hard_test_v1_candidate.jsonl"):
        frozen.update(jsonl_skus(path))
    split = set()
    for path in corpus_root.rglob("*.jsonl"):
        name = path.name
        if (
            name.endswith(("_train.jsonl", "_validation.jsonl", "_test.jsonl"))
            or name.startswith("qwen_incremental_")
            or "combined_gold_incremental" in path.parts
            or "field_conditioned_v1" in path.parts
        ):
            split.update(jsonl_skus(path))
    p0_families = set()
    for line in predictions.open(encoding="utf-8"):
        if not line.strip():
            continue
        row = json.loads(line)
        if row.get("sku") in {"3219003", "3220894", "3217699", "3223271", "3221705", "3221795", "3201998", "3213267", "3009588", "3209605"}:
            source = row.get("source") or {}
            p0_families.add((str(source.get("cat1") or ""), str(source.get("cat2") or "")))
    return frozen, split, p0_families


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a disjoint Stage-4 remediation review queue")
    parser.add_argument(
        "--output-suffix", default="", help="Optional version suffix; preserves prior queue artifacts",
    )
    parser.add_argument(
        "--per-failure-class", type=int, default=10,
        help="Balanced count for each remediation failure class (positive integer)",
    )
    parser.add_argument(
        "--class-quotas", default="",
        help="Optional comma-separated overrides, e.g. BRAND_FACT_LOSS=40,NUMERIC_MISSING=22",
    )
    args = parser.parse_args()
    if args.per_failure_class <= 0:
        raise SystemExit("PER_FAILURE_CLASS_MUST_BE_POSITIVE")
    try:
        quotas = parse_class_quotas(args.class_quotas, args.per_failure_class)
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc
    suffix = ("_" + args.output_suffix.strip()) if args.output_suffix.strip() else ""
    out_csv = RUN_DIR / f"stage4_targeted_remediation_review_queue{suffix}.csv"
    out_json = RUN_DIR / f"stage4_targeted_remediation_review_queue{suffix}.json"
    frozen, split, p0_families = load_excluded()
    dictionary_path = ROOT / "data/dictionary/product_dictionary.csv"
    brand_by_sku = {}
    if dictionary_path.exists():
        with dictionary_path.open(encoding="utf-8-sig", newline="") as handle:
            for item in csv.DictReader(handle):
                brand_by_sku[str(item.get("sku") or "").strip()] = str(item.get("brand_id") or "").strip()
    wb = load_workbook(MASTER, read_only=True, data_only=True)
    ws = wb["02_SKU_ES_CURRENT"]
    iterator = ws.iter_rows(values_only=True)
    headers = list(next(iterator))
    idx = {str(value): i for i, value in enumerate(headers)}
    rows = []
    for values in iterator:
        sku = str(values[idx["SKU"]] or "").strip()
        family = (str(values[idx["一级类目（西语）"]] or ""), str(values[idx["二级类目（西语）"]] or ""))
        if not sku or sku in frozen or sku in split or family in p0_families:
            continue
        name = str(values[idx["西班牙语品名"]] or "")
        spec = str(values[idx["规格（西语）"]] or "")
        desc = str(values[idx["描述（西语）"]] or "")
        details = str(values[idx["产品详情（西语）"]] or "")
        if not has_complete_training_source(name, spec, desc, details, family):
            continue
        all_text = " ".join((name, spec, desc, details))
        candidates = []
        brand = brand_by_sku.get(sku, "")
        if brand and brand.lower() in name.lower():
            candidates.append(("BRAND_FACT_LOSS", "name"))
        if NUM.search(desc + " " + spec + " " + details):
            candidates.append(("NUMERIC_MISSING", "description"))
        if len(NUM.findall(desc)) >= 2 or len(NUM.findall(details)) >= 3:
            candidates.append(("NUMERIC_HALLUCINATION", "description"))
        if TECH.search(all_text):
            candidates.append(("TECH_TOKEN_LOSS", "details"))
        if AMBIGUOUS.search(name):
            candidates.append(("PRODUCT_OBJECT_SEMANTIC_ERROR", "name"))
        for failure_class, field in candidates:
            rows.append({
                "sku": sku,
                "failure_class": failure_class,
                "target_field": field,
                "family_proxy": f"{family[0]} / {family[1]}",
                "source_name_es": name,
                "source_spec_es": spec,
                "source_description_es": desc,
                "source_details_es": details,
                "current_name_zh_for_review": "",
                "source_status": "NEEDS_OFFICIAL_CONFIRMATION",
                "gold_status": "NOT_GOLD",
                "human_disposition": "PENDING",
                "final_value": "",
                "reviewer": "",
                "reviewed_at": "",
                "family_isolation_status": "EXACT_P0_CAT_PAIR_EXCLUDED_REVIEW_REQUIRED",
                "retrain_eligible": "NO_UNTIL_HUMAN_CONFIRMED",
            })
    selected = []
    seen_skus = set()
    for failure_class in FAILURE_CLASSES:
        count = 0
        for row in rows:
            if row["failure_class"] != failure_class or row["sku"] in seen_skus:
                continue
            selected.append(row)
            seen_skus.add(row["sku"])
            count += 1
            if count >= quotas[failure_class]:
                break
    fields = list(selected[0]) if selected else []
    with out_csv.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(selected)
    report = {
        "report_version": "stage4-targeted-remediation-review-queue-2026-09-12-v1",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "evidence_only": True,
        "source_master": {"path": str(MASTER), "sha256": sha(MASTER)},
        "frozen_test_excluded": len(frozen),
        "combined_split_excluded": len(split),
        "rows": len(selected),
        "requested_per_failure_class": args.per_failure_class,
        "requested_class_quotas": quotas,
        "by_failure_class": {kind: sum(r["failure_class"] == kind for r in selected) for kind in FAILURE_CLASSES},
        "all_human_confirmation_required": True,
        "training_eligible_rows": 0,
        "production_writes": False,
        "exclusion_policy": "PREDICTIONS + stage4 test-only + hard test + all historical incremental/combined/field-conditioned JSONL corpora",
        "csv": {"path": str(out_csv), "sha256": sha(out_csv)},
    }
    out_json.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report["json"] = {"path": str(out_json), "sha256": sha(out_json)}
    out_json.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"rows": len(selected), "by_failure_class": report["by_failure_class"], "training_eligible_rows": 0}, ensure_ascii=False))


if __name__ == "__main__":
    main()
