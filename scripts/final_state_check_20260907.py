from pathlib import Path
import sqlite3
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
db = sqlite3.connect(root / "runtime" / "db" / "action_tracker.db")
blank = db.execute("""select count(*) from products p left join product_localizations es
    on es.official_sku=p.official_sku and es.language='es'
    where p.status='CURRENT' and trim(coalesce(es.cat2,''))=''""").fetchone()[0]
total = db.execute("select count(*) from products where status='CURRENT'").fetchone()[0]
db.close()
print({"current": total, "blank_cat2": blank})
for name in ("20260907Action商品全量_西班牙语版_不带图.xlsx", "20260907Action商品全量_西班牙语版_带图.xlsx"):
    ws = load_workbook(root / "runtime" / "exports" / name, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    h = {str(v): i for i, v in enumerate(rows[0])}
    body = rows[1:]
    cat2 = next(i for k, i in h.items() if k == "分类2")
    image = next(i for k, i in h.items() if "图片链接" in k)
    product = next(i for k, i in h.items() if "商品链接" in k)
    print(name, {"rows": len(body), "blank_cat2": sum(not str(r[cat2] or '').strip() for r in body),
                 "image_http": sum(str(r[image] or '').startswith('http') for r in body),
                 "product_http": sum(str(r[product] or '').startswith('http') for r in body)})
