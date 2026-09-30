"""Validate an owner-completed Stage-4 remediation workbook without training.

Only rows whose original Spanish source still matches the audited record, whose
owner disposition is complete, and whose final Chinese values clear the
deterministic fact guard are emitted as Human-confirmed remediation Gold.
This script never updates Master, dictionaries, model weights, or split files.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from action_tracker.services.hashing import localization_source_hash
from action_tracker.config import load_settings
from action_tracker.exporting.dictionary_join import load_dictionary_context

RUN_DIR = ROOT / "runtime/training/qwen3_8b/20260911"
DEFAULT_BOOK = RUN_DIR / "stage4_remediation200_final_owner_review_20260912.xlsx"
AUDIT_CSV = RUN_DIR / "stage4_remediation200_final_model_approved_20260912.csv"
FIELDS = ("name", "cat1", "cat2", "spec", "description", "details")
OWNER_DECISIONS = {"ACCEPT_AS_GOLD", "REVISE", "REJECT", "SOURCE_CONFLICT"}
SOURCE_COLUMNS = {"name": "es_name", "cat1": "es_cat1", "cat2": "es_cat2", "spec": "es_spec", "description": "es_description", "details": "es_details"}
FINAL_COLUMNS = {"name": "建议中文品名", "cat1": "建议中文一级类目", "cat2": "建议中文二级类目", "spec": "建议中文规格", "description": "建议中文描述", "details": "建议中文产品详情"}
_REVIEW_SPEC = importlib.util.spec_from_file_location(
    "stage4_incremental_reviewer", ROOT / "scripts/review_qwen_incremental_with_deepseek.py",
)
assert _REVIEW_SPEC and _REVIEW_SPEC.loader
_REVIEWER = importlib.util.module_from_spec(_REVIEW_SPEC)
_REVIEW_SPEC.loader.exec_module(_REVIEWER)
_CONTEXT = None


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_review_rows(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Gold 审核"]
    rows = sheet.iter_rows(values_only=True)
    for _ in range(3):
        next(rows)
    headers = [str(value or "").strip() for value in next(rows)]
    result = []
    for values in rows:
        record = {header: str(value or "").strip() for header, value in zip(headers, values)}
        if record.get("SKU"):
            result.append(record)
    return result


def audited_source(row: dict[str, str]) -> dict[str, str]:
    return {field: str(row.get(column) or "").strip() for field, column in SOURCE_COLUMNS.items()}


def finalized_target(row: dict[str, str]) -> dict[str, str]:
    return {field: str(row.get(column) or "").strip() for field, column in FINAL_COLUMNS.items()}


def source_hash(source: dict[str, str]) -> str:
    return localization_source_hash({
        "name_es": source["name"], "cat1_es": source["cat1"], "cat2_es": source["cat2"],
        "spec_es": source["spec"], "desc_es": source["description"], "details_es": source["details"],
    })


def guard_reasons(source: dict[str, str], target: dict[str, str], sku: str) -> list[str]:
    """Use the identical brand/proper-name allowance used in model review."""
    global _CONTEXT
    if _CONTEXT is None:
        _CONTEXT = load_dictionary_context(load_settings(ROOT / "config/settings.yaml"))
    reasons, _exceptions = _REVIEWER._guard_with_safe_normalizations(
        source, target, _REVIEWER._brand_phrases(_CONTEXT, sku, source["name"]),
    )
    return reasons


def classify_row(review: dict[str, str], audited: dict[str, str] | None) -> tuple[str, list[str], dict[str, str] | None]:
    """Return status, explicit reasons, and a possible Gold target."""
    if audited is None:
        return "BLOCKED_UNKNOWN_SKU", ["SKU_NOT_IN_AUDITED_FINAL_200"], None
    if review.get("规则检查") != "APPROVED_MODEL_REVIEWED":
        return "BLOCKED_WORKBOOK_STATUS", ["WORKBOOK_RULE_STATUS_CHANGED"], None
    decision = review.get("人工决定", "").strip().upper()
    if decision not in OWNER_DECISIONS:
        return "PENDING_OWNER", ["OWNER_DECISION_MISSING"], None
    if decision in {"REJECT", "SOURCE_CONFLICT"}:
        return "OWNER_NOT_APPROVED", [decision], None
    if not review.get("审核人") or not review.get("审核日期"):
        return "PENDING_OWNER", ["OWNER_REVIEWER_OR_DATE_MISSING"], None
    source = audited_source(audited)
    review_source = {
        "name": review.get("西语品名", ""), "cat1": review.get("西语一级类目", ""), "cat2": review.get("西语二级类目", ""),
        "spec": review.get("西语规格", ""), "description": review.get("西语描述", ""), "details": review.get("西语产品详情", ""),
    }
    if source != review_source:
        return "BLOCKED_SOURCE_CHANGED", ["WORKBOOK_SOURCE_DOES_NOT_MATCH_AUDIT"], None
    target = finalized_target(review)
    if not all(target.values()):
        return "BLOCKED_FINAL_VALUE_EMPTY", ["FINAL_TARGET_FIELD_EMPTY"], None
    reasons = guard_reasons(source, target, review["SKU"])
    if reasons:
        return "BLOCKED_FACT_GUARD", reasons, None
    return "HUMAN_CONFIRMED_GOLD", [], target


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate owner decisions for Stage-4 remediation candidates")
    parser.add_argument("--workbook", default=str(DEFAULT_BOOK))
    parser.add_argument("--output-prefix", default=str(RUN_DIR / "stage4_remediation200_owner_validated"))
    args = parser.parse_args()
    workbook = Path(args.workbook)
    prefix = Path(args.output_prefix)
    audit = {str(row.get("sku") or ""): row for row in read_csv(AUDIT_CSV) if row.get("final_status") == "APPROVED_MODEL_REVIEWED"}
    review = read_review_rows(workbook)
    seen = set()
    results, gold = [], []
    for row in review:
        sku = row["SKU"]
        status, reasons, target = classify_row(row, audit.get(sku))
        if sku in seen:
            status, reasons, target = "BLOCKED_DUPLICATE_SKU", ["DUPLICATE_WORKBOOK_SKU"], None
        seen.add(sku)
        results.append({"sku": sku, "status": status, "reasons": reasons, "decision": row.get("人工决定", "")})
        if status == "HUMAN_CONFIRMED_GOLD" and target is not None:
            source = audited_source(audit[sku])
            gold.append({
                "messages": [
                    {"role": "system", "content": "将 Action 西语商品字段标准化为简洁、忠实的中文。保持数字、单位、数量、尺寸和品牌/型号事实；不得臆造源数据没有的信息。输出固定六字段 JSON。"},
                    {"role": "user", "content": json.dumps(source, ensure_ascii=False, sort_keys=True)},
                    {"role": "assistant", "content": json.dumps(target, ensure_ascii=False, sort_keys=True)},
                ],
                "metadata": {
                    "sku": sku, "source_hash": source_hash(source), "gold_tier": "HUMAN_CONFIRMED_REMEDIATION_GOLD",
                    "owner_decision": row["人工决定"], "owner_reviewer": row["审核人"], "owner_reviewed_at": row["审核日期"],
                    "training_eligible": False,
                    "training_note": "Requires a separate leakage/split validation before any targeted retrain.",
                },
            })
    expected = set(audit)
    missing = sorted(expected - seen)
    if missing:
        results.extend({"sku": sku, "status": "BLOCKED_WORKBOOK_SKU_MISSING", "reasons": ["APPROVED_AUDIT_ROW_MISSING"], "decision": ""} for sku in missing)
    gold_path = prefix.with_name(prefix.name + "_gold.jsonl")
    manifest_path = prefix.with_name(prefix.name + "_manifest.json")
    gold_path.write_text("\n".join(json.dumps(row, ensure_ascii=False, sort_keys=True) for row in gold) + ("\n" if gold else ""), encoding="utf-8")
    counts = Counter(row["status"] for row in results)
    manifest = {
        "status": "OWNER_REVIEW_VALIDATED",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "workbook": {"path": str(workbook), "sha256": sha(workbook)},
        "audited_eligible_rows": len(audit),
        "workbook_rows": len(review),
        "status_counts": dict(sorted(counts.items())),
        "approved_gold_rows": len(gold),
        "all_owner_decisions_complete": not any(row["status"] == "PENDING_OWNER" for row in results),
        "source_integrity_pass": not any(row["status"] == "BLOCKED_SOURCE_CHANGED" for row in results),
        "gold_jsonl": {"path": str(gold_path), "sha256": sha(gold_path)},
        "row_results": results,
        "training_authorized": False,
        "production_writes": False,
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"gold": len(gold), "statuses": dict(sorted(counts.items())), "manifest": str(manifest_path)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
