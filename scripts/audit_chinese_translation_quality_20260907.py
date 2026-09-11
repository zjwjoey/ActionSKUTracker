"""Full quality audit for the 2026-09-07 Chinese export.

This is a read-only audit. It does not modify the Spanish or Chinese source
workbooks. The report separates brand-policy candidates from hard translation
and detail-structure defects.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

EXPORTS = Path(r"F:\ActionSKUTracker\runtime\exports")
ES_PATH = EXPORTS / "20260907Action商品全量_西班牙语版_不带图_最终修复版_分类清理版.xlsx"
ZH_PATH = EXPORTS / "20260907Action商品全量_中文版_GPT56Luna最终修复版_不带图.xlsx"
REPORT_JSON = EXPORTS / "20260907Action商品全量_中文翻译质量全量审查_20260907.json"
REPORT_XLSX = EXPORTS / "20260907Action商品全量_中文翻译质量全量审查_20260907.xlsx"


def load_rows(path: Path) -> dict[str, tuple]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = {str(r[1]).strip(): r for r in ws.iter_rows(min_row=2, values_only=True)}
    wb.close()
    return rows


def split_pairs(value: object) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for part in re.split(r";|；|\n", str(value or "")):
        part = part.strip()
        if not part:
            continue
        match = re.search(r"[:：\t]", part)
        if not match:
            pairs.append((part, ""))
            continue
        key, val = re.split(r"[:：\t]", part, maxsplit=1)
        pairs.append((key.strip(), val.strip()))
    return pairs


SPANISH = re.compile(
    r"\b(?:el|la|los|las|de|del|con|para|sin|una|un|unos|unas|varios|varias|"
    r"diferentes|talla|tallas|unidades?|piezas?|gramos?|metros?|litros?|paquete|"
    r"juego|calcetines|mallas|pantis|guantes|gorro|chocolate|cable|bolsa|discos|"
    r"gel|colores?|hogar|oficina|papelería|mascotas|juguetes|manualidades|limpieza|"
    r"deporte|número|tipo|color|material|cantidad|incluye|sí|no|sabor|cremoso|"
    r"picante|rápido|fácil|producto|cereal|arroz|azúcares|grasas|proteínas|sal|"
    r"energía|contenido|conservación|azúcar|saturados|carbohidratos|aromatizantes|"
    r"ingredientes|añadidos|pasta|fideos|rellenos|fuentes|calor|reciclado|"
    r"transparente|pizarra|adaptador|dispositivos|vez|estuche|uso|fuentes?)\b",
    re.I,
)


def numeric_tokens(value: object) -> set[str]:
    return set(re.findall(r"(?<![A-Za-z])\d+(?:[.,]\d+)?", str(value or "")))


def main() -> None:
    es = load_rows(ES_PATH)
    zh = load_rows(ZH_PATH)

    brand_marker: list[dict] = []
    title_untranslated: list[dict] = []
    description_residual: list[dict] = []
    detail_malformed: list[dict] = []
    detail_key_deficit: list[dict] = []
    detail_numeric_mismatch: list[dict] = []

    for sku, erow in es.items():
        zrow = zh.get(sku)
        if zrow is None:
            continue
        title = str(zrow[2] or "").strip()
        desc = str(zrow[9] or "").strip()
        detail = str(zrow[10] or "").strip()
        if "牌" in title and not any(x in title for x in ("扑克牌", "行李牌", "名牌")):
            brand_marker.append({"sku": sku, "title_zh": title, "title_es": erow[2]})
        if title and not re.search(r"[\u3400-\u9fff]", title):
            title_untranslated.append({"sku": sku, "title_zh": title, "title_es": erow[2]})
        if SPANISH.search(desc) or SPANISH.search(detail):
            description_residual.append({
                "sku": sku,
                "field": "描述/产品详情",
                "value_zh": (desc if SPANISH.search(desc) else detail),
            })

        es_pairs = split_pairs(erow[10])
        zh_pairs = split_pairs(detail)
        source_unique_keys = {k for k, _ in es_pairs if k}
        if len(zh_pairs) < len(source_unique_keys):
            detail_key_deficit.append({
                "sku": sku,
                "source_pair_count": len(es_pairs),
                "source_unique_key_count": len(source_unique_keys),
                "zh_pair_count": len(zh_pairs),
                "detail_es": erow[10],
                "detail_zh": detail,
            })
        malformed_entries = [
            f"{key}:{value}" for key, value in zh_pairs if not key or not value
        ]
        malformed_key_text = any(
            not key
            or not value
            or key.startswith(("/", ":"))
            or " 的 " in key
            or "用于" in key
            or "类型 的" in key
            for key, value in zh_pairs
        )
        if malformed_key_text:
            detail_malformed.append({
                "sku": sku,
                "malformed_entries": " | ".join(malformed_entries),
                "detail_es": erow[10],
                "detail_zh": detail,
            })

        source_nums = numeric_tokens(erow[10])
        zh_nums = numeric_tokens(detail)
        missing = sorted(n for n in source_nums if n not in zh_nums and n.replace(".", ",") not in zh_nums)
        if missing:
            detail_numeric_mismatch.append({
                "sku": sku,
                "missing_numeric_tokens": ", ".join(missing),
                "detail_es": erow[10],
                "detail_zh": detail,
            })

    summary = {
        "sku_count": len(es),
        "brand_marker_candidates": len(brand_marker),
        "titles_without_chinese": len(title_untranslated),
        "description_or_detail_spanish_residual": len(description_residual),
        "detail_key_count_deficits": len(detail_key_deficit),
        "malformed_detail_rows": len(detail_malformed),
        "detail_numeric_mismatch_rows": len(detail_numeric_mismatch),
        "source_empty_spec_rows": sum(not str(es[s][5] or "").strip() for s in es),
        "note": "Brand candidates are not automatically deleted; confirm against the current no-brand naming rule before applying.",
    }
    report = {
        "source_es": str(ES_PATH),
        "source_zh": str(ZH_PATH),
        "summary": summary,
        "brand_marker_candidates": brand_marker,
        "titles_without_chinese": title_untranslated,
        "description_or_detail_spanish_residual": description_residual,
        "detail_key_count_deficits": detail_key_deficit,
        "malformed_detail_rows": detail_malformed,
        "detail_numeric_mismatch_rows": detail_numeric_mismatch,
    }
    REPORT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["检查项目", "数量", "说明"])
    for key, value in summary.items():
        ws.append([key, value, "详见对应工作表" if key not in {"sku_count", "source_empty_spec_rows", "note"} else ""])
    sheets = {
        "BrandCandidates": (brand_marker, ["sku", "title_zh", "title_es"]),
        "TitlesNoChinese": (title_untranslated, ["sku", "title_zh", "title_es"]),
        "SpanishResidual": (description_residual, ["sku", "field", "value_zh"]),
        "DetailMalformed": (detail_malformed, ["sku", "malformed_entries", "detail_es", "detail_zh"]),
        "DetailKeyDeficit": (detail_key_deficit, ["sku", "source_pair_count", "source_unique_key_count", "zh_pair_count", "detail_es", "detail_zh"]),
        "DetailNumericMismatch": (detail_numeric_mismatch, ["sku", "missing_numeric_tokens", "detail_es", "detail_zh"]),
    }
    for name, (items, columns) in sheets.items():
        sh = wb.create_sheet(name)
        sh.append(columns)
        for item in items:
            sh.append([item.get(c, "") for c in columns])
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for cell in sh[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for col in sh.columns:
            letter = col[0].column_letter
            sh.column_dimensions[letter].width = min(80, max(14, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(REPORT_XLSX)
    print(json.dumps({"summary": summary, "json": str(REPORT_JSON), "xlsx": str(REPORT_XLSX)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
