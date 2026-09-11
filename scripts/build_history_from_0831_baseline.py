"""Build a bounded Action presence workbook from the user's 2026-08-31 baseline.

The supplied 08/31 workbook is a headerless 16-column matrix.  This script
keeps its rows and presence values unchanged, deliberately omits the April
snapshot, and appends only CURRENT SKUs from the current Master that are not
in the supplied baseline.  It does not alter Master or project history config.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASELINE = Path(r"D:\Users\Administrator\Desktop\Action商品上下架明细 26.08.31.xlsx")
MASTER = Path(r"F:\ActionSKUTracker\runtime\master\Action_Master.xlsx")
EXPORT_DIR = Path(r"F:\ActionSKUTracker\runtime\exports")
DATE_DIR = Path(r"F:\按日期整理\20260831")
FILE_NAME = "20260831_Action商品上下架明细.xlsx"
DATES = (
    "26.01.08", "26.06.28", "26.07.05", "26.07.12", "26.07.19",
    "26.07.26", "26.08.02", "26.08.10", "26.08.17", "26.08.24",
    "26.08.31",
)
HEADERS = ("序号", "编号", "中文品名", "图片链接", "商品链接", *DATES)


def sku_key(value: str) -> tuple[int, str]:
    text = str(value).strip()
    return (int(text), text) if text.isdigit() else (10**18, text)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def load_baseline() -> tuple[list[list[Any]], set[str]]:
    if not BASELINE.exists():
        raise FileNotFoundError(BASELINE)
    workbook = openpyxl.load_workbook(BASELINE, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not rows or any(len(row) != len(HEADERS) for row in rows):
        raise ValueError("BASELINE_SCHEMA_MISMATCH")
    skus = [str(row[1] or "").strip() for row in rows]
    if any(not sku for sku in skus):
        raise ValueError("BASELINE_EMPTY_SKU")
    if len(skus) != len(set(skus)):
        raise ValueError("BASELINE_DUPLICATE_SKU")
    if any(row[0] != index for index, row in enumerate(rows, start=1)):
        raise ValueError("BASELINE_SEQUENCE_MISMATCH")
    for row in rows:
        if any(row[5 + index] not in (0, 1) for index in range(len(DATES))):
            raise ValueError(f"BASELINE_BAD_PRESENCE:{row[1]}")
    return rows, set(skus)


def load_current_master(baseline_skus: set[str]) -> tuple[list[list[Any]], dict[str, int]]:
    if not MASTER.exists():
        raise FileNotFoundError(MASTER)
    workbook = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)
    try:
        sheet = workbook["01_SKU_ZH_CURRENT"]
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        index = {str(value): pos for pos, value in enumerate(header)}
        required = {"SKU", "中文品名", "图片链接", "商品链接", "当前状态"}
        if not required.issubset(index):
            raise ValueError("MASTER_SCHEMA_MISMATCH")
        rows: list[list[Any]] = []
        for raw in sheet.iter_rows(min_row=2, values_only=True):
            sku = str(raw[index["SKU"]] or "").strip()
            if not sku or sku in baseline_skus:
                continue
            if str(raw[index["当前状态"]] or "").strip().upper() != "CURRENT":
                continue
            rows.append([
                None, sku, raw[index["中文品名"]] or "", raw[index["图片链接"]] or "",
                raw[index["商品链接"]] or "", *([0] * len(DATES)),
            ])
    finally:
        workbook.close()
    rows.sort(key=lambda row: sku_key(row[1]))
    return rows, index


def write_workbook(path: Path, rows: list[list[Any]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "商品上下架明细"
    sheet.append(list(HEADERS))
    for number, row in enumerate(rows, start=1):
        row[0] = number
        sheet.append(row)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{sheet.max_row}"
    widths = {"A": 8, "B": 14, "C": 30, "D": 65, "E": 65}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for col in range(6, len(HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 11
        for row_no in range(2, sheet.max_row + 1):
            sheet.cell(row=row_no, column=col).number_format = "0"
    for row_no in range(2, sheet.max_row + 1):
        for col in (3, 4, 5):
            sheet.cell(row=row_no, column=col).alignment = Alignment(vertical="top", wrap_text=True)
        for col in (4, 5):
            value = sheet.cell(row=row_no, column=col).value
            if isinstance(value, str) and value.startswith(("http://", "https://")):
                sheet.cell(row=row_no, column=col).hyperlink = value
                sheet.cell(row=row_no, column=col).style = "Hyperlink"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.stem}.tmp.xlsx")
    workbook.save(temporary)
    workbook.close()
    temporary.replace(path)


def main() -> None:
    base_rows, base_skus = load_baseline()
    new_rows, _ = load_current_master(base_skus)
    rows = base_rows + new_rows
    output = EXPORT_DIR / FILE_NAME
    mirror = DATE_DIR / FILE_NAME
    write_workbook(output, rows)
    write_workbook(mirror, rows)
    audit = {
        "source_baseline": str(BASELINE),
        "source_master": str(MASTER),
        "baseline_sha256": sha256(BASELINE),
        "master_sha256": sha256(MASTER),
        "excluded_date_columns": ["26.04.05"],
        "date_columns": list(DATES),
        "baseline_sku_count": len(base_rows),
        "new_current_sku_count": len(new_rows),
        "output_sku_count": len(rows),
        "output_duplicate_sku_count": len(rows) - len({str(row[1]) for row in rows}),
        "new_sku_presence_policy": "0 for every baseline date; baseline is treated as complete through 26.08.31",
        "new_skus": [row[1] for row in new_rows],
        "output": str(output),
        "mirror_output": str(mirror),
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "qc_result": "PASS" if len(rows) == len({str(row[1]) for row in rows}) else "FAIL",
    }
    (EXPORT_DIR / f"{Path(FILE_NAME).stem}.audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(audit, ensure_ascii=False))


if __name__ == "__main__":
    main()
