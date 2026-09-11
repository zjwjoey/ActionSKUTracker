from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path("reports/Action_SKU_Tracker_Project_Status_2026-08-11.docx")

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell(cell, text, bold=False):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); r=p.add_run(str(text)); r.bold=bold; r.font.name='Calibri'; r.font.size=Pt(9)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; t.autofit=False
    for i,h in enumerate(headers):
        set_cell(t.rows[0].cells[i],h,True); shade(t.rows[0].cells[i], 'D9EAF7')
        if widths: t.rows[0].cells[i].width=Inches(widths[i])
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            set_cell(cells[i],v)
            if widths: cells[i].width=Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after=Pt(4)
    return t

def h(doc, text, level=1):
    p=doc.add_heading(text, level=level); p.paragraph_format.space_before=Pt(14); p.paragraph_format.space_after=Pt(6); return p

def bullet(doc,text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.add_run(text)

def main():
    OUT.parent.mkdir(exist_ok=True)
    d=Document(); s=d.sections[0]; s.top_margin=s.bottom_margin=Inches(0.8); s.left_margin=s.right_margin=Inches(0.85)
    normal=d.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(10); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.12
    title=d.add_paragraph(); title.paragraph_format.space_after=Pt(3); r=title.add_run('Action SKU Tracker - 项目现状与讨论报告'); r.bold=True; r.font.size=Pt(23); r.font.color.rgb=RGBColor(31,78,120)
    sub=d.add_paragraph('版本：2026-08-11  |  面向：技术评审 / 其他 AI 讨论  |  项目目录：F:\\ActionSKUTracker')
    sub.runs[0].italic=True; sub.runs[0].font.color.rgb=RGBColor(80,80,80)
    h(d,'一、结论摘要')
    d.add_paragraph('项目是一个本地化的 Action 西班牙官网商品监测系统。目标是依据 SKU Presence 证据做生命周期动作判断，而不是追求固定 SKU 总数。当前生命周期、QA 门禁和 dry-run 保护已稳定；本轮新增动态类目发现、站点结构快照、单实例锁、Madrid 业务日期与 Playwright 访问熔断。')
    table(d,['结论项','当前判断','说明'],[
        ['生命周期语义','稳定','known_skus.csv 是唯一跨日状态源；REAPPEARED 依赖前一状态 MISSING/OFFLINE。'],
        ['正式数据安全','稳定','dry-run 与 QA FAIL/BLOCKED 均不修改 Master、known_skus、offline_skus。'],
        ['采集可靠性','需继续观察','本次真实 dry-run 在详情页遭遇 Cloudflare，已正确 BLOCKED。'],
        ['正式运行许可','暂不建议','需在无 challenge 的完整 dry-run 中取得 QA PASS 后，再人工批准正式运行。'],
    ],[1.25,1.25,4.0])
    h(d,'二、系统架构与正式数据边界')
    d.add_paragraph('采集层（sitemap、一级类目 listing、Nuevo、Promoción）→ Presence / Coverage → SKU Monitor → Lifecycle Engine → Product Updater / Change Detector → Snapshot + Staging → QA Gate → 文件提交。')
    bullet(d,'正式文件：runtime/master/Action_Master.xlsx、runtime/state/known_skus.csv。offline_skus.csv 只从 known_skus[last_status=OFFLINE] 派生。')
    bullet(d,'冻结项目：src/action_tracker/database/ 仅保留脚手架；daily-run 不读写 SQLite。')
    bullet(d,'翻译：无供应商、无 API；保留已有中文，缺失中文回退西班牙语并标记 FALLBACK_ES / NOT_CONFIGURED。')
    h(d,'三、生命周期与证据规则')
    table(d,['状态/概念','规则','安全约束'],[
        ['NEW / FIRST_SEEN','今日存在，历史 known_skus 未见','与 REAPPEARED 互斥。'],
        ['ACTIVE','今日存在且非重新出现','类别移动或 URL 变化不改变 SKU 身份。'],
        ['MISSING_FIRST / CONTINUED','仅在有效 sitemap 或相关类别完整覆盖时推进','同日重复运行不再 +1。'],
        ['OFFLINE','连续有效缺失达到配置阈值','默认 3 次；offline_skus 为派生视图。'],
        ['REAPPEARED','今日存在、历史已知、前一状态为 MISSING 或 OFFLINE','清零 missing_count，记录事件。'],
        ['UNKNOWN','观测/解析/覆盖无效','绝不推进 missing_count，也不下架。'],
    ],[1.45,3.35,1.7])
    h(d,'四、当前采集实现与本轮可靠性加固')
    table(d,['能力','现状','备注'],[
        ['Sitemap','XML <loc> 容错解析，URL 正则提取 SKU','不依赖固定 URL 片段索引。'],
        ['类目发现','动态导航中的 /c/<slug>/ 为主；settings.yaml 为 fallback','不再把“15 类目”作为业务规则。'],
        ['分页','动态读取 GridPaginationLink','不硬编码页数。'],
        ['卡片/价格/标签解析','selector 集中于 monitor/listing.py、products/parser.py','不会泄漏至生命周期或 Excel 业务层。'],
        ['解析健康','扫描后 0 商品不视为可信空类目','防止 selector 失效制造批量下架。'],
        ['特殊页','Nuevo / Promoción 为辅助 Presence 与徽章来源','特殊页失败不应否定 sitemap + 主类目。'],
        ['访问控制','单 Playwright browser/context/page；NORMAL/DEGRADED/COOLDOWN/PROBE/BLOCKED','本轮明确不做裸 HTTP 4-6 并发或多浏览器并发。'],
        ['单实例','runtime/state/daily-run.lock','阻止 dry-run 与正式 run 同时采集；可回收陈旧锁。'],
        ['业务日期','Europe/Madrid','避免台湾本机跨日错误推进生命周期。'],
    ],[1.25,2.25,3.0])
    h(d,'五、Snapshot、Staging 与运行清单')
    d.add_paragraph('每次运行使用同一个 run_id：runtime/snapshots/YYYY-MM-DD/<run_id>/ 和 runtime/staging/<run_id>/。Snapshot 保存原始 sitemap/listing、presence_evidence、coverage、site_structure、category_structure、QA、run_report、run_manifest；Staging 保存生命周期、产品、价格、翻译与事件改动。run_manifest 包含 Madrid observation_date、启动时间、Git 提交、工作区是否 dirty、配置 hash 与访问控制状态。')
    h(d,'六、本次真实 dry-run 证据（2026-08-11）')
    table(d,['项目','结果'],[
        ['Run ID','2026-08-11_034523'],['模式','dry-run；单页 Playwright'],['Listing / Sitemap','listing 5,530；sitemap 本轮无法作为有效主证据'],['主类目覆盖','15 个主类目均完成并标记完整'],['候选生命周期','ACTIVE 5,527；NEW 2；REAPPEARED 1；MISSING_FIRST 10；MISSING_CONTINUED 3；OFFLINE 0；UNKNOWN 0'],['异常','Nuevo 第 5 页网络断连后恢复；详情补抓出现 Cloudflare challenge'],['最终 QA','BLOCKED；observation_complete=false'],['正式文件','Master、known_skus、offline_skus SHA-256 均未变化'],
    ],[2.0,4.5])
    h(d,'七、测试与提交')
    table(d,['项','结果'],[
        ['基线 Git','f1a8b30 - 修复 REAPPEARED 以前一状态为依据'],['本轮 Git','fbdc010 - 动态发现、运行安全与访问控制'],['自动测试','90 passed，0 failed'],['新增覆盖','动态发现/fallback、403/429 cooldown+probe、单实例锁、陈旧锁回收'],
    ],[1.5,5.0])
    h(d,'八、已知问题与建议讨论')
    bullet(d,'本次 dry-run 被 Cloudflare challenge 阻断，保护逻辑正确，但尚未得到“完整且 QA PASS”的新版本真实验证。')
    bullet(d,'详情抓取路径仍需进一步统一接入同一 AccessController，确保挑战发生时立即停止后续详情任务，而不仅由 QA 最终阻断。')
    bullet(d,'动态类目发现当前从公开导航链接提取一级 /c/<slug>/；建议讨论是否需增加更稳定的导航容器选择器或辅助结构来源。')
    bullet(d,'文件提交已有 staging + 原子替换，但 Master 与 state 的跨文件原子性只能报告 PARTIAL_COMMIT；可讨论是否需要未来的 commit-run 重试命令。')
    h(d,'九、讨论问题（建议直接转发）')
    for q in ['动态一级类目发现的导航链接策略是否足够稳健，是否需要 sitemap/页面 JSON 作为第二来源？','Cloudflare 挑战后，详情抓取是否应立即退出整个 run，还是保留 listing 证据但统一 QA BLOCKED？','是否需要实现 commit-run <run_id>，以便文件被 Excel 占用后仅重试提交、不重新抓取官网？','是否需要为 category rename/path change 建立跨 run 结构差异分类，而不进入商品 EVENT_HISTORY？','正式运行的最小验收条件是否应为：连续一次或多次无 challenge 的 dry-run QA PASS？']:
        bullet(d,q)
    h(d,'十、当前准备度')
    table(d,['指标','状态'],[
        ['SITE_STRUCTURE_READY','YES - 动态发现 + fallback + 快照已具备。'],
        ['COLLECTION_RESILIENCE_READY','PARTIAL - 熔断/锁/QA 阻断已具备；需一次无挑战成功运行验证。'],
        ['READY_FOR_FORMAL_RUN','NO - 等待人工确认及成功真实 dry-run。'],
    ],[2.2,4.3])
    d.save(OUT)

if __name__ == '__main__': main()
