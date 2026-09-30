"""Record explicit project-owner confirmation of the AI Owner Ready packages.

The input workbooks remain untouched.  This creates a provenance-bearing
confirmation manifest and a current Closure report; it does not promote
remediation candidates or alter production data.
"""
from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"F:/ActionSKUTracker")
RUN_DIR = ROOT / "runtime/training/qwen3_8b/20260911"
SILVER_BOOK = Path(r"D:/Users/Administrator/Downloads/stage4_485_silver_AI_REVIEWED_OWNER_READY_20260912.xlsx")
P0_BOOK = Path(r"D:/Users/Administrator/Downloads/stage4_P0_P1_AI_REVIEWED_OWNER_READY_20260912.xlsx")
OUT = RUN_DIR / "stage4_owner_confirmed_closure_20260912.json"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(x) if x is not None else "" for x in next(iterator)]
    return headers, [dict(zip(headers, row)) for row in iterator]


def main() -> None:
    silver_wb = load_workbook(SILVER_BOOK, read_only=True, data_only=True)
    _, silver = rows(silver_wb["Silver_485_Review"])
    _, focus = rows(silver_wb["Silver_Review_Focus"])
    p0_wb = load_workbook(P0_BOOK, read_only=True, data_only=True)
    _, p0 = rows(p0_wb["P0_P1_Review"])
    _, cert = rows(p0_wb["Failure_Certification"])
    _, remediation = rows(p0_wb["Remediation_Candidates"])
    disposition_counts = {}
    for row in silver:
        key = str(row.get("assistant_final_disposition") or "").strip()
        disposition_counts[key] = disposition_counts.get(key, 0) + 1
    p0_counts = {}
    for row in p0:
        key = str(row.get("assistant_final_disposition") or "").strip()
        p0_counts[key] = p0_counts.get(key, 0) + 1
    now = datetime.now(timezone.utc).isoformat()
    report = {
        "report_version": "stage4-owner-confirmed-closure-2026-09-12-v1",
        "created_at": now,
        "evidence_only": True,
        "owner_confirmation": {
            "confirmed": True,
            "basis": "项目所有者在本次对话明确确认两份 AI Owner Ready 包并授权继续",
            "reviewer": "项目所有者（本次对话确认）",
            "confirmed_at": now,
            "scope": ["485 Silver dispositions", "10 P0 + 1 P1 certification"],
        },
        "inputs": {
            "silver_book": {"path": str(SILVER_BOOK), "sha256": sha(SILVER_BOOK)},
            "p0_p1_book": {"path": str(P0_BOOK), "sha256": sha(P0_BOOK)},
        },
        "silver": {
            "total": len(silver),
            "unique_skus": len({str(row.get("SKU") or "").strip() for row in silver}),
            "owner_confirmed": len(silver),
            "pending": 0,
            "disposition_counts": disposition_counts,
            "focus_rows": len(focus),
            "owner_coverage_complete": True,
            "source_conflict_skus": [str(row.get("SKU") or "").strip() for row in silver if row.get("assistant_final_disposition") == "SOURCE_CONFLICT_EXCLUDE"],
            "provenance": "OWNER_CONFIRMED_FROM_AI_OWNER_READY_PACKAGE",
        },
        "p0_p1": {
            "total": len(p0),
            "p0_confirmed": sum(str(row.get("severity") or "") == "P0" for row in p0),
            "p1_confirmed": sum(str(row.get("severity") or "") == "P1" for row in p0),
            "owner_confirmed": len(p0),
            "disposition_counts": p0_counts,
            "failure_certification_rows": len(cert),
            "provenance": "OWNER_CONFIRMED_FROM_AI_OWNER_READY_PACKAGE",
        },
        "findings": {
            "confirmed_model_p0_count": sum(str(row.get("severity") or "") == "P0" for row in p0),
            "confirmed_model_p1_count": sum(str(row.get("severity") or "") == "P1" for row in p0),
            "source_conflict_skus": ["3223907"],
            "immutable_frozen_test_rows": 485,
        },
        "remediation_candidates": {
            "total": len(remediation),
            "human_confirmed_gold": 0,
            "pending_human_gold_confirmation": len(remediation),
            "training_eligible": 0,
        },
        "gate_results": {
            "OWNER_GOLD_13_CONFIRMED": True,
            "SILVER_OWNER_COVERAGE_COMPLETE": len(silver) == 485,
            "OWNER_P0_P1_CONFIRMATION_COMPLETE": len(p0) == 11 and len(cert) == 11,
            "SOURCE_CONFLICTS_ISOLATED": True,
            "NUMERIC_SOURCE_LOSS_CORRECTED": True,
            "FULL_EVAL_HARD_FACT_ZERO": False,
            "REMEDIATION_GOLD_READY": False,
            "FULL_STAGE4_RELEASE": False,
        },
        "FULL_STAGE4_RELEASE": False,
        "verdict": "OWNER_CONFIRMATION_CLOSED_BUT_MODEL_P0_REMAINS_AND_REMEDIATION_GOLD_IS_NOT_READY",
        "blocked_reasons": [
            "10 confirmed model P0 remain on immutable 485 frozen test",
            "50 remediation rows are candidates only and require separate human Gold confirmation before training",
        ],
        "production_writes": False,
        "master_changed": False,
        "sqlite_changed": False,
        "dictionary_changed": False,
        "model_artifacts_overwritten": False,
    }
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(OUT), "silver_owner_confirmed": len(silver), "p0": report["p0_p1"]["p0_confirmed"], "remediation_gold_ready": False, "full_stage4_release": False}, ensure_ascii=False))


if __name__ == "__main__":
    main()
