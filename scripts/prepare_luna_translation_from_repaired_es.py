"""Prepare remaining Chinese localization work from the repaired Spanish export."""
from __future__ import annotations
import json,re,sys
from pathlib import Path
if hasattr(sys.stdout,'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')
PROJECT=Path(__file__).resolve().parents[1]; sys.path.insert(0,str(PROJECT/'src'))
from openpyxl import load_workbook
from action_tracker.config import load_settings
CJK=re.compile(r'[\u3400-\u9fff]')
SP=re.compile(r'\b(?:el|la|los|las|de|del|con|para|sin|una|un|unos|unas|varios|varias|diferentes|talla|tallas|unidades?|piezas?|gramos?|metros?|litros?|paquete|juego|calcetines|mallas|pantis|guantes|gorro|chocolate|cable|bolsa|discos|gel|colores?|hogar|oficina|papelería|mascotas|juguetes|manualidades|limpieza|deporte|número|tipo|color|material|cantidad|incluye|sí|no)\b',re.I)
FIELDS=('name','cat1','cat2','spec','description','details')
ROW={'name':'标题','cat1':'分类1','cat2':'分类2','spec':'规格','description':'描述','details':'产品详情'}
SRC={'name':'标题','cat1':'分类1','cat2':'分类2','spec':'规格','description':'描述','details':'产品详情'}
ROOT=Path(r'F:\ActionSKUTracker\runtime')
ES=ROOT/'exports'/'20260907Action商品全量_西班牙语版_不带图_QC修复版.xlsx'
OLD=ROOT/'exports'/'20260907Action商品全量_西班牙语版_不带图.xlsx'
ZH=ROOT/'exports'/'20260907Action商品全量_中文版_GPT56Luna优化版_不带图.xlsx'
def t(v): return '' if v is None else str(v).strip()
def load(p):
    ws=load_workbook(p,read_only=True,data_only=True).active; h=[c.value for c in next(ws.iter_rows())]; return {t(r[1].value):dict(zip(h,[c.value for c in r])) for r in ws.iter_rows(min_row=2)}
def norm_detail(v):
    x=t(v); x=re.sub(r'^Especificaciones\s*','',x,flags=re.I).replace('|',';'); return re.sub(r'\s+',' ',x.replace('\t',': ')).strip()
def norm_desc(v): return re.sub(r'^\s*null\.\s*','',t(v),flags=re.I)
def source_changed(f,new,old):
    if f=='details': return norm_detail(new)!=norm_detail(old)
    if f=='description': return norm_desc(new)!=norm_desc(old)
    return t(new)!=t(old)
def needed(f,current,source,changed):
    if not source and f!='spec': return False
    if f=='spec' and not source: return True
    if changed: return True
    if not current or current==source: return True
    if f in {'cat1','cat2','name'}: return not CJK.search(current) or bool(SP.search(current))
    if f=='spec': return bool(SP.search(current))
    return not CJK.search(current) or bool(SP.search(current))
def main():
    es,old,zh=load(ES),load(OLD),load(ZH); rows=[]
    for sku in sorted(es,key=lambda x:(int(x) if x.isdigit() else 10**18,x)):
        needed_fields=[]
        for f in FIELDS:
            src=t(es[sku].get(SRC[f])); oldsrc=t(old[sku].get(SRC[f])); cur=t(zh[sku].get(ROW[f])); changed=source_changed(f,src,oldsrc)
            if needed(f,cur,src,changed): needed_fields.append(f)
        if needed_fields:
            rows.append({'sku':sku,'needed':needed_fields,'source':{f:t(es[sku].get(SRC[f])) for f in FIELDS},'current':{f:t(zh[sku].get(ROW[f])) for f in FIELDS},'old_source':{f:t(old[sku].get(SRC[f])) for f in FIELDS}})
    cfg=load_settings(); out=Path(cfg['paths']['temp'])/'gpt56_luna_repaired_inputs'; out.mkdir(parents=True,exist_ok=True); size=(len(rows)+2)//3
    for i in range(3): (out/f'batch_{i+1}.json').write_text(json.dumps(rows[i*size:(i+1)*size],ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'candidates':len(rows),'sizes':[len(rows[i*size:(i+1)*size]) for i in range(3)],'directory':str(out)},ensure_ascii=False))
if __name__=='__main__': main()
