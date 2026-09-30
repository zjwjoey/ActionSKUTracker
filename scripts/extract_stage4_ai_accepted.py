"""Extract AI-accepted Stage-4 rows as owner-pending candidate data.

The uploaded workbook is an external review artifact.  This script copies only
rows whose AI disposition is ACCEPT_AS_GOLD, preserving all six source and
target fields.  It does not promote rows to Gold, alter a split, or write
production data.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

import sys
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))
from action_tracker.services.hashing import localization_source_hash


FIELDS = ("name", "cat1", "cat2", "spec", "description", "details")
SOURCE_COLUMNS = {
    "name": "西语品名", "cat1": "西语一级类目", "cat2": "西语二级类目",
    "spec": "西语规格", "description": "西语描述", "details": "西语产品详情",
}
TARGET_COLUMNS = {
    "name": "建议中文品名", "cat1": "建议中文一级类目", "cat2": "建议中文二级类目",
    "spec": "建议中文规格", "description": "建议中文描述", "details": "建议中文产品详情",
}
AI_DISPOSITION = "AI审核结论"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def as_text(value: object) -> str:
    return str(value or "").strip()


def read_rows(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Gold 审核"]
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(i for i, row in enumerate(rows) if any(as_text(v) == "SKU" for v in row))
    headers = [as_text(v) for v in rows[header_index]]
    index = {header: i for i, header in enumerate(headers)}
    required = {"SKU", AI_DISPOSITION, *SOURCE_COLUMNS.values(), *TARGET_COLUMNS.values(), "AI审核理由"}
    missing = required - set(index)
    if missing:
        raise ValueError(f"WORKBOOK_COLUMNS_MISSING:{sorted(missing)}")
    result = []
    for values in rows[header_index + 1 :]:
        sku = as_text(values[index["SKU"]])
        if not sku:
            continue
        result.append({header: as_text(values[index[header]]) for header in required})
    return result


def extract(rows: list[dict[str, str]]) -> list[dict[str, object]]:
    accepted = []
    for row in rows:
        if row[AI_DISPOSITION] != "ACCEPT_AS_GOLD":
            continue
        source = {field: row[column] for field, column in SOURCE_COLUMNS.items()}
        target = {field: row[column] for field, column in TARGET_COLUMNS.items()}
        if not all(source.values()) or not all(target.values()):
            raise ValueError(f"ACCEPTED_ROW_INCOMPLETE:{row['SKU']}")
        accepted.append({
            "sku": row["SKU"],
            "source": source,
            "target": target,
            "ai_reason": row["AI审核理由"],
            "source_hash": localization_source_hash({
                "name_es": source["name"], "cat1_es": source["cat1"], "cat2_es": source["cat2"],
                "spec_es": source["spec"], "desc_es": source["description"], "details_es": source["details"],
            }),
        })
    return accepted


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract AI-accepted rows without Gold promotion")
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--manifest", required=True)
    args = parser.parse_args()
    workbook = Path(args.workbook)
    rows = read_rows(workbook)
    accepted = extract(rows)
    skus = [row["sku"] for row in accepted]
    if len(skus) != 177 or len(set(skus)) != len(skus):
        raise ValueError(f"EXPECTED_177_UNIQUE_ACCEPTED_ROWS:{len(skus)}:{len(set(skus))}")
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as handle:
        for row in accepted:
            payload = {
                "messages": [
                    {"role": "system", "content": "将 Action 西语商品字段标准化为简洁、忠实的中文。保持数字、单位、数量、尺寸和品牌/型号事实；不得臆造源数据没有的信息。输出固定六字段 JSON。"},
                    {"role": "user", "content": json.dumps(row["source"], ensure_ascii=False, sort_keys=True)},
                    {"role": "assistant", "content": json.dumps(row["target"], ensure_ascii=False, sort_keys=True)},
                ],
                "metadata": {
                    "sku": row["sku"], "source_hash": row["source_hash"], "ai_disposition": "ACCEPT_AS_GOLD",
                    "gold_status": "AI_ACCEPTED_OWNER_PENDING", "training_eligible": False,
                    "ai_reason": row["ai_reason"],
                },
            }
            handle.write(json.dumps(payload, ensure_ascii=False, sort_keys=True) + "\n")
    manifest = {
        "report_version": "stage4-ai-accepted-owner-pending-v1",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "source_workbook": {"path": str(workbook), "sha256": sha(workbook)},
        "rows": len(accepted), "unique_skus": len(set(skus)),
        "ai_disposition": "ACCEPT_AS_GOLD",
        "gold_rows": 0, "owner_confirmation_required": True,
        "training_eligible_rows": 0, "production_writes": False,
        "output": {"path": str(output), "sha256": sha(output)},
    }
    manifest_path = Path(args.manifest)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"rows": len(accepted), "unique_skus": len(set(skus)), "gold_rows": 0, "training_eligible_rows": 0}, ensure_ascii=False))


if __name__ == "__main__":
    main()
