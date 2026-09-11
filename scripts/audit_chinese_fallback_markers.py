from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
path = root / "runtime" / "exports" / "20260907Action商品全量_中文版_不带图.xlsx"
ws = load_workbook(path, read_only=True, data_only=True).active
rows = list(ws.iter_rows(values_only=True))
h = {str(v): i for i, v in enumerate(rows[0])}
remark = next(i for k, i in h.items() if k == "备注")
counter = Counter()
skus = set()
for r in rows[1:]:
    text = str(r[remark] or "")
    if "待审核" in text:
        skus.add(str(r[1]))
        for part in text.split("|"):
            if "待审核" in part:
                counter[part.strip()] += 1
labels = ("中文品名待审核", "中文分类1待审核", "中文分类2待审核", "中文规格待审核", "中文描述待审核", "中文产品详情待审核")
per_label = {label: sum(label in str(r[remark] or "") for r in rows[1:]) for label in labels}
print({"sku_with_fallback_marker": len(skus), "per_label": per_label})
