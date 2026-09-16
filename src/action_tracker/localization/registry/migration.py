from __future__ import annotations

import csv
import hashlib
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping


REQUIRED_SOURCE_COLUMNS = ("source_hash",)
APPROVED_STATUSES = {"APPROVE", "APPROVED", "HUMAN_APPROVED", "LOCKED", "GOLD", "CONTEXT_ONLY"}
REJECTED_STATUSES = {"REJECT", "REJECTED", "SOURCE_CONFLICT", "SOURCE_DAMAGED", "BLOCKED"}


def _sha(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def _read_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.casefold()
    if suffix in {".csv", ".tsv"}:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [dict(row) for row in csv.DictReader(handle, delimiter="\t" if suffix == ".tsv" else ",")]
    if suffix in {".json", ".jsonl"}:
        if suffix == ".jsonl":
            return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, list):
            return [dict(row) for row in payload if isinstance(row, Mapping)]
        if isinstance(payload, Mapping):
            rows = payload.get("rows") or payload.get("records") or payload.get("candidates") or []
            return [dict(row) for row in rows if isinstance(row, Mapping)]
        return []
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value or "").strip() for value in values[0]]
        return [{headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]} for row in values[1:]]
    raise ValueError(f"MIGRATION_INPUT_UNSUPPORTED:{path.suffix}")


def _text(row: Mapping[str, Any], *names: str) -> str:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _normalize_row(row: Mapping[str, Any], source_path: Path, ordinal: int) -> dict[str, Any]:
    shadow_scope = _text(row, "shadow_scope", "decision_scope", "scope")
    sku = _text(row, "sku", "SKU", "official_sku", "编号")
    if not sku and "GLOBAL" in shadow_scope.upper():
        sku = "__GLOBAL__"
    elif not sku and "CONTEXT" in shadow_scope.upper():
        sku = "__CONTEXT__"
    field = _text(row, "field_name", "field", "字段", "字段名")
    source_hash = _text(row, "source_hash", "reviewed_source_hash", "sourceHash")
    target = _text(row, "target_text", "target_value", "zh_value", "中文", "new_value", "suggested_value")
    status = _text(row, "status", "owner_decision", "approval_status", "review_status", "decision", "结论").upper() or "PENDING"
    if status == "OWNER_DECISION_RECORDED":
        status = "CONTEXT_ONLY" if "CONTEXT" in shadow_scope.upper() else ("REJECTED" if "NEVER_MATCH" in shadow_scope.upper() else "APPROVED")
    source_text = _text(row, "source_text", "source_value", "source", "西语")
    source_version = _text(row, "source_version_id", "source_run_id", "run_id")
    return {
        "source_file": str(source_path), "source_row": ordinal, "sku": sku,
        "field_name": field, "source_hash": source_hash, "target_text": target,
        "status": status, "source_text": source_text, "source_version_id": source_version,
        "shadow_scope": shadow_scope,
    }


def _expand_row(row: Mapping[str, Any]) -> list[Mapping[str, Any]]:
    """Expand owner-reviewed Gold records with proposed_zh/source objects."""
    proposed = row.get("proposed_zh")
    source = row.get("source")
    disposition = row.get("owner_disposition")
    if isinstance(proposed, Mapping) and isinstance(source, Mapping):
        result = []
        for field_name, target in proposed.items():
            decision = str(disposition.get(field_name, "") if isinstance(disposition, Mapping) else "").upper()
            if decision in {"REJECT", "REVIEW", "HOLD", "SOURCE_CONFLICT"}:
                continue
            result.append({"sku": row.get("sku"), "field_name": field_name, "source_hash": row.get("source_hash"), "target_text": target, "status": "GOLD", "source_text": source.get(field_name), "source_version_id": row.get("source_run_id")})
        return result
    return [row]


def build_migration_preview(input_paths: Iterable[Path], output_dir: Path, *, baseline_name: str = "origin-main") -> dict[str, Any]:
    """Build an auditable, no-write migration preview for TM/Gold/Patch files."""
    paths = [Path(path) for path in input_paths]
    rows: list[dict[str, Any]] = []
    input_manifest = []
    for path in paths:
        if not path.exists():
            raise FileNotFoundError(path)
        input_manifest.append({"path": str(path), "sha256": _sha(path), "size": path.stat().st_size})
        ordinal = 1
        for raw_row in _read_rows(path):
            for row in _expand_row(raw_row):
                ordinal += 1
                rows.append(_normalize_row(row, path, ordinal))

    eligible: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        reason = ""
        if not row["sku"] or not row["field_name"]:
            reason = "IDENTITY_MISSING"
        elif not row["source_hash"]:
            reason = "SOURCE_HASH_MISSING"
        elif not row["target_text"]:
            reason = "TARGET_EMPTY"
        elif row["status"] in REJECTED_STATUSES:
            reason = f"STATUS_{row['status']}"
        elif row["status"] not in APPROVED_STATUSES:
            reason = f"STATUS_NOT_APPROVED_{row['status']}"
        if reason:
            rejected.append({**row, "rejection_reason": reason})
            continue
        grouped[(row["sku"], row["field_name"], row["source_hash"])].append(row)

    for key, group in grouped.items():
        targets = {row["target_text"] for row in group}
        if len(targets) > 1:
            for row in group:
                conflicts.append({**row, "conflict_reason": "MULTIPLE_APPROVED_TARGETS_FOR_SOURCE_VERSION"})
            continue
        winner = dict(group[0])
        winner["migration_status"] = "ELIGIBLE"
        winner["context_count"] = len(group)
        eligible.append(winner)

    output_dir.mkdir(parents=True, exist_ok=True)
    headers = ["source_file", "source_row", "sku", "field_name", "source_hash", "target_text", "status", "source_text", "source_version_id", "migration_status", "context_count", "rejection_reason", "conflict_reason"]
    for name, data in (("eligible.csv", eligible), ("rejected.csv", rejected), ("conflicts.csv", conflicts)):
        path = output_dir / name
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
            writer.writeheader(); writer.writerows(data)

    manifest = {
        "schema_version": "TRANSLATION_MIGRATION_PREVIEW_V1",
        "baseline": baseline_name,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input_files": input_manifest,
        "input_row_count": len(rows),
        "eligible_count": len(eligible), "rejected_count": len(rejected), "conflict_count": len(conflicts),
        "production_writes": False, "sqlite_writes": False, "master_writes": False, "dictionary_writes": False,
        "outputs": {name: str(output_dir / name) for name in ("eligible.csv", "rejected.csv", "conflicts.csv")},
    }
    manifest["manifest_hash"] = hashlib.sha256(json.dumps(manifest, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest
